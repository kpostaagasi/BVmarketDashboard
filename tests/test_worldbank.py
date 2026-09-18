"""Dünya Bankası Pink Sheet (`ingest/worldbank.py`) istemcisi testleri."""

from __future__ import annotations

import io
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.worldbank import (
    LANDING_SAYFASI,
    SERI_TANIMLARI,
    _normalize,
    _sutun_indeksi,
    en_guncel_dosya_url,
    seri_cek,
)


def wb_seri(**kwargs):
    varsayilan = dict(id="emtia-tarim/x", wb_seri="gubre-endeksi", start_date=None)
    varsayilan.update(kwargs)
    return SimpleNamespace(**varsayilan)


# --- _normalize: dipnot/boşluk temizleme ---


def test_normalize_dipnot_yildizlarini_ve_bosluklari_temizler():
    assert _normalize("Fertilizers **") == "fertilizers"
    assert _normalize("Urea ") == "urea"
    assert _normalize("DAP") == "dap"


# --- _sutun_indeksi ---


def _kitap_baytlari(sayfalar: dict) -> bytes:
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    for ad, satirlar in sayfalar.items():
        ws = kitap.create_sheet(ad)
        for satir in satirlar:
            ws.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


def _kitabi_ac(sayfalar: dict):
    return openpyxl.load_workbook(io.BytesIO(_kitap_baytlari(sayfalar)), data_only=True)


def test_sutun_indeksi_dogru_sutunu_bulur():
    ws = _kitabi_ac({"Monthly Indices": [
        ["World Bank Commodity Price Data"],
        [None, "Total Index", "Energy"],
        [None, None, None, "Fertilizers **"],
        ["2026M08", 100.0, 200.0, 146.4],
    ]})["Monthly Indices"]
    assert _sutun_indeksi(ws, "Fertilizers") == 3


def test_sutun_indeksi_bulunamazsa_hata():
    ws = _kitabi_ac({"S": [["baslik"], [None, "Başka Sütun"]]})["S"]
    with pytest.raises(RuntimeError, match="Fertilizers"):
        _sutun_indeksi(ws, "Fertilizers")


# --- en_guncel_dosya_url ---


class SahteYanit:
    def __init__(self, content=b"", text=None, status_code=200):
        self.content = content
        self.text = text if text is not None else content.decode("utf-8", "replace")
        self.status_code = status_code


class SahteOturum:
    def __init__(self, yanit_haritasi):
        self.yanit_haritasi = yanit_haritasi
        self.cagrilar = []

    def get(self, url, timeout=None, headers=None):
        self.cagrilar.append(url)
        return self.yanit_haritasi[url]


def test_en_guncel_dosya_url_baglantiyi_bulur():
    html = '<a href="https://thedocs.worldbank.org/x/related/CMO-Historical-Data-Monthly.xlsx">İndir</a>'
    oturum = SahteOturum({LANDING_SAYFASI: SahteYanit(text=html)})
    url = en_guncel_dosya_url(session=oturum)
    assert url == "https://thedocs.worldbank.org/x/related/CMO-Historical-Data-Monthly.xlsx"


def test_en_guncel_dosya_url_baglanti_yoksa_hata():
    oturum = SahteOturum({LANDING_SAYFASI: SahteYanit(text="<html>boş</html>")})
    with pytest.raises(RuntimeError, match="bağlantısı bulunamadı"):
        en_guncel_dosya_url(session=oturum)


def test_en_guncel_dosya_url_http_hatasi_yukselir():
    oturum = SahteOturum({LANDING_SAYFASI: SahteYanit(status_code=500)})
    with pytest.raises(RuntimeError, match="HTTP 500"):
        en_guncel_dosya_url(session=oturum)


# --- seri_cek: ağ kabuğu + önbellek ---


DOSYA_URL = "https://thedocs.worldbank.org/x/related/CMO-Historical-Data-Monthly.xlsx"

_PINK_SHEET_SAYFALARI = {
    "Monthly Indices": [
        ["World Bank Commodity Price Data"],
        [None, "Total Index"],
        [None, None, "Fertilizers **"],
        ["2026M07", 100.0, 140.0],
        ["2026M08", 101.0, 146.4],
    ],
    "Monthly Prices": [
        ["Updated on September 2026"],
        [None],
        [None, "Crude oil", "DAP", "Urea "],
        ["2026M07", 60.0, 781.3, 400.0],
        ["2026M08", 61.0, 793.5, 390.0],
    ],
}


def _pink_sheet_oturumu(landing_html: str | None = None) -> SahteOturum:
    html = landing_html or f'<a href="{DOSYA_URL}">CMO-Historical-Data-Monthly.xlsx</a>'
    return SahteOturum({
        LANDING_SAYFASI: SahteYanit(text=html),
        DOSYA_URL: SahteYanit(content=_kitap_baytlari(_PINK_SHEET_SAYFALARI)),
    })


def test_seri_cek_gubre_endeksi_dogru_deger_dondurur():
    oturum = _pink_sheet_oturumu()
    df = seri_cek(wb_seri(wb_seri="gubre-endeksi"), onbellek={}, session=oturum)
    assert df.iloc[-1]["date"] == "2026-08-01"
    assert df.iloc[-1]["value"] == pytest.approx(146.4)


def test_seri_cek_urea_ve_dap_farkli_sayfadan_dogru_okur():
    oturum = _pink_sheet_oturumu()
    dap = seri_cek(wb_seri(wb_seri="dap"), onbellek={}, session=oturum)
    urea = seri_cek(wb_seri(wb_seri="urea"), onbellek={}, session=oturum)
    assert dap.iloc[-1]["value"] == pytest.approx(793.5)
    assert urea.iloc[-1]["value"] == pytest.approx(390.0)


def test_seri_cek_onbellegi_paylasir():
    oturum = _pink_sheet_oturumu()
    onbellek: dict = {}
    seri_cek(wb_seri(wb_seri="gubre-endeksi"), onbellek=onbellek, session=oturum)
    cagri_sayisi = len(oturum.cagrilar)
    seri_cek(wb_seri(wb_seri="urea"), onbellek=onbellek, session=oturum)
    assert len(oturum.cagrilar) == cagri_sayisi  # ikinci çağrı ağa çıkmamalı


def test_seri_cek_start_date_oncesini_kirpar():
    oturum = _pink_sheet_oturumu()
    df = seri_cek(wb_seri(wb_seri="gubre-endeksi", start_date="2026-08-01"), onbellek={}, session=oturum)
    assert list(df["date"]) == ["2026-08-01"]


def test_seri_cek_bilinmeyen_wb_serisinde_hata():
    oturum = _pink_sheet_oturumu()
    with pytest.raises(RuntimeError, match="Bilinmeyen wb_seri"):
        seri_cek(wb_seri(wb_seri="olmayan-seri"), onbellek={}, session=oturum)


def test_seri_tanimlari_katalog_enumuyla_birebir():
    from core.catalog import GECERLI_WB_SERILERI
    assert set(SERI_TANIMLARI.keys()) == GECERLI_WB_SERILERI
