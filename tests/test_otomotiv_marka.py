"""Otomotiv marka/dağıtıcı satış serileri: OSD araç-tipi/ihracat kırılımı
(`ingest/osd.py` ekleri) ve ODMD marka bazlı perakende (`ingest/odmd.py`)
testleri.

`firma_aylik_noktalari_arac_tipli`/`firma_bazli_dis_satis_tablolarindan`/
`marka_satirlari` PDF/XLSX'i pdfplumber/openpyxl'in kendi çıktısıyla aynı
biçimdeki DÜZ PYTHON YAPILARI ile test edilir (gerçek dosya yok) — tıpkı
`tests/test_osd.py`'nin `firma_aylik_noktalari`'yi test ettiği gibi. Ağ
kabuğu (`seri_cek`) testleri `monkeypatch` ile HTTP/PDF/XLSX katmanını
atlar. Değerler 2026-09-18'de canlı ölçülen gerçek OSD/ODMD verileriyle
birebir doğrulandı (bkz. yield raporu); burada yalnızca ayrıştırma/hesap
mantığı pinlenir.
"""

from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest

from ingest import odmd
from ingest.odmd import (
    _deger,
    _satirlari_ayikla,
    marka_satirlari,
    rapor_id_eslemesi,
)
from ingest.osd import (
    DIS_SATIS_TIP_SIRASI,
    degerlendirme_baglantilari,
    firma_aylik_noktalari_arac_tipli,
    firma_bazli_dis_satis_tablolarindan,
)
import ingest.osd as osd

# --- OSD: firma_aylik_noktalari_arac_tipli (üretim, araç tipi kırılımı) ---


def _tip_tablosu(*bolumler):
    """Her bölüm `(tip_toplam_satir_adi, [firma_satirlari])`; 14 sütun
    (tip/firma + 12 ay + toplam), tıpkı `tests/test_osd.py::_tablo` gibi."""
    basliklar = ["Tipler\nTypes"] + [f"AY{i}" for i in range(1, 13)] + ["TOPLAM"]
    satirlar = [basliklar]
    for toplam_adi, firmalar in bolumler:
        satirlar.extend(firmalar)
        satirlar.append([toplam_adi] + ["-"] * 13)
    return satirlar


def _firma_satiri(ad, *aylar):
    hucreler = list(aylar) + ["-"] * (12 - len(aylar))
    return [ad, *hucreler, "-"]


def test_firma_aylik_noktalari_arac_tipli_tip_bilgisini_korur():
    tablo = _tip_tablosu(
        ("K.KAMYON Toplam / L. Truck Total", [_firma_satiri("A.I.O.S.", "172")]),
        ("KAMYONET Toplam / Pick Up Total", [_firma_satiri("TOFAŞ", "5065")]),
    )
    noktalar = firma_aylik_noktalari_arac_tipli([tablo], 2026)
    assert ("A.I.O.S.", "K.KAMYON", "2026-01-01", 172.0) in noktalar
    assert ("TOFAŞ", "KAMYONET", "2026-01-01", 5065.0) in noktalar
    assert len(noktalar) == 2


def test_firma_aylik_noktalari_arac_tipli_firma_ayni_ayda_iki_bolumde_ayri_tutulur():
    """Bir firma iki farklı araç tipinde üretebilir; TOPLANMAZ, iki AYRI
    (firma, tip) noktası üretir — `firma_aylik_noktalari`nin aksine."""
    tablo = _tip_tablosu(
        ("OTOMOBİL Toplam / Pass.Car Total", [_firma_satiri("TOFAŞ", "4")]),
        ("KAMYONET Toplam / Pick Up Total", [_firma_satiri("TOFAŞ", "5065")]),
    )
    noktalar = firma_aylik_noktalari_arac_tipli([tablo], 2026)
    assert set(noktalar) == {
        ("TOFAŞ", "OTOMOBİL", "2026-01-01", 4.0),
        ("TOFAŞ", "KAMYONET", "2026-01-01", 5065.0),
    }


def test_firma_aylik_noktalari_arac_tipli_bos_ayi_atlar():
    tablo = _tip_tablosu(
        ("KAMYONET Toplam / Pick Up Total", [_firma_satiri("A.I.O.S.", "-", "5")]),
    )
    noktalar = firma_aylik_noktalari_arac_tipli([tablo], 2026)
    assert noktalar == [("A.I.O.S.", "KAMYONET", "2026-02-01", 5.0)]


def test_firma_aylik_noktalari_arac_tipli_pay_satirlarini_atlar():
    tablo = [
        ["Tipler\nTypes"] + [f"AY{i}" for i in range(1, 13)] + ["TOPLAM"],
        _firma_satiri("A.I.O.S.", "172"),
        ["Pay / Share %"] + ["-"] * 13,
        ["K.KAMYON Toplam / L. Truck Total"] + ["-"] * 13,
        ["Toplam Pay / Total Share"] + ["-"] * 13,
    ]
    noktalar = firma_aylik_noktalari_arac_tipli([tablo], 2026)
    assert noktalar == [("A.I.O.S.", "K.KAMYON", "2026-01-01", 172.0)]


# --- OSD: Değerlendirme Raporu indeksi (ihracat kaynağı) ---

_DEGERLENDIRME_INDEKS_HTML = """
<a href="/saved-files\\PDF\\2026\\09\\14\\08-2026-OSD_Aylik_Degerlendirme_Raporu.pdf">AĞUSTOS 2026</a>
<a href="/saved-files\\PDF\\2023\\01\\16\\12-2022_OSD_Aylik_Degerlendirme_Raporu.pdf">2022 ARALIK</a>
"""


def test_degerlendirme_baglantilari_tire_ve_alt_cizgiyi_kabul_eder():
    b = degerlendirme_baglantilari(_DEGERLENDIRME_INDEKS_HTML)
    assert b["2026.08"].endswith("08-2026-OSD_Aylik_Degerlendirme_Raporu.pdf")
    assert b["2022.12"].endswith("12-2022_OSD_Aylik_Degerlendirme_Raporu.pdf")


def test_degerlendirme_baglantilari_ters_boluyu_duzeltir():
    b = degerlendirme_baglantilari(_DEGERLENDIRME_INDEKS_HTML)
    assert "\\" not in b["2026.08"]
    assert b["2026.08"].startswith("https://www.osd.org.tr/saved-files/PDF/")


def test_degerlendirme_baglantilari_hicbiri_yoksa_yukselir():
    with pytest.raises(RuntimeError, match="Değerlendirme Raporu"):
        degerlendirme_baglantilari("<html><body>hiçbir şey</body></html>")


# --- OSD: Dış Satışlar (ihracat) firma tabloları ---


def _firmalar_tablosu(*satirlar, toplam):
    basliklar = ["Firmalar / Firms"] + [str(i) for i in range(1, 13)] + ["Toplam / Total"]
    gövde = [[ad, *degerler, "-"] for ad, *degerler in satirlar]
    return [basliklar, *gövde, ["TOPLAM - Total", *toplam, "-"]]


def _yedi_tablo(degistir: dict[str, list]) -> list:
    """`DIS_SATIS_TIP_SIRASI` sırasıyla 7 boş (tek firmalı, tümü 0) tablo
    üretir; `degistir` verilen tip(ler)i override eder."""
    varsayilan = {
        tip: _firmalar_tablosu(("FIRMA", *(["0"] * 12)), toplam=["0"] * 12)
        for tip in DIS_SATIS_TIP_SIRASI
    }
    varsayilan.update(degistir)
    return [varsayilan[tip] for tip in DIS_SATIS_TIP_SIRASI]


def test_firma_bazli_dis_satis_tablolarindan_tip_atar_ve_toplar():
    tablolar = _yedi_tablo({
        "OTOBÜS": _firmalar_tablosu(("A.I.O.S.", "19", *(["0"] * 11)), toplam=["19"] + ["0"] * 11),
    })
    noktalar = firma_bazli_dis_satis_tablolarindan(tablolar, 2026)
    assert ("A.I.O.S.", "OTOBÜS", "2026-01-01", 19.0) in noktalar


def test_firma_bazli_dis_satis_tablolarindan_ozet_doymayan_ayi_atlar():
    tablolar = _yedi_tablo({
        "MİDİBÜS": _firmalar_tablosu(("A.I.O.S.", "-", "44", *(["-"] * 10)), toplam=["0", "44"] + ["0"] * 10),
    })
    noktalar = firma_bazli_dis_satis_tablolarindan(tablolar, 2026)
    midibus_noktalari = [n for n in noktalar if n[0] == "A.I.O.S." and n[1] == "MİDİBÜS"]
    assert midibus_noktalari == [("A.I.O.S.", "MİDİBÜS", "2026-02-01", 44.0)]


def test_firma_bazli_dis_satis_tablolarindan_toplam_tutmazsa_yukselir():
    """Firma satırları TOPLAM satırıyla uyuşmuyorsa (şablon kayması) hata
    verilmeli — sessizce yanlış veri üretilmemeli."""
    tablolar = _yedi_tablo({
        "OTOBÜS": _firmalar_tablosu(("A.I.O.S.", "19", *(["0"] * 11)), toplam=["999"] + ["0"] * 11),
    })
    with pytest.raises(RuntimeError, match="öz-doğrulama"):
        firma_bazli_dis_satis_tablolarindan(tablolar, 2026)


def test_firma_bazli_dis_satis_tablolarindan_tablo_sayisi_yanlissa_yukselir():
    eksik = _yedi_tablo({})[:6]  # yalnızca 6 tablo — beklenen 7
    with pytest.raises(RuntimeError, match="7"):
        firma_bazli_dis_satis_tablolarindan(eksik, 2026)


# --- OSD: seri_cek ağ kabuğu (osd_arac_tipi / osd_veri_tipi dispatch) ---


def test_seri_cek_arac_tipi_verilmezse_mevcut_davranis_degismez(monkeypatch):
    """`osd_arac_tipi`/`osd_veri_tipi` taşımayan (mevcut 13 seri gibi) bir
    SimpleNamespace ile çağrıldığında hâlâ eski toplam-yolu çalışmalı —
    `getattr(..., None)` geriye dönük uyumluluğu garanti eder."""
    monkeypatch.setattr(osd, "_indeks_cek", lambda session=None: {"2026.07": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "_bulteni_ayristir",
        lambda baytlar, anahtar: [("TOFAŞ", "2026-01-01", 5069.0)],
    )
    seri = SimpleNamespace(id="otomotiv/tofas", osd_firma="TOFAŞ", start_date=None)
    df = osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 4))
    assert df["value"].iloc[0] == 5069.0


def test_seri_cek_arac_tipi_verilince_tip_filtreler(monkeypatch):
    monkeypatch.setattr(osd, "_indeks_cek", lambda session=None: {"2026.07": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "_bultendeki_tipli_noktalar",
        lambda baytlar, anahtar: [
            ("A.I.O.S.", "K.KAMYON", "2026-01-01", 172.0),
            ("A.I.O.S.", "OTOBÜS", "2026-01-01", 36.0),
        ],
    )
    seri = SimpleNamespace(
        id="otomotiv/asuzu-hafif-kamyon-uretim", osd_firma="A.I.O.S.",
        osd_arac_tipi="K.KAMYON", osd_veri_tipi=None, osd_eski_adlar=None,
        start_date=None,
    )
    df = osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 4))
    assert list(df["value"]) == [172.0]


def test_seri_cek_arac_tipi_firma_hic_yoksa_yukselir(monkeypatch):
    monkeypatch.setattr(osd, "_indeks_cek", lambda session=None: {"2026.07": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "_bultendeki_tipli_noktalar",
        lambda baytlar, anahtar: [("A.I.O.S.", "K.KAMYON", "2026-01-01", 172.0)],
    )
    seri = SimpleNamespace(
        id="otomotiv/yok", osd_firma="YOK BÖYLE FİRMA", osd_arac_tipi="OTOBÜS",
        osd_veri_tipi=None, osd_eski_adlar=None, start_date=None,
    )
    with pytest.raises(RuntimeError, match="YOK BÖYLE FİRMA"):
        osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 4))


def test_seri_cek_ihracat_firma_toplamini_tiplerde_toplar(monkeypatch):
    monkeypatch.setattr(osd, "_dis_satis_indeks_cek", lambda session=None: {"2026.08": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "dis_satis_noktalari",
        lambda baytlar: [
            ("A.I.O.S.", "OTOBÜS", "2026-08-01", 19.0),
            ("A.I.O.S.", "MİDİBÜS", "2026-08-01", 44.0),
        ],
    )
    seri = SimpleNamespace(
        id="otomotiv/asuzu-toplam-ihracat", osd_firma="A.I.O.S.", osd_arac_tipi=None,
        osd_veri_tipi="ihracat", osd_eski_adlar=None, start_date=None,
    )
    df = osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 18))
    assert df["value"].iloc[0] == pytest.approx(63.0)


def test_seri_cek_ihracat_tek_tip_filtreler(monkeypatch):
    monkeypatch.setattr(osd, "_dis_satis_indeks_cek", lambda session=None: {"2026.08": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "dis_satis_noktalari",
        lambda baytlar: [
            ("OTOKAR", "KAMYON", "2026-08-01", 16.0),
            ("OTOKAR", "OTOBÜS", "2026-08-01", 113.0),
        ],
    )
    seri = SimpleNamespace(
        id="otomotiv/otokar-kamyon-ihracat", osd_firma="OTOKAR", osd_arac_tipi="KAMYON",
        osd_veri_tipi="ihracat", osd_eski_adlar=None, start_date=None,
    )
    df = osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 18))
    assert list(df["value"]) == [16.0]


def test_seri_cek_ihracat_firma_hic_yoksa_yukselir(monkeypatch):
    monkeypatch.setattr(osd, "_dis_satis_indeks_cek", lambda session=None: {"2026.08": "u"})
    monkeypatch.setattr(osd, "_pdf_indir", lambda url, session=None: b"x")
    monkeypatch.setattr(
        osd, "dis_satis_noktalari",
        lambda baytlar: [("OTOKAR", "KAMYON", "2026-08-01", 16.0)],
    )
    seri = SimpleNamespace(
        id="otomotiv/yok", osd_firma="YOK BÖYLE FİRMA", osd_arac_tipi=None,
        osd_veri_tipi="ihracat", osd_eski_adlar=None, start_date=None,
    )
    with pytest.raises(RuntimeError, match="YOK BÖYLE FİRMA"):
        osd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 18))


# ================= ODMD =================

# --- Liste sayfası ayrıştırma ---

_ODMD_LISTE_HTML = """
<table id="TableReports">
<tbody><tr class="tr_r_rep">
<td><input type="radio" name="report" onclick="sRep(6163)" class="radio_reports"></td>
<td class="td_reports_name" onclick="return dRep('?primary_id=6163')">
2026 Yılı (Ocak-Ağustos) Perakende Satışlar (Yerli &amp; İthal)
</td>
</tr>
<tr class="tr_r_sep"><td></td></tr><tr class="tr_r_rep">
<td><input type="radio" name="report" onclick="sRep(6162)" class="radio_reports"></td>
<td class="td_reports_name" onclick="return dRep('?primary_id=6162')">
2026 Ağustos Perakende Satışlar (Yerli &amp; İthal)
</td>
</tr>
</tbody></table>
"""


def test_satirlari_ayikla_yillik_kumulatif_satiri_atlar():
    satirlar = _satirlari_ayikla(_ODMD_LISTE_HTML)
    assert satirlar == [("6162", "2026.08")]


def test_rapor_id_eslemesi_birden_cok_sayfayi_birlestirir():
    sayfa2 = _ODMD_LISTE_HTML.replace("6162", "6147").replace("Ağustos", "Temmuz")
    eslesme = rapor_id_eslemesi([_ODMD_LISTE_HTML, sayfa2])
    assert eslesme == {"2026.08": "6162", "2026.07": "6147"}


def test_rapor_id_eslemesi_hicbiri_yoksa_yukselir():
    with pytest.raises(RuntimeError, match="ODMD rapor listesi"):
        rapor_id_eslemesi(["<html></html>"])


# --- XLSX ayrıştırma ---


def _sahte_odmd_xlsx(satirlar, marka_hucresi="MARKA") -> bytes:
    """`marka_satirlari`nin beklediği sayfa şablonunu üretir (openpyxl ile
    bellekte, dosyaya yazmadan) — gerçek ODMD dosyasının 2026-09-18'de
    ölçülen sütun düzenini (Oto/HT/Toplam × Yerli/İthal/Toplam) birebir
    taşır."""
    import io as _io

    kitap = openpyxl.Workbook()
    sayfa = kitap.active
    sayfa.append(["BAŞLIK", None, None, None, None, None, None, None, None, None])
    sayfa.append([marka_hucresi, "OTOMOBİL", None, None, "HAFİF TİCARİ", None, None, "TOPLAM", None, None])
    sayfa.append([None, "YERLİ", "İTHAL", "TOPLAM", "YERLİ", "İTHAL", "TOPLAM", "YERLİ", "İTHAL", "TOPLAM"])
    sayfa.append([None] * 10)
    for satir in satirlar:
        sayfa.append(satir)
    toplam = [sum(s[i] for s in satirlar) for i in (3, 6, 9)]
    sayfa.append(["TOPLAM:", None, None, toplam[0], None, None, toplam[1], None, None, toplam[2]])
    buf = _io.BytesIO()
    kitap.save(buf)
    return buf.getvalue()


def test_marka_satirlari_toplam_kolonlarini_okur():
    baytlar = _sahte_odmd_xlsx([
        ["FIAT", 1905, 220, 2125, 630, 3424, 4054, 2535, 3644, 6179],
        ["AUDI", 0, 1275, 1275, 0, 0, 0, 0, 1275, 1275],
    ])
    markalar = marka_satirlari(baytlar)
    assert markalar["FIAT"] == {"otomobil": 2125.0, "hafif_ticari": 4054.0, "toplam": 6179.0}
    assert markalar["AUDI"]["toplam"] == 1275.0


def test_marka_satirlari_marka_etiketi_bos_olsa_da_calisir():
    """Eski dosyalarda (2021 öncesi) MARKA hücresinin metni boş — yapısal
    olarak OTOMOBİL/HAFİF TİCARİ/TOPLAM sütun üçlüsüne bakılır."""
    baytlar = _sahte_odmd_xlsx([["FIAT", 1905, 220, 2125, 630, 3424, 4054, 2535, 3644, 6179]],
                                marka_hucresi=None)
    markalar = marka_satirlari(baytlar)
    assert markalar["FIAT"]["toplam"] == 6179.0


def test_marka_satirlari_toplam_tutmazsa_yukselir():
    baytlar = _sahte_odmd_xlsx([["FIAT", 1905, 220, 2125, 630, 3424, 4054, 2535, 3644, 6179]])
    # TOPLAM satırını bozmak için sahte dosyayı yeniden aç ve değiştir.
    import io as _io

    kitap = openpyxl.load_workbook(_io.BytesIO(baytlar))
    sayfa = kitap.active
    sayfa.cell(row=sayfa.max_row, column=10, value=99999)
    buf = _io.BytesIO()
    kitap.save(buf)
    with pytest.raises(RuntimeError, match="öz-doğrulama"):
        marka_satirlari(buf.getvalue())


# --- _deger: tam/yarım ağırlık toplama ---


def test_deger_birden_cok_markayi_toplar():
    seri = SimpleNamespace(odmd_marka=("FIAT", "PEUGEOT"), odmd_yarim_marka=None, odmd_kategori="toplam")
    markalar = {"FIAT": {"toplam": 6179.0}, "PEUGEOT": {"toplam": 4775.0}}
    assert _deger(seri, markalar) == pytest.approx(10954.0)


def test_deger_yarim_agirlikli_markayi_yarilar():
    seri = SimpleNamespace(odmd_marka=("VOLKSWAGEN",), odmd_yarim_marka=("SKODA",), odmd_kategori="toplam")
    markalar = {"VOLKSWAGEN": {"toplam": 7420.0}, "SKODA": {"toplam": 3055.0}}
    assert _deger(seri, markalar) == pytest.approx(7420.0 + 0.5 * 3055.0)


def test_deger_marka_dosyada_yoksa_sifir_sayilir():
    """Marka o ayın dosyasında hiç yoksa (henüz pazara girmemiş, ör. CUPRA
    2021 öncesi) katkısı 0 sayılır — hata değil."""
    seri = SimpleNamespace(odmd_marka=("FIAT", "CUPRA"), odmd_yarim_marka=None, odmd_kategori="toplam")
    markalar = {"FIAT": {"toplam": 6179.0}}
    assert _deger(seri, markalar) == pytest.approx(6179.0)


# --- ODMD: seri_cek ağ kabuğu ---


def test_odmd_seri_cek_onbellegi_paylasir(monkeypatch):
    indirilenler = []

    def sahte_indir(rapor_id, session=None):
        indirilenler.append(rapor_id)
        return b"sahte-xlsx"

    def sahte_ayristir(baytlar):
        return {"FIAT": {"otomobil": 2125.0, "hafif_ticari": 4054.0, "toplam": 6179.0}}

    monkeypatch.setattr(odmd, "_indeks_cek", lambda session=None: {"2026.08": "6162"})
    monkeypatch.setattr(odmd, "_dosya_indir", sahte_indir)
    monkeypatch.setattr(odmd, "marka_satirlari", sahte_ayristir)

    onbellek: dict = {}
    fiat_toplam = SimpleNamespace(id="otomotiv/toaso-fiat-toplam", odmd_marka=("FIAT",),
                                   odmd_yarim_marka=None, odmd_kategori="toplam", start_date=None)
    fiat_oto = SimpleNamespace(id="otomotiv/toaso-fiat-otomobil", odmd_marka=("FIAT",),
                                odmd_yarim_marka=None, odmd_kategori="otomobil", start_date=None)

    df1 = odmd.seri_cek(fiat_toplam, onbellek=onbellek, bugun=date(2026, 9, 18))
    df2 = odmd.seri_cek(fiat_oto, onbellek=onbellek, bugun=date(2026, 9, 18))

    assert len(indirilenler) == 1  # ikinci seri önbellekten
    assert df1["value"].iloc[0] == 6179.0
    assert df2["value"].iloc[0] == 2125.0


def test_odmd_seri_cek_kategori_disina_gecmis_yili_atlar(monkeypatch):
    monkeypatch.setattr(
        odmd, "_indeks_cek",
        lambda session=None: {"2026.08": "6162", "2015.03": "999"},
    )
    monkeypatch.setattr(odmd, "_dosya_indir", lambda rapor_id, session=None: b"x")
    monkeypatch.setattr(
        odmd, "marka_satirlari",
        lambda baytlar: {"FIAT": {"otomobil": 1.0, "hafif_ticari": 1.0, "toplam": 2125.0}},
    )
    seri = SimpleNamespace(id="otomotiv/toaso-fiat-toplam", odmd_marka=("FIAT",),
                            odmd_yarim_marka=None, odmd_kategori="toplam", start_date=None)
    df = odmd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 18))
    assert list(df["date"]) == ["2026-08-01"]  # 2015 VARSAYILAN_GECMIS_YIL dışında


def test_odmd_seri_cek_marka_hic_bulunamazsa_yukselir(monkeypatch):
    """Katalogdaki marka adı hiçbir ayda hiç görülmezse (yazım hatası)
    hata verilmeli."""
    monkeypatch.setattr(odmd, "_indeks_cek", lambda session=None: {"2026.08": "6162"})
    monkeypatch.setattr(odmd, "_dosya_indir", lambda rapor_id, session=None: b"x")
    monkeypatch.setattr(
        odmd, "marka_satirlari",
        lambda baytlar: {"FIAT": {"otomobil": 1.0, "hafif_ticari": 1.0, "toplam": 2.0}},
    )
    seri = SimpleNamespace(id="otomotiv/yok", odmd_marka=("YOK BÖYLE MARKA",),
                            odmd_yarim_marka=None, odmd_kategori="toplam", start_date=None)
    with pytest.raises(RuntimeError, match="YOK BÖYLE MARKA"):
        odmd.seri_cek(seri, onbellek={}, bugun=date(2026, 9, 18))
