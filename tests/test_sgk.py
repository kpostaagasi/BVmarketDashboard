"""SGK Aylık Sağlık İstatistik Bülteni istemcisi testleri."""

import io
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.sgk import (
    _ay_no,
    eczane_verisini_ayikla,
    hastane_verisini_ayikla,
    seri_cek,
)


def _hastane_kitabi(satirlar):
    """`satirlar`: [(yil_ya_da_None, ay_metni, m2, m3, mOzel, mUni, mToplam,
    f2, f3, fOzel, fUni, fToplam)] — Tablo 21 biçiminde bir sayfa üretir."""
    kitap = openpyxl.Workbook()
    ws = kitap.active
    ws.title = "21.SGK Hastane Aylar (Baş.Türü)"
    for _ in range(7):
        ws.append([None] * 17)
    ws.cell(4, 1, "TABLO 21- SAĞLIK HİZMET SUNUCUSU TÜRÜNE GÖRE FATURA BİLGİLERİ, 2012-2026")
    for satir in satirlar:
        ws.append(list(satir) + [None] * (17 - len(satir)))
    tampon = io.BytesIO()
    kitap.save(tampon)
    tampon.seek(0)
    return openpyxl.load_workbook(tampon, data_only=True, read_only=True)


def _eczane_kitabi(bloklar):
    """`bloklar`: [[(yil_ya_da_None, ay_metni, recete, fatura)]] — her alt liste
    bir 5-sütunluk bloğun satırlarıdır (bloklar yan yana yerleştirilir)."""
    kitap = openpyxl.Workbook()
    ws = kitap.active
    ws.title = "23.SGK-Kamu Reçete Aylar"
    for _ in range(6):
        ws.append([None] * (5 * len(bloklar)))
    uzunluk = max(len(b) for b in bloklar)
    for i in range(uzunluk):
        satir = []
        for blok in bloklar:
            if i < len(blok):
                yil, ay, recete, fatura = blok[i]
                satir += [yil, ay, recete, fatura, None]
            else:
                satir += [None] * 5
        ws.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    tampon.seek(0)
    return openpyxl.load_workbook(tampon, data_only=True, read_only=True)


# --- _ay_no: ay adı ayrıştırma ---


def test_ay_no_standart_bicim():
    assert _ay_no("Nisan - April") == 4


def test_ay_no_bosluksuz_tire():
    assert _ay_no("Kasım -November") == 11


def test_ay_no_taninmayan_metin_none_doner():
    assert _ay_no("GENEL TOPLAM\nGeneral Total") is None


# --- hastane_verisini_ayikla ---


def test_hastane_verisini_ayikla_dogru_sutunlari_okur():
    kitap = _hastane_kitabi([
        (2026, "Ocak - January", 100, 50, 20, 10, 180, 1000, 500, 200, 100, 1800),
        (None, "Şubat - February", 101, 51, 21, 11, 184, 1010, 510, 210, 110, 1840),
        (None, "GENEL TOPLAM\nGeneral Total", 201, 101, 41, 21, 364, 2010, 1010, 410, 210, 3640),
    ])
    sonuc = hastane_verisini_ayikla(kitap)
    assert sonuc[(2026, 1)]["muracaat"][4] == 20  # özel müracaat
    assert sonuc[(2026, 1)]["fatura"][11] == 1800  # toplam fatura
    assert (2026, "GENEL") not in sonuc  # toplam satırı ayrı bir dönem değil
    assert len(sonuc) == 2  # GENEL TOPLAM atlanmış


def test_hastane_verisini_ayikla_veri_yoksa_hata():
    kitap = _hastane_kitabi([])
    with pytest.raises(RuntimeError, match="hiç veri noktası"):
        hastane_verisini_ayikla(kitap)


# --- eczane_verisini_ayikla: çoklu blok, kronolojik olmayan yıl sırası ---


def test_eczane_verisini_ayikla_tek_blok():
    kitap = _eczane_kitabi([[
        (2026, "Nisan - April", 46066.56, 52107881.113),
        (None, "Mayıs - May", None, None),
    ]])
    sonuc = eczane_verisini_ayikla(kitap)
    assert sonuc[(2026, 4)] == {2: 46066.56, 3: 52107881.113}
    assert (2026, 5) not in sonuc  # boş hücreler atlanmış


def test_eczane_verisini_ayikla_paralel_bloklar_kronolojik_olmayan_sira():
    """Gerçek SGK dosyasında blok0=2015/2017/2019, blok1=2016/2018/2020 gibi
    çapraz yıllar olabilir — parser blok SIRASINI değil, her bloğun kendi
    YIL hücresini izlemeli."""
    kitap = _eczane_kitabi([
        [(2019, "Ocak - January", 10, 100)],
        [(2018, "Ocak - January", 20, 200)],
    ])
    sonuc = eczane_verisini_ayikla(kitap)
    assert sonuc[(2019, 1)] == {2: 10, 3: 100}
    assert sonuc[(2018, 1)] == {2: 20, 3: 200}


def test_eczane_verisini_ayikla_veri_yoksa_hata():
    kitap = _eczane_kitabi([[]])
    with pytest.raises(RuntimeError, match="hiç veri noktası"):
        eczane_verisini_ayikla(kitap)


# --- seri_cek: onbellek paylaşımı + metrik yönlendirme ---


class _SahteYanit:
    def __init__(self, json_data=None, content=b"", headers=None):
        self._json = json_data
        self.content = content
        self.headers = headers or {}

    def json(self):
        return self._json

    def raise_for_status(self):
        pass

    @property
    def text(self):
        return '<input name="__RequestVerificationToken" type="hidden" value="tok123" />'


class _SahteOturum:
    def __init__(self, xlsx_bytes):
        self.xlsx_bytes = xlsx_bytes
        self.post_cagrilari = 0
        self.get_cagrilari = 0

    def get(self, url, **kwargs):
        self.get_cagrilari += 1
        if "downloadfilestatistic" in url:
            return _SahteYanit(
                content=self.xlsx_bytes,
                headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            )
        return _SahteYanit()

    def post(self, url, **kwargs):
        self.post_cagrilari += 1
        if "YilListesi" in url:
            return _SahteYanit(json_data=[{"year": 2026}])
        return _SahteYanit(json_data=[{"month": 4, "monthName": "Nisan"}])


def _birlesik_kitap_bytes():
    kitap = openpyxl.Workbook()
    ws1 = kitap.active
    ws1.title = "21.SGK Hastane Aylar (Baş.Türü)"
    for _ in range(7):
        ws1.append([None] * 17)
    ws1.append([2026, "Nisan - April", 100, 50, 20, 10, 180, 1000, 500, 200, 100, 1800])
    ws2 = kitap.create_sheet("23.SGK-Kamu Reçete Aylar")
    for _ in range(6):
        ws2.append([None] * 5)
    ws2.append([2026, "Nisan - April", 46066.56, 52107881.113, None])
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


def test_seri_cek_hastane_metrigi():
    onbellek: dict = {}
    seri = SimpleNamespace(sgk_metrik="hastane-ozel-muracaat")
    df = seri_cek(seri, onbellek=onbellek, session=_SahteOturum(_birlesik_kitap_bytes()))
    assert list(df["date"]) == ["2026-04-01"]
    assert list(df["value"]) == [20]


def test_seri_cek_eczane_metrigi_ayni_onbellegi_paylasir():
    """6 SGK serisi (4 hastane + 2 eczane) TEK XLSX indirmesini paylaşır."""
    onbellek: dict = {}
    oturum = _SahteOturum(_birlesik_kitap_bytes())
    seri_cek(SimpleNamespace(sgk_metrik="hastane-ozel-muracaat"), onbellek=onbellek, session=oturum)
    seri_cek(SimpleNamespace(sgk_metrik="eczane-recete-sayisi"), onbellek=onbellek, session=oturum)
    assert oturum.get_cagrilari == 2  # sayfa (token) + tek xlsx indirmesi, iki serilik


def test_seri_cek_bilinmeyen_metrikte_hata():
    onbellek: dict = {}
    seri = SimpleNamespace(sgk_metrik="olmayan-metrik")
    with pytest.raises(RuntimeError, match="bilinmeyen sgk_metrik"):
        seri_cek(seri, onbellek=onbellek, session=_SahteOturum(_birlesik_kitap_bytes()))
