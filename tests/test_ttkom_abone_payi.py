"""Türk Telekom abone payı genişletmesi (`ingest/ttkom.py::_payi_serisi`)
testleri — mevcut `tests/test_telekom.py` bu yeni fonksiyonu kapsamıyor.

Mobil Faturalı Abone Payı ve Sabit Genişbant Fiber Abone Payı, "Abone
Verileri" sayfasındaki İKİ HAM alt küme/toplam abone sayısı alanından
hesaplanan gerçek oranlardır (tahmini değil) — bkz. modül docstring'i.
"""

import io
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.ttkom import ABONE_PAYI_KAYNAK, _payi_serisi, seri_cek


def _kitabi_ac(sayfalar: dict):
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    for ad, satirlar in sayfalar.items():
        sayfa = kitap.create_sheet(ad)
        for satir in satirlar:
            sayfa.append(list(satir))
    tampon = io.BytesIO()
    kitap.save(tampon)
    return openpyxl.load_workbook(io.BytesIO(tampon.getvalue()), data_only=True)


def _abone_verileri_sayfasi():
    return [
        [None, "Periyod", "2025 3Ç", "2026 2Ç"],
        [None, "Mobil Toplam Abone Sayısı (mn)", 30.808, 32.74],
        [None, "Mobil Faturalı Abone Sayısı (mn)", 23.972, 26.296],
        [None, "Fiber Abone Sayısı (mn)", 14.233, 14.488],
        [None, "Genişbant Toplam Abone Sayısı (mn)", 15.474, 15.401],
    ]


def test_payi_serisi_mobil_faturali_orani_hesaplar():
    kitap = _kitabi_ac({"Abone Verileri": _abone_verileri_sayfasi()})
    sonuc = _payi_serisi(kitap, "mobil-faturali-abone-payi")
    assert sonuc["2026 2Ç"] == pytest.approx(26.296 / 32.74 * 100, abs=1e-6)
    assert sonuc["2025 3Ç"] == pytest.approx(23.972 / 30.808 * 100, abs=1e-6)


def test_payi_serisi_fiber_orani_hesaplar():
    kitap = _kitabi_ac({"Abone Verileri": _abone_verileri_sayfasi()})
    sonuc = _payi_serisi(kitap, "sabit-genisbant-fiber-abone-payi")
    assert sonuc["2026 2Ç"] == pytest.approx(14.488 / 15.401 * 100, abs=1e-6)


def test_payi_serisi_toplam_eksikse_o_donemi_atlar():
    sayfalar = {
        "Abone Verileri": [
            [None, "Periyod", "2025 3Ç", "2026 2Ç"],
            [None, "Mobil Faturalı Abone Sayısı (mn)", 23.972, 26.296],
            # Mobil Toplam Abone Sayısı satırı yok -> pay hiçbir dönem için hesaplanamaz
        ],
    }
    kitap = _kitabi_ac(sayfalar)
    assert _payi_serisi(kitap, "mobil-faturali-abone-payi") == {}


def test_seri_cek_payi_metrigini_ceyrek_ilk_ayina_damgalar():
    kitap = _kitabi_ac({"Abone Verileri": _abone_verileri_sayfasi()})
    onbellek = {"kitap": kitap}
    seri = SimpleNamespace(
        id="telekom/ttkom-mobil-faturali-abone-payi",
        ttkom_metrik="mobil-faturali-abone-payi",
        start_date=None,
    )
    df = seri_cek(seri, onbellek=onbellek)
    assert "2026-04-01" in list(df["date"])  # 2Ç -> çeyreğin ilk ayı
    satir = df[df["date"] == "2026-04-01"].iloc[0]
    assert satir["value"] == pytest.approx(26.296 / 32.74 * 100, abs=1e-6)


def test_abone_payi_kaynak_iki_metrik_tanimlar():
    assert set(ABONE_PAYI_KAYNAK) == {
        "mobil-faturali-abone-payi", "sabit-genisbant-fiber-abone-payi",
    }
