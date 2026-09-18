"""ifo Institute (`ingest/ifo.py`) istemcisi testleri."""

from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.ifo import (
    DOSYA_SABLONU,
    SERI_ETIKETLERI,
    _ay_geri,
    _en_guncel_kitabi_indir,
    _sutun_indeksi,
    seri_cek,
)


def ifo_seri(**kwargs):
    varsayilan = dict(id="avrupa/x", ifo_seri="iklim", start_date=None)
    varsayilan.update(kwargs)
    return SimpleNamespace(**varsayilan)


# --- _ay_geri: yıl sınırını doğru geçer ---


def test_ay_geri_ayni_yil_icinde_geri_gider():
    assert _ay_geri(2026, 9, 1) == (2026, 8)


def test_ay_geri_ocaktan_bir_geri_onceki_aralik():
    assert _ay_geri(2026, 1, 1) == (2025, 12)


def test_ay_geri_sifir_ayni_ayi_doner():
    assert _ay_geri(2026, 9, 0) == (2026, 9)


def test_ay_geri_birden_fazla_yil_sinirini_gecer():
    assert _ay_geri(2026, 2, 14) == (2024, 12)


# --- _sutun_indeksi: soldaki (Endeks grubu) eşleşmeyi seçer ---


def _kitap_baytlari(sayfalar: dict) -> bytes:
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    for ad, satirlar in sayfalar.items():
        ws = kitap.create_sheet(ad)
        for satir in satirlar:
            ws.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


def _kitabi_ac(sayfalar: dict):
    return openpyxl.load_workbook(io.BytesIO(_kitap_baytlari(sayfalar)), data_only=True)


def _ifo_sayfasi():
    """Gerçek dosyadaki gibi 'Business Climate' etiketi İKİ kez geçer:
    önce Endeks grubunda (sol), sonra Denge grubunda (sağ)."""
    return _kitabi_ac({"ifo Business Climate": [
        ["ifo Business Climate Germany"],
        [None, "Germany"],
        [None, "Index, 2015 = 100, seasonally adjusted", None, None, "Balances, seasonally adjusted"],
        ["Month/year", "Business Climate", "Business Situation", "Business Expectations",
         "Business Climate", "Business Situation", "Business Expectations"],
        [None],
        [" 08/2026", 88.8, 88.5, 89.1, -5.9, 1.7, -13.2],
    ]})["ifo Business Climate"]


def test_sutun_indeksi_soldaki_endeks_grubunu_secer():
    ws = _ifo_sayfasi()
    assert _sutun_indeksi(ws, "Business Climate") == 1  # Denge grubu (sütun 4) değil


def test_sutun_indeksi_durum_ve_beklenti_dogru_sutunlari_bulur():
    ws = _ifo_sayfasi()
    assert _sutun_indeksi(ws, "Business Situation") == 2
    assert _sutun_indeksi(ws, "Business Expectations") == 3


def test_sutun_indeksi_bulunamazsa_hata():
    ws = _kitabi_ac({"S": [["baslik"], [None, "Başka Etiket"]]})["S"]
    with pytest.raises(RuntimeError, match="Business Climate"):
        _sutun_indeksi(ws, "Business Climate")


# --- _en_guncel_kitabi_indir: soft-404 atlanır, gerçek XLSX bulunur ---


class SahteYanit:
    def __init__(self, content=b"", status_code=200):
        self.content = content
        self.status_code = status_code


class SahteOturum:
    """Var olmayan ay HTTP 200 ile küçük bir "kabuk" döner (gerçek WAF
    davranışı, ölçüldü 2026-09-18); yalnızca gerçek XLSX'ler ZIP sihirli
    baytıyla (`PK\x03\x04`) başlar."""

    def __init__(self, gercek_dosyalar: dict[str, bytes]):
        self.gercek_dosyalar = gercek_dosyalar
        self.cagrilar = []

    def get(self, url, timeout=None, headers=None):
        self.cagrilar.append(url)
        if url in self.gercek_dosyalar:
            return SahteYanit(content=self.gercek_dosyalar[url])
        return SahteYanit(content=b"<html>kabuk</html>")  # soft-404


def _xlsx_baytlari() -> bytes:
    return _kitap_baytlari({"ifo Business Climate": [
        ["Month/year", "Business Climate", "Business Situation", "Business Expectations"],
        [" 08/2026", 88.8, 88.5, 89.1],
    ]})


def test_en_guncel_kitabi_indir_en_yeni_ayi_bulur():
    url_agustos = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url_agustos: _xlsx_baytlari()})
    kitap = _en_guncel_kitabi_indir(session=oturum, bugun=date(2026, 8, 25))
    assert kitap["ifo Business Climate"]["A2"].value == " 08/2026"


def test_en_guncel_kitabi_indir_yayimlanmamis_ayi_atlayip_gerideki_ayi_bulur():
    """Eylül henüz yayımlanmadıysa (soft-404) Ağustos'a düşer."""
    url_agustos = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url_agustos: _xlsx_baytlari()})
    kitap = _en_guncel_kitabi_indir(session=oturum, bugun=date(2026, 9, 18))
    assert kitap["ifo Business Climate"]["A2"].value == " 08/2026"
    # Önce Eylül denenmiş olmalı (soft-404), sonra Ağustos.
    assert oturum.cagrilar[0] == DOSYA_SABLONU.format(yil=2026, ay=9)
    assert oturum.cagrilar[1] == url_agustos


def test_en_guncel_kitabi_indir_hicbiri_bulunamazsa_hata():
    oturum = SahteOturum({})
    with pytest.raises(RuntimeError, match="bulunamadı"):
        _en_guncel_kitabi_indir(session=oturum, bugun=date(2026, 9, 18))


# --- seri_cek: ağ kabuğu + önbellek ---


def test_seri_cek_dogru_deger_dondurur():
    url = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url: _xlsx_baytlari()})
    df = seri_cek(ifo_seri(ifo_seri="iklim"), onbellek={}, session=oturum, bugun=date(2026, 9, 18))
    assert df.iloc[-1]["date"] == "2026-08-01"
    assert df.iloc[-1]["value"] == pytest.approx(88.8)


def test_seri_cek_durum_ve_beklenti_dogru_sutunu_okur():
    url = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url: _xlsx_baytlari()})
    durum = seri_cek(ifo_seri(ifo_seri="durum"), onbellek={}, session=oturum, bugun=date(2026, 9, 18))
    beklenti = seri_cek(ifo_seri(ifo_seri="beklenti"), onbellek={}, session=oturum, bugun=date(2026, 9, 18))
    assert durum.iloc[-1]["value"] == pytest.approx(88.5)
    assert beklenti.iloc[-1]["value"] == pytest.approx(89.1)


def test_seri_cek_onbellegi_paylasir():
    url = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url: _xlsx_baytlari()})
    onbellek: dict = {}
    seri_cek(ifo_seri(ifo_seri="iklim"), onbellek=onbellek, session=oturum, bugun=date(2026, 9, 18))
    cagri_sayisi = len(oturum.cagrilar)
    seri_cek(ifo_seri(ifo_seri="durum"), onbellek=onbellek, session=oturum, bugun=date(2026, 9, 18))
    assert len(oturum.cagrilar) == cagri_sayisi  # ikinci çağrı ağa çıkmamalı


def test_seri_cek_start_date_oncesini_kirpar():
    kitap_bytes = _kitap_baytlari({"ifo Business Climate": [
        ["Month/year", "Business Climate"],
        [" 07/2026", 86.7],
        [" 08/2026", 88.8],
    ]})
    url = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url: kitap_bytes})
    df = seri_cek(
        ifo_seri(ifo_seri="iklim", start_date="2026-08-01"),
        onbellek={}, session=oturum, bugun=date(2026, 9, 18),
    )
    assert list(df["date"]) == ["2026-08-01"]


def test_seri_cek_bilinmeyen_ifo_serisinde_hata():
    url = DOSYA_SABLONU.format(yil=2026, ay=8)
    oturum = SahteOturum({url: _xlsx_baytlari()})
    with pytest.raises(RuntimeError, match="Bilinmeyen ifo_seri"):
        seri_cek(ifo_seri(ifo_seri="olmayan-seri"), onbellek={}, session=oturum, bugun=date(2026, 9, 18))


def test_seri_etiketleri_katalog_enumuyla_birebir():
    from core.catalog import GECERLI_IFO_SERILERI
    assert set(SERI_ETIKETLERI.keys()) == GECERLI_IFO_SERILERI
