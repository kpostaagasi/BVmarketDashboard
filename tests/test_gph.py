"""GPH (Global Ports Holding) Aylık Trafik İstatistikleri istemcisi testleri."""

import io
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.gph import (
    _guncel_dosya_url,
    _sekme_degerini_oku,
    seri_cek,
    sekme_donemi,
    tum_noktalari_ayikla,
)


def _sekme_ekle(kitap, ad, konsolide_yolcu, konsolide_sefer, ke_yolcu=None, ke_sefer=None):
    """`ke_yolcu`/`ke_sefer` None ise "Konsolide Edilmeyen Limanlar" bölümü
    hiç yazılmaz (2026-01 öncesi şablonu taklit eder)."""
    ws = kitap.create_sheet(ad)
    for _ in range(7):
        ws.append([None] * 8)
    ws.cell(8, 2, "Kruvaziyer Limanları")
    etiket_seferi = "Toplam Sefer Sayısı" + (" (Konsolide)" if ke_yolcu is not None else "")
    etiket_yolcu = "Toplam Yolcu Sayısı" + (" (Konsolide)" if ke_yolcu is not None else "")
    ws.append([None, etiket_seferi, None, None, konsolide_sefer])
    ws.append([None, etiket_yolcu, None, None, konsolide_yolcu])
    if ke_yolcu is not None:
        ws.append([None, "Konsolide Edilmeyen Limanlar", None, None, None])
        ws.append([None, None, "Seferler", None, ke_sefer])
        ws.append([None, None, "Yolcu Sayısı", None, ke_yolcu])
        ws.append([None, "Toplam Seferler (Tüm GPH Portföyü)", None, None, konsolide_sefer + (ke_sefer or 0)])
        ws.append([None, "Toplam Yolcu Sayısı (Tüm GPH Portföyü)", None, None, konsolide_yolcu + ke_yolcu])
    return ws


def _kitap_bytes(sekmeler: list[tuple]) -> bytes:
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    for args in sekmeler:
        _sekme_ekle(kitap, *args)
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


# --- sekme_donemi: sekme adından (yıl, ay) çıkarımı ---


def test_sekme_donemi_tam_yil():
    assert sekme_donemi("Ağustos-2026") == (2026, 8)


def test_sekme_donemi_kisa_yil():
    assert sekme_donemi("Tem-22") == (2022, 7)


def test_sekme_donemi_taninmayan_sekme_none_doner():
    assert sekme_donemi("Mart-23_Eski Raporlama") is None
    assert sekme_donemi("Notlar") is None


# --- _sekme_degerini_oku: iki dönem şablonu (Konsolide Edilmeyen var/yok) ---


def test_sekme_degerini_oku_2026_oncesi_sablon():
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    ws = _sekme_ekle(kitap, "Kasim-2025", konsolide_yolcu=1596402, konsolide_sefer=669)
    assert _sekme_degerini_oku(ws, "yolcu-konsolide") == 1596402
    assert _sekme_degerini_oku(ws, "sefer-konsolide") == 669
    assert _sekme_degerini_oku(ws, "yolcu-konsolide-edilmeyen") is None  # bu dönemde yok


def test_sekme_degerini_oku_2026_sonrasi_sablon():
    kitap = openpyxl.Workbook()
    kitap.remove(kitap.active)
    ws = _sekme_ekle(
        kitap, "Agustos-2026", konsolide_yolcu=1768845, konsolide_sefer=518,
        ke_yolcu=397324, ke_sefer=123,
    )
    assert _sekme_degerini_oku(ws, "yolcu-konsolide") == 1768845
    assert _sekme_degerini_oku(ws, "yolcu-konsolide-edilmeyen") == 397324
    # "Tüm GPH Portföyü" satırı "Toplam Yolcu Sayısı" ile başlasa da HARİÇ tutulmalı
    assert _sekme_degerini_oku(ws, "yolcu-konsolide") != 1768845 + 397324


# --- tum_noktalari_ayikla: birden çok sekme, tarih sıralaması ---


def test_tum_noktalari_ayikla_tarihe_gore_siralar():
    kitap = openpyxl.load_workbook(
        io.BytesIO(_kitap_bytes([
            ("Agustos-2026", 1768845, 518),
            ("Temmuz-2026", 1798335, 530),
        ])),
        data_only=True, read_only=True,
    )
    noktalar = tum_noktalari_ayikla(kitap, "yolcu-konsolide")
    assert noktalar == [("2026-07-01", 1798335.0), ("2026-08-01", 1768845.0)]


def test_tum_noktalari_ayikla_veri_yoksa_hata():
    kitap = openpyxl.Workbook()
    kitap.active.title = "Notlar"  # sekme_donemi ile eşleşmeyen tek sekme
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap = openpyxl.load_workbook(tampon, data_only=True, read_only=True)
    with pytest.raises(RuntimeError, match="hiçbir sekmede"):
        tum_noktalari_ayikla(kitap, "yolcu-konsolide")


# --- _guncel_dosya_url: IR sayfasından ilk .xlsx bağlantısı ---


class _SahteYanit:
    def __init__(self, text="", content=b""):
        self.text = text
        self.content = content

    def raise_for_status(self):
        pass


class _SahteOturum:
    def __init__(self, sayfa_html: str, dosya_icerigi: bytes):
        self.sayfa_html = sayfa_html
        self.dosya_icerigi = dosya_icerigi

    def get(self, url, **kwargs):
        if url.endswith(".xlsx"):
            return _SahteYanit(content=self.dosya_icerigi)
        return _SahteYanit(text=self.sayfa_html)


def test_guncel_dosya_url_ilk_baglantiyi_secer():
    html = (
        '<a href=" https://x.com/gph-aylik-trafik-agustos-2026.xlsx">İNDİR</a>'
        '<a href=" https://x.com/GPH_AylikTrafik_Temmuz2026.xlsx">İNDİR</a>'
    )
    oturum = _SahteOturum(html, b"")
    assert _guncel_dosya_url(oturum) == "https://x.com/gph-aylik-trafik-agustos-2026.xlsx"


def test_guncel_dosya_url_baglanti_yoksa_hata():
    oturum = _SahteOturum("<html>hiç link yok</html>", b"")
    with pytest.raises(RuntimeError, match="xlsx bağlantısı"):
        _guncel_dosya_url(oturum)


# --- seri_cek: uçtan uca, önbellek paylaşımı ---


def test_seri_cek_dosya_url_onbellekten_okur():
    icerik = _kitap_bytes([("Agustos-2026", 1768845, 518)])
    onbellek = {"dosya_url": "https://x.com/f.xlsx"}
    oturum = _SahteOturum("", icerik)
    df = seri_cek(SimpleNamespace(gph_metrik="yolcu-konsolide"), onbellek=onbellek, session=oturum)
    assert list(df["value"]) == [1768845.0]


def test_seri_cek_bilinmeyen_metrikte_hata():
    onbellek = {"dosya_url": "https://x.com/f.xlsx"}
    oturum = _SahteOturum("", _kitap_bytes([("Agustos-2026", 1, 1)]))
    with pytest.raises(RuntimeError, match="bilinmeyen gph_metrik"):
        seri_cek(SimpleNamespace(gph_metrik="olmayan"), onbellek=onbellek, session=oturum)
