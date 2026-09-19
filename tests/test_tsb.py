"""TSB (Türkiye Sigorta Birliği) Prim Üretimleri Sıralama istemcisi testleri."""

from datetime import date
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest

from ingest.tsb import ILK_DESTEKLENEN_YIL, dosya_listesi, kumulatif_seri, seri_cek, sheet_degerleri


def tsb_seri(**kwargs):
    varsayilan = dict(
        tsb_alt_kategori="prim-adet",
        tsb_rapor="Prim Üretimleri Sıralama",
        tsb_sheet="Hayatdışı",
        tsb_sirket_kodu=1020,
        start_date=None,
    )
    varsayilan.update(kwargs)
    return SimpleNamespace(**varsayilan)


def _workbook_baytlari(sheet_satirlari: dict[str, list[tuple]]) -> bytes:
    """`{sheet_adi: [(Sıralama, Şirket Adı, Şirket Kodu, Toplam Üretim, Pazar Payı)]}`
    biçiminden gerçek şablonu taklit eden bir workbook üretir (5 metadata
    satırı + başlık satırı + veri satırları)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, satirlar in sheet_satirlari.items():
        ws = wb.create_sheet(sheet)
        ws.append(["ŞİRKETLERİN PAZAR PAY VE TOPLAM PRİM DEĞİŞİMLERİ"])
        ws.append([sheet])
        ws.append(["01.01.2026-31.08.2026"])
        ws.append([])
        ws.append([None, None, None, "2026-08"])
        ws.append(["Sıralama", "Şirket Adı", "Şirket Kodu", "Toplam Üretim (TL)", "Pazar Payı %"])
        for satir in satirlar:
            ws.append(list(satir))
    arabellek = BytesIO()
    wb.save(arabellek)
    return arabellek.getvalue()


class SahteYanit:
    def __init__(self, status_code=200, json_veri=None, content=b""):
        self.status_code = status_code
        self._json = json_veri
        self.content = content

    def json(self):
        return self._json


class SahteOturum:
    """`sayfalar`: {alt_kategori: [[dosya1, dosya2, ...], [...sayfa2...]]}.
    `dosyalar`: {FilePath: workbook_baytlari}."""

    def __init__(self, sayfalar, dosyalar):
        self.sayfalar = sayfalar
        self.dosyalar = dosyalar
        self.cagrilar = []

    def get(self, url, params=None, timeout=None):
        self.cagrilar.append(url)
        if url.endswith("/Statistic/GetAllStatistics"):
            alt_kategori = params["SubCategoryUrl"]
            sayfa_no = params["pageId"]
            sayfalar = self.sayfalar.get(alt_kategori, [])
            sonuc = sayfalar[sayfa_no - 1] if sayfa_no <= len(sayfalar) else []
            return SahteYanit(json_veri={"Result": sonuc})
        for yol, icerik in self.dosyalar.items():
            if url.endswith(yol):
                return SahteYanit(content=icerik)
        return SahteYanit(status_code=404)


def _dosya_kaydi(ad, yil, ay, yol):
    return {
        "FileName": ad, "FilePath": yol, "PeriodYear": yil, "PeriodMonth": ay,
    }


# --- sheet_degerleri: workbook satır ayrıştırma ---


def test_sheet_degerleri_sirket_kodundan_deger_okur():
    baytlar = _workbook_baytlari({
        "Hayatdışı": [(1, "Türkiye Sigorta AŞ", 1020, 120434918659.72, 0.1476)],
    })
    wb = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    degerler = sheet_degerleri(wb, "Hayatdışı")
    assert degerler[1020] == pytest.approx(120434918659.72)


def test_sheet_degerleri_sheet_yoksa_hata():
    baytlar = _workbook_baytlari({"Hayat": [(1, "X", 1, 1.0, 1.0)]})
    wb = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    with pytest.raises(RuntimeError, match="Hayatdışı"):
        sheet_degerleri(wb, "Hayatdışı")


def test_sheet_degerleri_bos_satirlari_atlar():
    """Boş 'Şirket Kodu' hücreli satır (ör. dipnot) sessizce atlanır."""
    baytlar = _workbook_baytlari({
        "Hayatdışı": [
            (1, "Türkiye Sigorta AŞ", 1020, 100.0, 0.5),
            (None, "*Not: geçici rakamlar", None, None, None),
        ],
    })
    wb = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    degerler = sheet_degerleri(wb, "Hayatdışı")
    assert degerler == {1020: 100.0}


def test_sheet_degerleri_kodsuz_sektor_toplamini_9003e_normalize_eder():
    """Eski dosyalarda (ör. 2022-03) 'SEKTÖR TOPLAMI' satırının Şirket Kodu
    hücresi boş — isimden tanınıp 9003'e normalize edilmeli."""
    baytlar = _workbook_baytlari({
        "Hayat": [
            (1, "Türkiye Hayat ve Emeklilik AŞ", 3018, 500.0, 0.2),
            (None, "SEKTÖR TOPLAMI", None, 5793809742.64, 1.0),
        ],
    })
    wb = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    degerler = sheet_degerleri(wb, "Hayat")
    assert degerler[9003] == pytest.approx(5793809742.64)


def test_sheet_degerleri_kodsuz_baska_satiri_atlar():
    """Kodu boş ve adı 'SEKTÖR TOPLAMI' OLMAYAN satır (ör. dipnot) atlanır,
    yanlışlıkla 9003'e eşlenmez."""
    baytlar = _workbook_baytlari({
        "Hayat": [
            (1, "Türkiye Hayat ve Emeklilik AŞ", 3018, 500.0, 0.2),
            (None, "*geçici rakam notu", None, 999.0, None),
        ],
    })
    wb = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    degerler = sheet_degerleri(wb, "Hayat")
    assert degerler == {3018: 500.0}


# --- dosya_listesi: sayfalama ---


def test_dosya_listesi_erken_yila_ulasinca_durur():
    sayfalar = {
        "prim-adet": [
            [_dosya_kaydi("3 Prim Üretimleri Sıralama 2026-08", 2026, 8, "/a.xlsx")],
            [_dosya_kaydi("3 Prim Üretimleri Sıralama 2020-01", 2020, 1, "/b.xlsx")],
            [_dosya_kaydi("3 Prim Üretimleri Sıralama 2015-01", 2015, 1, "/c.xlsx")],
        ]
    }
    oturum = SahteOturum(sayfalar, {})
    dosyalar = dosya_listesi("prim-adet", en_eski_yil=2021, session=oturum)
    # 2020 sayfası tamamen sınırın altında (2020 < 2021) — orada durur,
    # 2015 sayfasına inmez.
    assert len(dosyalar) == 2
    assert oturum.cagrilar.count("https://www.tsb.org.tr/Statistic/GetAllStatistics") == 2


def test_dosya_listesi_bos_alt_kategoride_hata():
    oturum = SahteOturum({"prim-adet": [[]]}, {})
    with pytest.raises(RuntimeError, match="prim-adet"):
        dosya_listesi("prim-adet", en_eski_yil=2020, session=oturum)


def test_kumulatif_seri_nfd_dosya_adiyla_nfc_raporu_eslestirir():
    """2021-03 öncesi yeniden indekslenmiş kayıtlarda FileName NFD Unicode
    kullanıyor (ör. 'Ü' = 'U' + birleşen çift nokta); normalize edilmeden
    substring araması sessizce hiç eşleşmiyordu."""
    nfd_ad = "Prim U\u0308retimleri Sıralama 2018-12"  # NFD: U + combining diaeresis
    assert "Ü" not in nfd_ad  # doğrulama: gerçekten NFD, NFC "Ü" içermiyor
    dosyalar = {
        "/nfd.xlsx": _workbook_baytlari({
            "Hayatdışı": [(1, "Türkiye Sigorta AŞ", 1020, 100.0, 0.5)],
        }),
    }
    sayfalar = {"prim-adet": [[_dosya_kaydi(nfd_ad, 2018, 12, "/nfd.xlsx")]]}
    oturum = SahteOturum(sayfalar, dosyalar)
    kumulatif = kumulatif_seri(
        tsb_seri(start_date="2018-01-01"), onbellek={}, session=oturum
    )
    assert kumulatif == {"2018-12-01": 100.0}


def test_kumulatif_seri_xls_donemine_inmez():
    """`ILK_DESTEKLENEN_YIL`nin altındaki dönemler (eski .xls dosyaları)
    `start_date` boş bırakılsa bile hiç indirilmeye ÇALIŞILMAZ (workbook
    açma denemesi olsaydı BadZipFile ile patlardı)."""
    dosyalar = {
        "/eski.xlsx": _workbook_baytlari({"Hayatdışı": [(1, "X", 1020, 1.0, 1.0)]}),
    }
    sayfalar = {
        "prim-adet": [[
            _dosya_kaydi("3 Prim Üretimleri Sıralama 2014-01", 2014, 1, "/eski.xlsx"),
        ]]
    }
    oturum = SahteOturum(sayfalar, dosyalar)
    kumulatif = kumulatif_seri(tsb_seri(start_date=None), onbellek={}, session=oturum)
    assert kumulatif == {}
    assert not any(c.endswith("/eski.xlsx") for c in oturum.cagrilar)
    assert ILK_DESTEKLENEN_YIL == 2015


# --- seri_cek: kümülatif→aylık dönüşüm ve ağ kabuğu ---


def _iki_aylik_kurulum():
    dosyalar = {
        "/a08.xlsx": _workbook_baytlari({
            "Hayatdışı": [
                (1, "Türkiye Sigorta AŞ", 1020, 120434918659.72, 0.1476),
                (None, "SEKTÖR TOPLAMI", 9003, 815731863400.10, 0.8394),
            ],
        }),
        "/a07.xlsx": _workbook_baytlari({
            "Hayatdışı": [
                (1, "Türkiye Sigorta AŞ", 1020, 107113524712.56, 0.1450),
                (None, "SEKTÖR TOPLAMI", 9003, 723794475018.95, 0.8410),
            ],
        }),
    }
    sayfalar = {
        "prim-adet": [[
            _dosya_kaydi("3 Prim Üretimleri Sıralama 2026-08", 2026, 8, "/a08.xlsx"),
            _dosya_kaydi("3 Prim Üretimleri Sıralama 2026-07", 2026, 7, "/a07.xlsx"),
        ]]
    }
    return SahteOturum(sayfalar, dosyalar)


def test_seri_cek_kumulatifi_aylik_farka_cevirir():
    oturum = _iki_aylik_kurulum()
    df = seri_cek(tsb_seri(start_date="2020-01-01"), onbellek={}, session=oturum)
    agustos = df[df["date"] == "2026-08-01"]["value"].iloc[0]
    assert agustos == pytest.approx(120434918659.72 - 107113524712.56)


def test_seri_cek_sektor_toplami_kodu_9003():
    oturum = _iki_aylik_kurulum()
    df = seri_cek(tsb_seri(tsb_sirket_kodu=9003, start_date="2020-01-01"), onbellek={}, session=oturum)
    agustos = df[df["date"] == "2026-08-01"]["value"].iloc[0]
    assert agustos == pytest.approx(815731863400.10 - 723794475018.95)


def test_seri_cek_onbellek_workbooku_sirketler_arasinda_paylasir():
    """Aynı workbook'u okuyan iki şirket serisi dosyayı yalnızca bir kez indirir."""
    oturum = _iki_aylik_kurulum()
    onbellek = {}
    seri_cek(tsb_seri(tsb_sirket_kodu=1020, start_date="2020-01-01"), onbellek=onbellek, session=oturum)
    cagri_sayisi = len(oturum.cagrilar)
    seri_cek(tsb_seri(tsb_sirket_kodu=9003, start_date="2020-01-01"), onbellek=onbellek, session=oturum)
    # Dosya listesi + workbook indirmeleri tekrarlanmaz; yalnızca çağrı sayısı artmaz.
    assert len(oturum.cagrilar) == cagri_sayisi


def test_seri_cek_onbellek_workbooku_farkli_sheetler_arasinda_paylasir():
    """Aynı ayın workbook'unu FARKLI sheet isteyen iki seri de tek indirme
    üzerinden okur (workbook indirme, sheet çözümünden AYRI önbelleklenir)."""
    dosyalar = {
        "/karma.xlsx": _workbook_baytlari({
            "Hayatdışı": [(1, "Türkiye Sigorta AŞ", 1020, 100.0, 0.5)],
            "Hayat": [(1, "Türkiye Hayat ve Emeklilik AŞ", 3018, 50.0, 0.5)],
        }),
    }
    sayfalar = {"prim-adet": [[_dosya_kaydi("3 Prim Üretimleri Sıralama 2026-08", 2026, 8, "/karma.xlsx")]]}
    oturum = SahteOturum(sayfalar, dosyalar)
    onbellek = {}
    kumulatif_seri(tsb_seri(tsb_sheet="Hayatdışı", tsb_sirket_kodu=1020, start_date="2020-01-01"), onbellek=onbellek, session=oturum)
    indirme_sayisi = sum(1 for c in oturum.cagrilar if c.endswith("/karma.xlsx"))
    kumulatif_seri(tsb_seri(tsb_sheet="Hayat", tsb_sirket_kodu=3018, start_date="2020-01-01"), onbellek=onbellek, session=oturum)
    assert sum(1 for c in oturum.cagrilar if c.endswith("/karma.xlsx")) == indirme_sayisi == 1


def test_seri_cek_start_date_oncesini_kirpar():
    oturum = _iki_aylik_kurulum()
    df = seri_cek(tsb_seri(start_date="2026-08-01"), onbellek={}, session=oturum)
    assert list(df["date"]) == ["2026-08-01"]


def test_seri_cek_eslesen_rapor_yoksa_hata():
    oturum = _iki_aylik_kurulum()
    with pytest.raises(RuntimeError, match="eşleşen dosya yok"):
        seri_cek(tsb_seri(tsb_rapor="Olmayan Rapor", start_date="2020-01-01"), onbellek={}, session=oturum)


def test_seri_cek_bilinmeyen_sirket_kodunda_hata():
    oturum = _iki_aylik_kurulum()
    with pytest.raises(RuntimeError, match="9999"):
        seri_cek(tsb_seri(tsb_sirket_kodu=9999, start_date="2020-01-01"), onbellek={}, session=oturum)


def test_seri_cek_http_hatasinda_yukselir():
    class HataliOturum:
        def get(self, url, params=None, timeout=None):
            return SahteYanit(status_code=500)

    with pytest.raises(RuntimeError, match="HTTP 500"):
        seri_cek(tsb_seri(start_date="2020-01-01"), onbellek={}, session=HataliOturum())
