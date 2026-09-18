"""TİM ülke grupları bülteni istemcisi (`ulke_grubu_seri_cek`) testleri.

Bültenin ayrıştırma katmanı (`ulke_grubu_noktalari`, `_eski_sablon_grup_noktalari`,
`_ulke_grubu_dosyasini_ayikla`) gerçek XLSX'e karşı ölçüldü (bkz. modülün
docstring'i) — burada hem yapısal (minimal XLSX fixture) hem davranışsal
(monkeypatch ile `_ulke_grubu_dosyasini_ayikla` beslenerek `ulke_grubu_seri_cek`
düzeyi: aylık sütun seçimi, 404 atlama, önbellek paylaşımı, eksik grup) testler var.
"""

import io
from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest

import ingest.tim as tim
from ingest.tim import (
    ULKE_GRUBU_ESKI_SAYFA_ADI,
    ULKE_GRUBU_YENI_SAYFA_ADI,
    _eski_sablon_grup_noktalari,
    _ulke_grubu_dosyasini_ayikla,
    ulke_grubu_bazinda_url,
    ulke_grubu_noktalari,
    ulke_grubu_seri_cek,
)


def grup_seri(**kwargs):
    varsayilan = dict(
        id="ihracat-ulke-grubu/ab-toplam",
        kaynak_tipi="tim_ulke_grubu",
        tim_ulke_grubu="Avrupa Birliği Ülkeleri",
        start_date=None,
    )
    varsayilan.update(kwargs)
    return SimpleNamespace(**varsayilan)


class SahteYanit:
    def __init__(self, status_code, content=b""):
        self.status_code = status_code
        self.content = content


class SahteOturum:
    def __init__(self, yanitlar):
        self.yanitlar = yanitlar
        self.cagrilar = []

    def get(self, url, timeout=None):
        self.cagrilar.append(url)
        return self.yanitlar.get(url, SahteYanit(404, b"<html>404</html>"))


def agustos_2026_noktalari():
    """Ölçülen gerçek sütun etiketleriyle sahte çıktı: doğru aylık sütun
    ("1 - 31 AĞUSTOS / 2026") + tuzaklar (önceki yıl, DEĞ., YTD, önceki ay)."""
    return {
        "Avrupa Birliği Ülkeleri": {
            "1 - 31 AĞUSTOS / 2025": 8110139.3,
            "1 - 31 AĞUSTOS / 2026": 7984727.72,
            "1 - 31 AĞUSTOS / DEĞ.": -0.0155,
            "1 - 31 TEMMUZ / 2026": 9268124.41,
            "1 - 31 TEMMUZ / DEĞ.": -0.1385,
            "1 OCAK  -  31 AĞUSTOS / 2025": 69320829.31,
            "1 OCAK  -  31 AĞUSTOS / 2026": 71788515.05,
        },
        "Ortadoğu Ülkeleri": {
            "1 - 31 AĞUSTOS / 2025": 2602620.2,
            "1 - 31 AĞUSTOS / 2026": 2520174.61,
            "1 - 31 TEMMUZ / 2026": 2453693.53,
        },
        "TOPLAM": {
            "1 - 31 AĞUSTOS / 2025": 18561853.03,
            "1 - 31 AĞUSTOS / 2026": 20200170.6,
            "1 - 31 TEMMUZ / 2026": 22028915.64,
        },
    }


# --- Aylık sütun seçimi: YTD, DEĞ. ve aynı yılın önceki ayı elenir ---


def test_aylik_sutun_secimi_ytd_ve_deg_sutunlarini_almaz(monkeypatch):
    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", lambda b, a: agustos_2026_noktalari())
    oturum = SahteOturum({ulke_grubu_bazinda_url(2026, 8): SahteYanit(200, b"xlsx")})

    df = ulke_grubu_seri_cek(grup_seri(), session=oturum, bugun=date(2026, 8, 15))

    assert df["date"].iloc[0] == "2026-08-01"
    assert df["value"].iloc[0] == pytest.approx(7984727.72)
    assert df["value"].iloc[0] not in (8110139.3, 9268124.41, 71788515.05)


def test_farkli_grup_dogru_deger_dondurur(monkeypatch):
    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", lambda b, a: agustos_2026_noktalari())
    oturum = SahteOturum({ulke_grubu_bazinda_url(2026, 8): SahteYanit(200, b"xlsx")})

    df = ulke_grubu_seri_cek(
        grup_seri(tim_ulke_grubu="TOPLAM"), session=oturum, bugun=date(2026, 8, 15),
    )

    assert df["value"].iloc[0] == pytest.approx(20200170.6)


# --- Eksik grup: sessiz atlama yasak ---


def test_eksik_grup_hata_yukseltir(monkeypatch):
    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", lambda b, a: agustos_2026_noktalari())
    oturum = SahteOturum({ulke_grubu_bazinda_url(2026, 8): SahteYanit(200, b"xlsx")})
    seri = grup_seri(tim_ulke_grubu="Uzayda Bir Grup")

    with pytest.raises(RuntimeError, match="grup bulunamadı"):
        ulke_grubu_seri_cek(seri, session=oturum, bugun=date(2026, 8, 15))


# --- 404: o ay atlanır, seri kalan aylarla üretilir ---


def test_404_ay_atlanir_seri_kalan_aylarla_uretilir(monkeypatch):
    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", lambda b, a: agustos_2026_noktalari())
    oturum = SahteOturum({ulke_grubu_bazinda_url(2026, 7): SahteYanit(200, b"xlsx")})

    df = ulke_grubu_seri_cek(grup_seri(), session=oturum, bugun=date(2026, 7, 15))

    assert len(df) == 1
    assert df["date"].iloc[0] == "2026-07-01"


def test_hicbir_ay_yoksa_hata_yukseltir():
    oturum = SahteOturum({})  # her URL 404

    with pytest.raises(RuntimeError, match="hiç nokta bulunamadı"):
        ulke_grubu_seri_cek(grup_seri(), session=oturum, bugun=date(2023, 1, 15))


# --- Önbellek paylaşımı: aynı ay iki grup için tek indirme/ayrıştırma ---


def test_onbellek_ayni_ayi_iki_grup_icin_yeniden_ayiklamaz(monkeypatch):
    cagri_sayisi = SimpleNamespace(n=0)

    def sayan_ayikla(b, a):
        cagri_sayisi.n += 1
        return agustos_2026_noktalari()

    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", sayan_ayikla)
    oturum = SahteOturum({ulke_grubu_bazinda_url(2026, 8): SahteYanit(200, b"xlsx")})

    onbellek: dict = {}
    ulke_grubu_seri_cek(
        grup_seri(tim_ulke_grubu="Avrupa Birliği Ülkeleri"),
        onbellek=onbellek, session=oturum, bugun=date(2026, 8, 15),
    )
    ilk = cagri_sayisi.n
    ulke_grubu_seri_cek(
        grup_seri(tim_ulke_grubu="Ortadoğu Ülkeleri"),
        onbellek=onbellek, session=oturum, bugun=date(2026, 8, 15),
    )

    assert cagri_sayisi.n == ilk


# --- start_date süzmesi ---


def test_start_date_oncesini_kirpar(monkeypatch):
    monkeypatch.setattr(tim, "_ulke_grubu_dosyasini_ayikla", lambda b, a: agustos_2026_noktalari())
    oturum = SahteOturum({
        ulke_grubu_bazinda_url(2026, 7): SahteYanit(200, b"xlsx"),
        ulke_grubu_bazinda_url(2026, 8): SahteYanit(200, b"xlsx"),
    })

    df = ulke_grubu_seri_cek(
        grup_seri(start_date="2026-08-01"), session=oturum, bugun=date(2026, 8, 15),
    )

    assert len(df) == 1
    assert df["date"].iloc[0] == "2026-08-01"


# --- YENİ şablon (tek anahtarlı ÜLKE GRUP) yapısal testleri ---


def _yeni_sablon_baytlari(sayfa_adi, satirlar):
    kitap = openpyxl.Workbook()
    sayfa = kitap.active
    sayfa.title = sayfa_adi
    sayfa.append(["31.08.2026 Konsolide Ülke Gruplarına Göre İhracat  (1000 $)"])
    sayfa.append([])
    sayfa.append([None, "1 - 31 AĞUSTOS", None, None, "1 - 31 TEMMUZ", None])
    sayfa.append(["ULKE GRUP", 2025, 2026, "DEĞ.", 2026, "DEĞ."])
    for satir in satirlar:
        sayfa.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap.close()
    return tampon.getvalue()


def test_yeni_sablon_grup_toplami_dogru_okunur(monkeypatch):
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 2)
    baytlar = _yeni_sablon_baytlari(ULKE_GRUBU_YENI_SAYFA_ADI, [
        ["Afrika Ülkeleri", 100.0, 120.0, 0.2, 110.0, 0.09],
        ["Avrupa Birliği Ülkeleri", 900.0, 880.0, -0.02, 950.0, -0.07],
        ["TOPLAM", 1000.0, 1000.0, 0.0, 1060.0, -0.06],
    ])

    noktalar = ulke_grubu_noktalari(baytlar, "2026.08", ULKE_GRUBU_YENI_SAYFA_ADI)

    assert noktalar["Afrika Ülkeleri"]["1 - 31 AĞUSTOS / 2026"] == 120.0
    assert noktalar["TOPLAM"]["1 - 31 AĞUSTOS / 2026"] == 1000.0


def test_yeni_sablon_toplam_uyusmazsa_hata_verir(monkeypatch):
    """Şablon kaymasını yakalamak için TOPLAM, grupların toplamıyla uzlaşmalı."""
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 2)
    baytlar = _yeni_sablon_baytlari(ULKE_GRUBU_YENI_SAYFA_ADI, [
        ["Afrika Ülkeleri", 100.0, 120.0, 0.2, 110.0, 0.09],
        ["Avrupa Birliği Ülkeleri", 900.0, 880.0, -0.02, 950.0, -0.07],
        ["TOPLAM", 1000.0, 999999.0, 0.0, 1060.0, -0.06],
    ])

    with pytest.raises(RuntimeError, match="uyuşmuyor"):
        ulke_grubu_noktalari(baytlar, "2026.08", ULKE_GRUBU_YENI_SAYFA_ADI)


def test_yeni_sablon_asgari_grup_sayisinin_altinda_hata_verir():
    baytlar = _yeni_sablon_baytlari(ULKE_GRUBU_YENI_SAYFA_ADI, [
        ["Afrika Ülkeleri", 100.0, 120.0, 0.2, 110.0, 0.09],
        ["TOPLAM", 100.0, 120.0, 0.2, 110.0, 0.09],
    ])

    with pytest.raises(RuntimeError, match="şablon değişmiş olabilir"):
        ulke_grubu_noktalari(baytlar, "2026.08", ULKE_GRUBU_YENI_SAYFA_ADI)


def test_yeni_sablon_yanlis_baslik_hata_verir():
    kitap = openpyxl.Workbook()
    sayfa = kitap.active
    sayfa.title = ULKE_GRUBU_YENI_SAYFA_ADI
    sayfa.append(["31.08.2026 Konsolide Ülke Gruplarına Göre İhracat  (1000 $)"])
    sayfa.append([])
    sayfa.append([None, "1 - 31 AĞUSTOS"])
    sayfa.append(["BAŞKA SÜTUN", 2025, 2026])
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap.close()

    with pytest.raises(RuntimeError, match="ULKE GRUP' başlığı yok"):
        ulke_grubu_noktalari(tampon.getvalue(), "2026.08", ULKE_GRUBU_YENI_SAYFA_ADI)


# --- ESKİ şablon (SEKTÖR×ÜLKEGRUP çapraz tablo) yapısal testleri ---


def _eski_sablon_baytlari(sayfa_adi, satirlar):
    kitap = openpyxl.Workbook()
    sayfa = kitap.active
    sayfa.title = sayfa_adi
    sayfa.append(["31.01.2023 Konsolide Ülke Guruplarına Göre Sektörel İhracat  (1000 $)"])
    sayfa.append([])
    sayfa.append([None, None, "1 - 31 OCAK"])
    sayfa.append(["SEKTÖR", "ULKEGRUP", 2022, 2023, "DEĞ."])
    for satir in satirlar:
        sayfa.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap.close()
    return tampon.getvalue()


def test_eski_sablon_grup_toplami_sektor_satirlarinin_toplamidir(monkeypatch):
    """Bu bültende grup toplamı yayımlanmıyor — TOPLAM olmayan sektör
    satırlarının toplamıyla türetilir (il/ülke'nin 'TOPLAM sektör' mantığıyla
    aynı)."""
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 2)
    baytlar = _eski_sablon_baytlari(ULKE_GRUBU_ESKI_SAYFA_ADI, [
        ["Çelik", "Afrika Ülkeleri", 10.0, 12.0, 0.2],
        ["Tekstil ve Hammaddeleri", "Afrika Ülkeleri", 20.0, 24.0, 0.2],
        ["Çelik", "Avrupa Birliği Ülkeleri", 100.0, 90.0, -0.1],
        ["Tekstil ve Hammaddeleri", "Avrupa Birliği Ülkeleri", 200.0, 180.0, -0.1],
        # Her sektörün TÜM gruplar toplamı ("TOPLAM" ÜLKEGRUP) — gerçek grup
        # değil, ayrıştırıcı bunu grup listesine KATMAMALI.
        ["Çelik", "TOPLAM", 110.0, 102.0, -0.07],
        ["Tekstil ve Hammaddeleri", "TOPLAM", 220.0, 204.0, -0.07],
    ])

    noktalar = _eski_sablon_grup_noktalari(baytlar, "2023.01", ULKE_GRUBU_ESKI_SAYFA_ADI)

    assert noktalar["Afrika Ülkeleri"]["1 - 31 OCAK / 2023"] == pytest.approx(36.0)
    assert noktalar["Avrupa Birliği Ülkeleri"]["1 - 31 OCAK / 2023"] == pytest.approx(270.0)
    assert "TOPLAM" not in [g for g in noktalar if g == "TOPLAM"] or noktalar["TOPLAM"]["1 - 31 OCAK / 2023"] == pytest.approx(306.0)


def test_eski_sablon_deg_sutunlari_tasinmaz(monkeypatch):
    """Oran sütunları sektörler arası anlamsızca toplanmasın diye hiç okunmaz."""
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 1)
    baytlar = _eski_sablon_baytlari(ULKE_GRUBU_ESKI_SAYFA_ADI, [
        ["Çelik", "Afrika Ülkeleri", 10.0, 12.0, 0.2],
        ["Tekstil ve Hammaddeleri", "Afrika Ülkeleri", 20.0, 24.0, 0.2],
    ])

    noktalar = _eski_sablon_grup_noktalari(baytlar, "2023.01", ULKE_GRUBU_ESKI_SAYFA_ADI)

    assert all(not etiket.endswith(" / DEĞ.") for etiket in noktalar["Afrika Ülkeleri"])


# --- Dispatcher: sayfa ADI değil dosyanın GERÇEK başlığı belirleyici ---


def test_dispatcher_sayfa_adi_yaniltici_olsa_da_gercek_basligi_okur(monkeypatch):
    """Ölçüldü (2023-05 dosyası): 'GUNLUK_SEKTOR_ULKEGRUBU' adında ama
    içeriği basit YENİ şablon — sayfa ADINA değil başlık hücresine bakılmalı."""
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 2)
    baytlar = _yeni_sablon_baytlari(ULKE_GRUBU_ESKI_SAYFA_ADI, [
        ["Afrika Ülkeleri", 100.0, 120.0, 0.2, 110.0, 0.09],
        ["Avrupa Birliği Ülkeleri", 900.0, 880.0, -0.02, 950.0, -0.07],
        ["TOPLAM", 1000.0, 1000.0, 0.0, 1060.0, -0.06],
    ])

    noktalar = _ulke_grubu_dosyasini_ayikla(baytlar, "2023.05")

    assert noktalar["Afrika Ülkeleri"]["1 - 31 AĞUSTOS / 2026"] == 120.0


def test_dispatcher_eski_sablonu_dogru_yonlendirir(monkeypatch):
    monkeypatch.setattr(tim, "ASGARI_ULKE_GRUBU_SAYISI", 2)
    baytlar = _eski_sablon_baytlari(ULKE_GRUBU_ESKI_SAYFA_ADI, [
        ["Çelik", "Afrika Ülkeleri", 10.0, 12.0, 0.2],
        ["Tekstil ve Hammaddeleri", "Afrika Ülkeleri", 20.0, 24.0, 0.2],
        ["Çelik", "Avrupa Birliği Ülkeleri", 100.0, 90.0, -0.1],
        ["Tekstil ve Hammaddeleri", "Avrupa Birliği Ülkeleri", 200.0, 180.0, -0.1],
    ])

    noktalar = _ulke_grubu_dosyasini_ayikla(baytlar, "2023.01")

    assert noktalar["Afrika Ülkeleri"]["1 - 31 OCAK / 2023"] == pytest.approx(36.0)


def test_dispatcher_bilinmeyen_sayfa_hata_verir():
    kitap = openpyxl.Workbook()  # varsayılan sayfa adı "Sheet"
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap.close()

    with pytest.raises(RuntimeError, match="bilinen hiçbir sayfa adı yok"):
        _ulke_grubu_dosyasini_ayikla(tampon.getvalue(), "2023.01")


def test_dispatcher_taninmayan_baslik_hata_verir():
    kitap = openpyxl.Workbook()
    sayfa = kitap.active
    sayfa.title = ULKE_GRUBU_YENI_SAYFA_ADI
    sayfa.append(["başlık"])
    sayfa.append([])
    sayfa.append([None])
    sayfa.append(["BEKLENMEYEN"])
    tampon = io.BytesIO()
    kitap.save(tampon)
    kitap.close()

    with pytest.raises(RuntimeError, match="tanınmayan başlık hücresi"):
        _ulke_grubu_dosyasini_ayikla(tampon.getvalue(), "2023.01")
