"""Pegasus Yatırımcı İlişkileri aylık trafik bülteni istemcisi.

Tek XLSX dosyası 2019'dan bugüne TÜM geçmişi taşıyor (TİM/OSD'nin aksine
yıl/ay başına ayrı dosya yok) — bu yüzden `seri_cek` her koşuda dosyayı
BİR KEZ indirip ayrıştırır; 18 seri (3 segment × 6 ölçüt) aynı önbelleği
paylaşır, aksi halde 18 indirme olurdu.

`TRAFİK` sayfasının AYLIK bloğu (satır 2 yıl, satır 3 ay başlığı, satır
4-21 üç segment × altı ölçüt) okunur; segment etiketi (A sütunu) yalnızca
o segmentin İLK satırında yazılıdır, sonraki beş satırda boştur — bu
yüzden "son görülen segment" satır satır taşınır. KÜMÜLATİF blok (satır
25'ten itibaren) ve satır 3'teki AYLIK bloğun ardından gelen yıllık %
değişim sütunu (tamamen boş bir ayraç sütunundan sonra gelir) okunmaz.

Ölçüldü (2026-09-18): Ağustos 2026 Toplam/İç Hat/Dış Hat misafir sayısı
4.51 / 1.61 / 2.90 mn — pegasusyatirimciiliskileri.com kart özetiyle
birebir eşleşiyor.
"""

from __future__ import annotations

import io
import re

import openpyxl
import pandas as pd
import pdfplumber
import requests

from core.catalog import GECERLI_PGSUS_METRIKLERI, GECERLI_PGSUS_OLCUTLERI, GECERLI_PGSUS_SEGMENTLERI
from ingest.ir_sunum import tr_sayi

UC = (
    "https://www.pegasusyatirimciiliskileri.com/medium/image/"
    "pegasus-trafik-karbon-emisyonu-verileri-2019-2026-excel_1571/view.aspx"
)
ZAMAN_ASIMI = 60
SAYFA_ADI = "TRAFİK"

AY_ADLARI = (
    "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
    "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
)

# Ham ölçüt etiketi -> katalogdaki kısa anahtar (core.catalog'daki
# GECERLI_PGSUS_OLCUTLERI ile birebir eşleşmeli).
OLCUT_ESLEME = {
    "Misafir sayısı, mn": "misafir",
    "Konma": "konma",
    "Koltuk sayısı, mn": "koltuk",
    "Doluluk Oranı": "doluluk",
    "ASK (mln km)": "ask",
    "Konma başına Misafir": "konma-basina-misafir",
}
assert set(OLCUT_ESLEME.values()) == GECERLI_PGSUS_OLCUTLERI


def _ay_sutunlarini_bul(satir_yil: tuple, satir_ay: tuple) -> dict[int, str]:
    """C sütunundan başlayıp ilk tamamen boş (yıl, ay) çiftinde duran
    yıl×ay ızgarasını `{sütun indeksi: "YYYY-MM-01"}` olarak döner.

    AYLIK bloğunun ardından bir ayraç (tamamen boş) sütun gelip, onun
    ardından yıllık % değişim gibi türetilmiş bir özet sütunu geliyor —
    ayraç görülünce taramak durur, o özet sütunu hiç okunmaz.
    """
    sutunlar: dict[int, str] = {}
    onceki: tuple[int, int] | None = None
    for sutun in range(2, len(satir_yil)):
        yil, ay_adi = satir_yil[sutun], satir_ay[sutun]
        if yil is None and ay_adi is None:
            break
        if yil is None or ay_adi is None:
            raise RuntimeError(
                f"Pegasus trafik bülteni: eksik yıl/ay başlığı (sütun {sutun})"
            )
        if ay_adi not in AY_ADLARI:
            raise RuntimeError(
                f"Pegasus trafik bülteni: bilinmeyen ay adı {ay_adi!r} (sütun {sutun})"
            )
        if not isinstance(yil, int) or isinstance(yil, bool):
            raise RuntimeError(
                f"Pegasus trafik bülteni: yıl sayı değil {yil!r} (sütun {sutun})"
            )
        ay = AY_ADLARI.index(ay_adi) + 1
        simdiki = (yil, ay)
        if onceki is not None and simdiki <= onceki:
            raise RuntimeError(
                f"Pegasus trafik bülteni: yıl/ay sırası bozuk (sütun {sutun}: "
                f"{simdiki} <= {onceki})"
            )
        onceki = simdiki
        sutunlar[sutun] = f"{yil}-{ay:02d}-01"
    if not sutunlar:
        raise RuntimeError("Pegasus trafik bülteni: yıl/ay ızgarası bulunamadı")
    return sutunlar


def trafik_noktalari(baytlar: bytes) -> dict[tuple[str, str], dict[str, float]]:
    """XLSX baytlarından `{(segment, ölçüt): {"YYYY-MM-01": değer}}` çıkarır.

    Yalnızca AYLIK blok (satır 4-21) okunur; KÜMÜLATİF blok (satır 25+)
    dokunulmaz. Boş hücre (henüz yayımlanmamış ay) sıfıra çevrilmeden
    atlanır.
    """
    kitap = openpyxl.load_workbook(io.BytesIO(baytlar), data_only=True)
    try:
        if SAYFA_ADI not in kitap.sheetnames:
            raise RuntimeError(
                f"Pegasus trafik bülteninde '{SAYFA_ADI}' sayfası yok: {kitap.sheetnames}"
            )
        sayfa = kitap[SAYFA_ADI]
        satirlar = list(sayfa.iter_rows(values_only=True))
        if len(satirlar) < 3 or satirlar[2][0] != "AYLIK":
            raise RuntimeError(
                "Pegasus trafik bülteni: A3 hücresi 'AYLIK' değil — şablon değişmiş olabilir"
            )
        sutunlar = _ay_sutunlarini_bul(satirlar[1], satirlar[2])

        noktalar: dict[tuple[str, str], dict[str, float]] = {}
        mevcut_segment: str | None = None
        for satir in satirlar[3:]:
            segment_ham, olcut_ham = satir[0], satir[1]
            if segment_ham is None and olcut_ham is None:
                break  # AYLIK blok bitti (KÜMÜLATİF başlığından önceki boş satırlar)
            if segment_ham is not None:
                if segment_ham not in GECERLI_PGSUS_SEGMENTLERI:
                    raise RuntimeError(
                        f"Pegasus trafik bülteni: bilinmeyen segment {segment_ham!r}"
                    )
                mevcut_segment = segment_ham
            if mevcut_segment is None:
                raise RuntimeError(
                    "Pegasus trafik bülteni: segment etiketinden önce ölçüt satırı geldi"
                )
            if olcut_ham not in OLCUT_ESLEME:
                raise RuntimeError(
                    f"Pegasus trafik bülteni: bilinmeyen ölçüt {olcut_ham!r} "
                    f"(segment {mevcut_segment!r})"
                )
            olcut = OLCUT_ESLEME[olcut_ham]
            aylik: dict[str, float] = {}
            for sutun, tarih in sutunlar.items():
                hucre = satir[sutun]
                if hucre is None:
                    continue  # yayımlanmamış ay — sıfır değil, atlanır
                if isinstance(hucre, bool) or not isinstance(hucre, (int, float)):
                    raise RuntimeError(
                        f"Pegasus trafik bülteni: sayısal olmayan hücre "
                        f"{mevcut_segment}/{olcut}/{tarih}: {hucre!r}"
                    )
                aylik[tarih] = float(hucre)
            anahtar = (mevcut_segment, olcut)
            if anahtar in noktalar:
                raise RuntimeError(
                    f"Pegasus trafik bülteni: yinelenen segment/ölçüt satırı {anahtar!r}"
                )
            noktalar[anahtar] = aylik

        beklenen = {
            (segment, olcut)
            for segment in GECERLI_PGSUS_SEGMENTLERI
            for olcut in GECERLI_PGSUS_OLCUTLERI
        }
        if set(noktalar) != beklenen:
            eksik = beklenen - set(noktalar)
            raise RuntimeError(
                f"Pegasus trafik bülteni: eksik segment/ölçüt satırı {sorted(eksik)}"
            )
        return noktalar
    finally:
        kitap.close()


def _dosya_indir(session=None) -> bytes:
    http = session or requests
    yanit = http.get(UC, headers={"User-Agent": "Mozilla/5.0"}, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"Pegasus trafik bülteni HTTP {yanit.status_code}")
    return yanit.content


def seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse dosya indirme + ayrıştırma koşu boyunca
    paylaşılır: 18 seri aynı tek dosyayı okuduğu için yoksa 18 indirme
    olurdu.
    """
    onbellek = {} if onbellek is None else onbellek
    if "noktalar" not in onbellek:
        baytlar = _dosya_indir(session)
        onbellek["noktalar"] = trafik_noktalari(baytlar)
    noktalar = onbellek["noktalar"]

    anahtar = (seri.pgsus_segment, seri.pgsus_olcut)
    if anahtar not in noktalar:
        raise RuntimeError(
            f"Pegasus trafik bülteninde segment/ölçüt bulunamadı: {anahtar!r} ({seri.id})"
        )

    df = pd.DataFrame(sorted(noktalar[anahtar].items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)


# --- Çeyreklik Yatırımcı Sunumu (kaynak_tipi: pgsus_ir) ---
#
# Kaynak gerçekleri 2026-09-20'de canlı ölçüldü (2Ç'26 sunumu,
# "2026 2. Çeyrek Yatırımcı Sunumu", 26 sayfa):
#
# - "Yatırımcı Sunumlarımız" sayfası (`SUNUMLARIMIZ_SAYFASI`) en yeni
#   çeyrekten en eskiye sıralı `"<YIL> <Ç>. Çeyrek Yatırımcı Sunumu"`
#   bağlantıları listeler; URL slug'ı ÇEYREĞİ GÜVENİLİR YANSITMIYOR (ör.
#   2026 1Ç dosyasının slug'ı "pgsusciptur_1676" — yıl/çeyrek bilgisi YOK).
#   Bu yüzden yıl/çeyrek LİNK METNİNDEN ayrıştırılır, URL'den değil. HTML
#   şablonunda `class="ico-pdf %>"` gibi işlenmemiş bir kalıp `>` karakteri
#   içeriyor — `[^<]*` kullanılır (`[^>]*` bu yüzden erken durur).
# - "OPERASYONEL VE FİNANSAL VERİLER" sayfası (bkz. `_BUYUK_TABLO_SATIRLARI`)
#   TEK tabloda 10 ÇEYREK (1Ç24..cari çeyrek) taşıyor — THY'nin aksine tam
#   geçmiş TEK PDF'ten çıkıyor. Bu tablo İNGİLİZCE sayı biçimi kullanıyor
#   (virgül binlik, nokta ondalık — ör. "1,091" = 1091, "4.49" = 4.49);
#   şirketin diğer tüm sayfaları (ve THY'nin tamamı) TÜRKÇE biçim kullanır
#   (`tr_sayi`) — KARIŞTIRILMAMALI, bu yüzden ayrı `_intl_sayi`.
# - "BİLANÇO YAPISI" sayfası yalnızca ÜÇ dönem taşır (31 Aralık 2024/2025 +
#   cari çeyrek sonu) — İNGİLİZCE biçim burada da geçerli.
# - "NAKİT POZİSYONU" sayfasındaki "Artı nakit pozisyonu ... sonunda X
#   milyon Euro'dur" cümlesi TEK nokta verir (yalnızca cari dönem düz
#   cümlede; önceki dönemler yalnızca grafik değer etiketi olarak gömülü,
#   metin akışı güvenilmez — okunmaz).
# - "FİLO" sayfasında tip bazlı tablo ("Toplam" sütunu) + "X,X yıl ortalama
#   yaş" cümlesi + "Sipariş planı: N A320neo (...), N A321neo (...) ve N
#   Boeing MAX-10 (...)" cümlesi aynı sayfada. Sipariş cümlesi filo
#   tablosuyla YATAY olarak iç içe iki sütunlu sayfa düzeninde geldiği için
#   metin akışında ARAYA tablo satırları karışabiliyor (ölçüldü) — bu
#   yüzden durum ifadesi ("hepsi teslim alındı" / "N adet teslim alındı" /
#   "teslimatlar henüz başlamadı") ayrıştırılırken parantez içindeki TÜM
#   metin arandığı için bu karışma sonucu etkilemiyor.
# - Ölçüldü (2026-09-20): Satış Gelirleri 865 / FAVÖK 80 / Net Kâr -92 /
#   RASK 4,49 / Net Borç 3.156 / Artı Nakit 548 / A321neo filo 70 / Ort.
#   Yaş 5,5 / A321neo sipariş bakiyesi 39 — marketvisuals.net kart
#   özetleriyle birebir eşleşiyor.

SUNUMLARIMIZ_SAYFASI = "https://www.pegasusyatirimciiliskileri.com/tr/operasyonel-ve-finansal-veriler/yatirimci-sunumlarimiz"
_SUNUM_BAGLANTISI = re.compile(r'<a href="([^"]+)"[^<]*>(\d{4}) (\d)\. Çeyrek Yatırımcı Sunumu</a>')

_CEYREK_ILK_AY = {1: "01", 2: "04", 3: "07", 4: "10"}
_AY_CEYREGI = {
    "Ocak": 1, "Şubat": 1, "Mart": 1, "Nisan": 2, "Mayıs": 2, "Haziran": 2,
    "Temmuz": 3, "Ağustos": 3, "Eylül": 3, "Ekim": 4, "Kasım": 4, "Aralık": 4,
}


def _intl_sayi(ham: str) -> float:
    """Uluslararası biçimli sayıyı float'a çevirir: virgül binlik ayıracı,
    nokta ondalık (`tr_sayi`'nin TAM TERSİ biçim) — yalnızca "OPERASYONEL VE
    FİNANSAL VERİLER" ve "BİLANÇO YAPISI" tabloları için (bkz. modül
    docstring'i)."""
    return float(ham.replace(",", ""))


def _ceyrek_tarihi(yil: int, ceyrek: int) -> str:
    """Katalog sözleşmesi: çeyreksel seri tarihi çeyreğin İLK ayı."""
    return f"{yil}-{_CEYREK_ILK_AY[ceyrek]}-01"


def _ceyrekleri_uret(yil: int, ceyrek: int, adet: int) -> list[tuple[int, int]]:
    """`(yil, ceyrek)`ten geriye doğru `adet` çeyreği ESKİDEN YENİYE üretir."""
    sonuc = []
    y, c = yil, ceyrek
    for _ in range(adet):
        sonuc.append((y, c))
        c -= 1
        if c == 0:
            c, y = 4, y - 1
    return list(reversed(sonuc))


def sunum_bilgisi(session=None) -> tuple[str, int, int]:
    """Yatırımcı Sunumlarımız sayfasından en güncel çeyreklik sunumun
    `(url, yıl, çeyrek)` bilgisini döner; yıl/çeyrek LİNK METNİNDEN
    ayrıştırılır (URL slug'ı güvenilir değil — bkz. modül docstring'i)."""
    http = session or requests
    yanit = http.get(SUNUMLARIMIZ_SAYFASI, headers={"User-Agent": "Mozilla/5.0"}, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"Pegasus sunumlar sayfası HTTP {yanit.status_code} döndü")
    m = _SUNUM_BAGLANTISI.search(yanit.text)
    if not m:
        raise RuntimeError("Pegasus sunumlar sayfasında çeyreklik yatırımcı sunumu bağlantısı bulunamadı")
    url = m.group(1)
    return (url if url.startswith("http") else f"https://www.pegasusyatirimciiliskileri.com{url}"), int(m.group(2)), int(m.group(3))


def _sayfa_metni_bul(pdf, icerir: str) -> tuple[int, str]:
    """Verilen alt dizeyi içeren İLK sayfayı bulur; `(sayfa indeksi, metin)`
    döner."""
    for i, sayfa in enumerate(pdf.pages):
        metin = sayfa.extract_text() or ""
        if icerir in metin:
            return i, metin
    raise RuntimeError(f"Pegasus sunumunda {icerir!r} içeren sayfa bulunamadı (şablon değişmiş olabilir)")


# "OPERASYONEL VE FİNANSAL VERİLER" tablosu: etiket -> pgsus_metrik, 13
# alan (12A24, 12A25, %değ, 1Ç24..2Ç26) yakalanır; yalnızca son 10'u
# (1Ç24..2Ç26 çeyreklik sütunlar) kullanılır.
_BUYUK_TABLO_SATIRLARI = {
    "Toplam Gelirler (mln Euro)": "satis-gelirleri",
    "Yan Gelirler (mln Euro)": "yan-gelirler",
    "FAVÖK (mln Euro)*": "favok",
    "Net kar/zarar": "net-kar",
    "RASK, (€c)": "rask",
    "CASK, (€c)": "cask",
}
_BUYUK_TABLO_BASLIK_DESENI = re.compile(
    r"12A 12A % değ\. 1Ç 2Ç 3Ç 4Ç 1Ç 2Ç 3Ç 4Ç 1Ç 2Ç % değ\."
)
_BUYUK_TABLO_ALAN = r"(\S+)"


def _buyuk_tablo_noktalari(metin: str, ceyrekler: list[tuple[int, int]]) -> dict[str, dict[str, float]]:
    """10 çeyreklik özet tablosundan `{pgsus_metrik: {tarih: değer}}` çıkarır."""
    if not _BUYUK_TABLO_BASLIK_DESENI.search(metin):
        raise RuntimeError("Pegasus sunumunda OPERASYONEL VE FİNANSAL VERİLER başlık sütunları beklenenden farklı")
    sonuc: dict[str, dict[str, float]] = {}
    for etiket, metrik in _BUYUK_TABLO_SATIRLARI.items():
        desen = re.compile(re.escape(etiket) + r"\s+" + r"\s+".join([_BUYUK_TABLO_ALAN] * 13))
        m = desen.search(metin)
        if not m:
            raise RuntimeError(f"Pegasus sunumunda '{metrik}' için tablo satırı bulunamadı ({etiket!r})")
        ceyrek_degerleri = m.groups()[3:13]
        if len(ceyrek_degerleri) != len(ceyrekler):
            raise RuntimeError(
                f"Pegasus sunumunda '{metrik}' çeyrek sayısı uyuşmuyor: "
                f"{len(ceyrek_degerleri)} değer, {len(ceyrekler)} çeyrek"
            )
        sonuc[metrik] = {
            _ceyrek_tarihi(yil, ceyrek): _intl_sayi(deger)
            for (yil, ceyrek), deger in zip(ceyrekler, ceyrek_degerleri)
        }
    return sonuc


# "BİLANÇO YAPISI" sayfası: üç dönem sütunu. Ham metinde satır sırası
# "31 Aralık 31 Aralık 30 <Ay>" / "Milyon Euro" / "<yıl1> <yıl2> <yıl3>"
# şeklinde geliyor (sütun başlıkları ve "Milyon Euro" etiketi farklı metin
# bloklarında, tablo çizgisiz) — ay adı üçüncü (cari) döneme ait.
_BILANCO_BASLIK_DESENI = re.compile(
    r"31 Aralık 31 Aralık 30 (\w+)\s*Milyon Euro\s*(\d{4}) (\d{4}) (\d{4})"
)
_NET_BORC_DESENI = re.compile(r"Net Borç, mn Euro\s+(\S+)\s+(\S+)\s+(\S+)")


def _net_borc_noktalari(metin: str) -> dict[str, float]:
    baslik = _BILANCO_BASLIK_DESENI.search(metin)
    if not baslik:
        raise RuntimeError("Pegasus sunumunda BİLANÇO YAPISI başlık sütunları (31 Aralık.../30 <Ay>...) bulunamadı")
    ay3, yil1, yil2, yil3 = baslik.groups()
    ceyrek3 = _AY_CEYREGI.get(ay3)
    if ceyrek3 is None:
        raise RuntimeError(f"Pegasus sunumunda bilinmeyen ay adı: {ay3!r}")
    tarihler = [f"{yil1}-10-01", f"{yil2}-10-01", _ceyrek_tarihi(int(yil3), ceyrek3)]

    m = _NET_BORC_DESENI.search(metin)
    if not m:
        raise RuntimeError("Pegasus sunumunda BİLANÇO YAPISI 'Net Borç, mn Euro' satırı bulunamadı")
    return dict(zip(tarihler, (_intl_sayi(g) for g in m.groups())))


# "NAKİT POZİSYONU" sayfası: yalnızca cari dönem cümlesi.
_ARTI_NAKIT_DESENI = re.compile(r"Artı nakit pozisyonu (\w+) (\d{4}) sonunda ([\d.,]+) milyon Euro")


def _arti_nakit_noktasi(metin: str, yil: int, ceyrek: int) -> dict[str, float]:
    m = _ARTI_NAKIT_DESENI.search(metin)
    if not m:
        raise RuntimeError("Pegasus sunumunda NAKİT POZİSYONU 'Artı nakit pozisyonu ... sonunda' cümlesi bulunamadı")
    ay_adi, cumle_yil, deger = m.groups()
    ceyrek_kontrol = _AY_CEYREGI.get(ay_adi)
    if ceyrek_kontrol is None:
        raise RuntimeError(f"Pegasus sunumunda bilinmeyen ay adı: {ay_adi!r}")
    if (int(cumle_yil), ceyrek_kontrol) != (yil, ceyrek):
        raise RuntimeError(
            f"Pegasus sunumunda Artı Nakit cümlesinin dönemi ({cumle_yil} {ay_adi}) "
            f"sunum meta verisiyle ({yil} Ç{ceyrek}) uyuşmuyor"
        )
    return {_ceyrek_tarihi(yil, ceyrek): tr_sayi(deger)}


# "FİLO" sayfası: tip bazlı Toplam sütunu + ortalama yaş + sipariş bakiyesi.
_FILO_TIP_DESENLERI = {
    "filo-b737-800": re.compile(r"Boeing 737-800\s+[\d-]+\s+[\d-]+\s+[\d-]+\s+(\d+)"),
    "filo-a320-ceo": re.compile(r"Airbus A320ceo\s+[\d-]+\s+[\d-]+\s+[\d-]+\s+(\d+)"),
    "filo-a320-neo": re.compile(r"Airbus A320neo\s+[\d-]+\s+[\d-]+\s+[\d-]+\s+(\d+)"),
    "filo-a321-neo": re.compile(r"Airbus A321neo\s+[\d-]+\s+[\d-]+\s+[\d-]+\s+(\d+)"),
}
_FILO_TOPLAM_DESENI = re.compile(r"^Toplam\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)$", re.M)
_FILO_YAS_DESENI = re.compile(r"([\d,]+)\s*yıl ortalama yaş")
_SIPARIS_DESENI = re.compile(
    r"(\d+)\s+A320neo\s*\(([^)]*)\).*?"
    r"(\d+)\s+A321neo\s*\(([^)]*)\).*?"
    r"(\d+)\s+Boeing MAX-10\s*\(([^)]*)\)",
    re.DOTALL,
)


def _teslim_sayisi(durum_metni: str, siparis_adedi: int) -> int:
    if "hepsi teslim alındı" in durum_metni:
        return siparis_adedi
    if "henüz başlamadı" in durum_metni:
        return 0
    m = re.search(r"(\d+)\s+adet", durum_metni)
    if m:
        return int(m.group(1))
    raise RuntimeError(f"Pegasus sunumunda sipariş teslim durumu tanınamadı: {durum_metni!r}")


def _filo_noktalari(metin: str, bu_tarih: str) -> dict[str, dict[str, float]]:
    sonuc: dict[str, dict[str, float]] = {}

    tip_toplamlari = {}
    for metrik, desen in _FILO_TIP_DESENLERI.items():
        m = desen.search(metin)
        if not m:
            raise RuntimeError(f"Pegasus sunumunda Filo tablosunda '{metrik}' satırı bulunamadı")
        tip_toplamlari[metrik] = int(m.group(1))
        sonuc[metrik] = {bu_tarih: float(m.group(1))}
    genel = _FILO_TOPLAM_DESENI.search(metin)
    if not genel:
        raise RuntimeError("Pegasus sunumunda Filo tablosunda 'Toplam' satırı bulunamadı")
    if sum(tip_toplamlari.values()) != int(genel.group(4)):
        raise RuntimeError(
            f"Pegasus sunumunda Filo tip toplamları ({sum(tip_toplamlari.values())}) "
            f"Toplam satırı ({genel.group(4)}) ile uyuşmuyor"
        )

    yas_m = _FILO_YAS_DESENI.search(metin)
    if not yas_m:
        raise RuntimeError("Pegasus sunumunda 'X yıl ortalama yaş' cümlesi bulunamadı")
    sonuc["filo-ortalama-yas"] = {bu_tarih: tr_sayi(yas_m.group(1))}

    siparis_m = _SIPARIS_DESENI.search(metin)
    if not siparis_m:
        raise RuntimeError("Pegasus sunumunda 'Sipariş planı' cümlesi bulunamadı")
    a320neo_adet, a320neo_durum, a321neo_adet, a321neo_durum, max10_adet, max10_durum = siparis_m.groups()
    sonuc["siparis-bakiye-a320neo"] = {
        bu_tarih: float(int(a320neo_adet) - _teslim_sayisi(a320neo_durum, int(a320neo_adet)))
    }
    sonuc["siparis-bakiye-a321neo"] = {
        bu_tarih: float(int(a321neo_adet) - _teslim_sayisi(a321neo_durum, int(a321neo_adet)))
    }
    sonuc["siparis-bakiye-max10"] = {
        bu_tarih: float(int(max10_adet) - _teslim_sayisi(max10_durum, int(max10_adet)))
    }
    return sonuc


def _sunumu_ayristir(pdf, yil: int, ceyrek: int) -> dict[str, dict[str, float]]:
    """Açılmış sunum PDF'inden TÜM pgsus_metrik noktalarını tek geçişte
    çıkarır: `{pgsus_metrik: {tarih: değer}}`."""
    bu_tarih = _ceyrek_tarihi(yil, ceyrek)
    noktalar: dict[str, dict[str, float]] = {}

    _, metin = _sayfa_metni_bul(pdf, "OPERASYONEL VE FİNANSAL VERİLER")
    ceyrekler = _ceyrekleri_uret(yil, ceyrek, 10)
    noktalar.update(_buyuk_tablo_noktalari(metin, ceyrekler))

    _, metin = _sayfa_metni_bul(pdf, "BİLANÇO YAPISI")
    noktalar["net-borc"] = _net_borc_noktalari(metin)

    _, metin = _sayfa_metni_bul(pdf, "NAKİT POZİSYONU")
    noktalar["arti-nakit"] = _arti_nakit_noktasi(metin, yil, ceyrek)

    _, metin = _sayfa_metni_bul(pdf, "Sipariş planı")
    noktalar.update(_filo_noktalari(metin, bu_tarih))

    eksik = GECERLI_PGSUS_METRIKLERI - set(noktalar)
    if eksik:
        raise RuntimeError(f"Pegasus sunumunda çıkarılamayan metrikler: {sorted(eksik)}")
    return noktalar


def _sunum_noktalarini_getir(onbellek: dict, session=None) -> dict[str, dict[str, float]]:
    """PDF indirme + ayrıştırmayı önbellekler: 16 seri aynı tek dosyayı
    okuduğu için yoksa 16 indirme olurdu."""
    if "noktalar" in onbellek:
        return onbellek["noktalar"]
    http = session or requests
    url, yil, ceyrek = sunum_bilgisi(session)
    yanit = http.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"Pegasus yatırımcı sunumu indirilemedi ({url}): HTTP {yanit.status_code}")
    with pdfplumber.open(io.BytesIO(yanit.content)) as pdf:
        onbellek["noktalar"] = _sunumu_ayristir(pdf, yil, ceyrek)
    return onbellek["noktalar"]


def sunum_seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Pegasus çeyreklik yatırımcı sunumundan tek bir `pgsus_metrik`
    serisini çeker. `onbellek` verilirse PDF indirme + ayrıştırma koşu
    boyunca paylaşılır (bkz. `_sunum_noktalarini_getir`)."""
    onbellek = {} if onbellek is None else onbellek
    tum_noktalar = _sunum_noktalarini_getir(onbellek, session=session)
    noktalar = tum_noktalar.get(seri.pgsus_metrik)
    if not noktalar:
        raise RuntimeError(f"Pegasus sunumunda pgsus_metrik bulunamadı: {seri.pgsus_metrik!r} ({seri.id})")
    df = pd.DataFrame(sorted(noktalar.items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
