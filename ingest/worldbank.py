"""Dünya Bankası "Pink Sheet" (Emtia Fiyat Verisi) Excel istemcisi.

Kaynak: https://www.worldbank.org/en/research/commodity-markets — sayfa her
ay güncellenen `CMO-Historical-Data-Monthly.xlsx` dosyasına bağlantı verir;
dosyanın URL'si aylık yayın numarasını (GUID) içerdiği için sabit değildir,
bu yüzden önce arşiv sayfası taranıp güncel bağlantı bulunur (bkz.
`ingest/ttkom.py::en_guncel_dosya_url` ile aynı desen).

Dosyada iki ilgili sayfa var:
- "Monthly Indices": endeksler (2010=100), ör. Gübre Endeksi.
- "Monthly Prices": tekil emtia fiyatları (USD), ör. Urea, DAP.

Her iki sayfada da veri satırları `YYYYMmm` biçiminde ay etiketiyle başlar
(ör. "2026M08"); tarih ayın 1'ine damgalanır. Sütun başlıkları çok satırlı
ve `**`/boşluk gibi dipnot işaretleri taşıdığından karşılaştırma
normalize edilerek yapılır (bkz. `_normalize`).
"""

from __future__ import annotations

import io
import re

import openpyxl
import pandas as pd
import requests

from core.catalog import GECERLI_WB_SERILERI, Seri

LANDING_SAYFASI = "https://www.worldbank.org/en/research/commodity-markets"
ZAMAN_ASIMI = 60
_BASLIKLAR = {"User-Agent": "Mozilla/5.0"}

_DOSYA_BAGLANTISI = re.compile(r'href="([^"]+CMO-Historical-Data-Monthly\.xlsx)"', re.I)
_AY_ETIKETI = re.compile(r"^(\d{4})M(\d{2})$")

# wb_seri -> (sayfa adı, sütun başlığı — normalize edilmiş haliyle karşılaştırılır)
SERI_TANIMLARI = {
    "gubre-endeksi": ("Monthly Indices", "Fertilizers"),
    "urea": ("Monthly Prices", "Urea"),
    "dap": ("Monthly Prices", "DAP"),
    # Ölçüldü (2026-09-19): marketvisuals'ın "commodity_other.html" (Business
    # Analytiq kaynaklı, ücretli) sayfasındaki "Doğal Kauçuk - ABD
    # (Gerçekleşen)" kartı ile bu sütun 13 aylık pencerede oran 0,98–1,05
    # bandında (ortalama ~1,01, trend sapması yok) — aynı küresel TSR20
    # benchmark'ı (SGX/SICOM). "Avrupa" kartı ile Rubber RSS3 arasındaki
    # oran ise 0,70–0,95 aralığında SÜREKLİ AÇILIYOR (aynı seri değil, eşleme
    # YOK — bu yüzden yalnızca TSR20/ABD eklendi, RSS3/Avrupa eklenmedi).
    "kaucuk-tsr20": ("Monthly Prices", "Rubber, TSR20"),
}
assert set(SERI_TANIMLARI.keys()) == GECERLI_WB_SERILERI


def _normalize(metin: str) -> str:
    return re.sub(r"[\s*]+$", "", metin or "").strip().lower()


def en_guncel_dosya_url(session=None) -> str:
    """Pink Sheet açılış sayfasındaki güncel xlsx bağlantısı."""
    http = session or requests
    yanit = http.get(LANDING_SAYFASI, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"Dünya Bankası Pink Sheet sayfası HTTP {yanit.status_code}")
    eslesme = _DOSYA_BAGLANTISI.search(yanit.text)
    if not eslesme:
        raise RuntimeError(
            "Dünya Bankası Pink Sheet sayfasında CMO-Historical-Data-Monthly.xlsx "
            "bağlantısı bulunamadı — sayfa yapısı değişmiş olabilir"
        )
    return eslesme.group(1)


def _kitabi_getir(onbellek: dict, session=None):
    if "kitap" in onbellek:
        return onbellek["kitap"]
    http = session or requests
    url = en_guncel_dosya_url(session=session)
    yanit = http.get(url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"Pink Sheet Excel'i indirilemedi ({url}): HTTP {yanit.status_code}")
    onbellek["kitap"] = openpyxl.load_workbook(io.BytesIO(yanit.content), data_only=True)
    return onbellek["kitap"]


def _sutun_indeksi(ws, etiket: str) -> int:
    """İlk 10 satırı tarayıp `etiket`i (normalize edilmiş) taşıyan sütunu bulur."""
    hedef = _normalize(etiket)
    for satir in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        for i, hucre in enumerate(satir):
            if isinstance(hucre, str) and _normalize(hucre) == hedef:
                return i
    raise RuntimeError(f"Pink Sheet '{ws.title}' sayfasında '{etiket}' sütunu bulunamadı")


def seri_cek(seri: Seri, *, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse tek Excel dosyası (gübre endeksi + urea + DAP aynı
    dosyayı paylaşır) koşu boyunca bir kez indirilir/ayrıştırılır.
    """
    onbellek = {} if onbellek is None else onbellek
    kitap = _kitabi_getir(onbellek, session=session)
    wb_seri = seri.wb_seri
    if wb_seri not in SERI_TANIMLARI:
        raise RuntimeError(f"Bilinmeyen wb_seri: {wb_seri!r} (seri={seri.id})")
    sayfa_adi, etiket = SERI_TANIMLARI[wb_seri]
    ws = kitap[sayfa_adi]
    sutun = _sutun_indeksi(ws, etiket)

    noktalar: list[tuple[str, float]] = []
    for satir in ws.iter_rows(values_only=True):
        eslesme = _AY_ETIKETI.match(satir[0]) if isinstance(satir[0], str) else None
        if not eslesme:
            continue
        deger = satir[sutun] if sutun < len(satir) else None
        if deger is None or not isinstance(deger, (int, float)):
            continue
        yil, ay = eslesme.groups()
        noktalar.append((f"{yil}-{ay}-01", float(deger)))

    if not noktalar:
        raise RuntimeError(f"Pink Sheet '{wb_seri}' için veri noktası bulunamadı (seri={seri.id})")

    df = pd.DataFrame(sorted(noktalar), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
