"""EİB (Ege İhracatçı Birlikleri) Su Ürünleri ve Hayvansal Mamuller
İhracatçıları Birliği (ESÜHMİB) aylık ihracat istatistiği istemcisi.

Kaynak sayfası (`https://www.eib.org.tr/Sayfa.Asp?id=AF5247B95C`, "İstatistikler")
JS ile render ediliyor (SPA) — ama arkasındaki içerik `requests` ile de
çekilebiliyor: sayfanın kendi `assets/api.js`'i içeriği
`GET /Sayfa_Icerik_JSON.Asp?SI_Id=AF5247B95C` ile çekiyor ve dönen JSON'un
`icerik` alanı, yıl başına `<h3>YIL</h3>` ile bölünmüş, ay başına kısa
bağlantılar (`eib.li/XXXXX`) taşıyan düz HTML. Bir kısa bağlantı bozuk geldi
(ölçüldü, Şubat 2024: `http://https//eib.li/BE6FE`) — `_eib_li_duzelt`
kod kısmını regex ile çıkarıp temiz URL'i yeniden kurar.

Her ayın "EİB Aylık" bağlantısı KÜMÜLATİFTİR: örn. Ağustos 2026 dosyası
Ocak–Ağustos 2026'nın TAMAMINI (ay ay) taşır, tek dosya indirmek o yılın
o ana kadarki tüm aylarını verir (ölçüldü: Ağustos dosyasının ham veri
sayfasında "Aylar" sütununda OCAK..AĞUSTOS'un hepsi var). Bu yüzden yıl
başına TEK indirme yeterli — tim.py'deki "yıl başına tek bülten" deseniyle
aynı mantık: geçmiş yıllar için ARALIK, cari yıl için yayımlanan en son ay.

Ham veri her yılın dosyasında AYRI bir sayfada ("GTIP ULKE GB TARIH" ya da
"Sayfa1" — sayfa ADI yıldan yıla değişiyor, bu yüzden `_ham_veri_sayfasini_bul`
başlık satırının ilk hücresi "AYLAR" olan sayfayı arar) satır bazlı GTİP
kayıtları olarak geliyor: `Aylar, ürün grubu, alt grup 1, ..., AGIRLIK
(KG), FOBUSD, ...`. "ürün grubu" == "SU ÜRÜNLERİ" olan satırların "alt
grup 1"i Levrek/Çipura/Türk Somonu/Alabalık/Kaya Levreği/Diğer Su
Ürünleri/(nadiren) Orkinos gibi ürünlere ayrılıyor.

**Format kırılması (ölçüldü 2026-09-18): "alt grup 1" kırılımı yalnızca
2025 Ocak'tan itibaren var.** 2024 dosyasında "alt grup 1" sütunu SU
ÜRÜNLERİ'yi tekrarlıyor (gerçek kırılım yok) ve ham veri sayfası farklı bir
şema taşıyor (ekstra "alt grup 3" sütunu, TİM/EPDK tarzı bir format
değişimi). Bu yüzden `ILK_YIL = 2025` — TİM/EPDK'nın "yalnızca stabil
formatın geçerli olduğu pencere" ilkesiyle aynı karar.

Ürün kırılımı: Levrek/Çipura/Türk Somonu/Alabalık/Kaya Levreği DOĞRUDAN
izlenir; geri kalan HER ŞEY (kaynağın kendi "Diğer Su Ürünleri" satırı dahil,
nadiren görülen Orkinos gibi kalemler dahil) `DİĞER SU ÜRÜNLERİ` = SU
ÜRÜNLERİ toplamı − 5 adlandırılmış ürün olarak HESAPLANIR — kaynak yeni bir
ürün eklerse (Orkinos gibi) adaptör değişmeden otomatik "diğer"e düşer.
Ölçüldü (Ağustos 2026): Levrek 56.4262 mn$, Çipura 42.9848 mn$ — MarketVisuals
`eib_balikcilik.html` referans değerleriyle birebir.

Hayvansal 8 ana grup (Kanatlı, Yumurta, Süt ve Süt Ürünleri, Sosis ve
Benzeri Ürünler, Bal, Diğer, Canlı Hayvan, Kırmızı Et ve Sakatat) doğrudan
"ürün grubu" adlarıyla izlenir. `HAYVANSAL_TOPLAM` bu 8 grubun TOPLANMASIYLA
hesaplanır (kaynağın "Genel Toplam"ı SU ÜRÜNLERİ'ni de içerdiğinden
doğrudan kullanılamaz — ihracat-il/ihracat-ulke'deki "toplamı adaptör
hesaplar" deseniyle aynı).

# ponytail: "DİĞER ETLER" grubu bilinçli DIŞARIDA bırakıldı — yıl içinde
yalnızca birkaç ayda satırı var (ör. Haziran 2026), kalan aylarda satırın
kendisi YOK (sıfır değil, yokluk — kaynak sıfır satırı hiç yazmıyor). Sürekli
bir seri olarak izlenemeyecek kadar seyrek; HAYVANSAL_TOPLAM o birkaç ayda
birkaç bin dolar eksik kalabilir. Kaynak sürekli aylık satır yayımlamaya
başlarsa `HAYVANSAL_GRUPLARI`'na eklenir.

Ağustos 2026 ölçümü (kaynakla birebir): SU ÜRÜNLERİ toplam FOBUSD
146.033.708,42 (146,0337 mn$), AGIRLIK 17.242.615,26 kg (17,2426 bin ton);
HAYVANSAL_TOPLAM FOBUSD 29.949.152,00 (29,9492 mn$), AGIRLIK 19.326.789,59 kg
(19,3268 bin ton); Kanatlı 14.893.488,38 (14,8935 mn$); Yumurta
8.610.178,29 (8,6102 mn$); Süt ve Süt Ürünleri 2.982.044,51 (2,9820 mn$).

Ham AGIRLIK KG'dir — EPDK'daki gibi adaptör TONA çevirir (/1000); "Bin ton"
gösterimi katalogdaki `olcek: 0.001` ile. FOBUSD zaten USD; "Milyon USD"
gösterimi `olcek: 0.000001` ile.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from io import BytesIO

import openpyxl
import pandas as pd
import requests

from core.catalog import GECERLI_EIB_KALEMLERI, GECERLI_EIB_OLCUTLERI

ISTATISTIK_SI_ID = "AF5247B95C"
ICERIK_URL = "https://www.eib.org.tr/Sayfa_Icerik_JSON.Asp"
ZAMAN_ASIMI = 60

# Kaynak HTML'inin `<strong>Ay: </strong>` etiketleriyle yazdığı biçim
# (başlık-harfli); XLSX'teki "Aylar" sütunu BÜYÜK harfli — ikisi burada
# eşlenir, Türkçe İ/I büyütme tuzağına (`str.upper()`) düşmeden.
AY_ESLEME = {
    "Ocak": "OCAK", "Şubat": "ŞUBAT", "Mart": "MART", "Nisan": "NİSAN",
    "Mayıs": "MAYIS", "Haziran": "HAZİRAN", "Temmuz": "TEMMUZ",
    "Ağustos": "AĞUSTOS", "Eylül": "EYLÜL", "Ekim": "EKİM",
    "Kasım": "KASIM", "Aralık": "ARALIK",
}
AY_ADLARI = tuple(AY_ESLEME.values())
AY_NO = {ad: i + 1 for i, ad in enumerate(AY_ADLARI)}

# Su ürünleri kırılımı yalnızca 2025 Ocak'tan itibaren var (bkz. modül
# docstring'i) — 2024 ve öncesi farklı/eksik şema taşıyor.
ILK_YIL = 2025

SU_URUNLERI_GRUBU = "SU ÜRÜNLERİ"
SU_URUNLERI_ADLARI = ("LEVREK", "ÇİPURA", "TÜRK SOMONU", "ALABALIK", "KAYA LEVREĞİ")
DIGER_SU_URUNLERI = "DİĞER SU ÜRÜNLERİ"

HAYVANSAL_GRUPLARI = (
    "KANATLI", "YUMURTA", "SÜT VE SÜT ÜRÜNLERİ",
    "SOSİS VE BENZERİ ÜRÜNLER (KIRMIZI ET VE KANATLI)",
    "BAL", "DİĞER", "CANLI HAYVAN", "KIRMIZI ET VE SAKATAT",
)
HAYVANSAL_TOPLAM = "HAYVANSAL_TOPLAM"

assert set(SU_URUNLERI_ADLARI) | {SU_URUNLERI_GRUBU, DIGER_SU_URUNLERI} | set(HAYVANSAL_GRUPLARI) | {HAYVANSAL_TOPLAM} == GECERLI_EIB_KALEMLERI
assert {"fobusd", "agirlik"} == GECERLI_EIB_OLCUTLERI


def _eib_li_duzelt(ham_url: str) -> str:
    """Bazı kaynak bağlantıları bozuk geliyor (ölçüldü, Şubat 2024:
    `http://https//eib.li/BE6FE`). Kısa linkin kod kısmı regex ile çıkarılıp
    temiz `https://eib.li/<kod>` yeniden kurulur."""
    m = re.search(r"eib\.li/([A-Za-z0-9]+)", ham_url)
    if not m:
        raise RuntimeError(f"EİB: tanınmayan bağlantı biçimi: {ham_url!r}")
    return f"https://eib.li/{m.group(1)}"


def yillara_ayir(icerik_html: str) -> dict[str, str]:
    """`<h3>YIL</h3>` başlıklarına göre yıl bazlı HTML dilimlerine böler."""
    parcalar = re.split(r"<h3>(\d{4})</h3>", icerik_html)
    return {parcalar[i]: parcalar[i + 1] for i in range(1, len(parcalar), 2)}


def ay_baglantilarini_cikar(yil_html: str) -> dict[str, dict[str, str]]:
    """Bir yılın HTML diliminden `{AY_ADI(XLSX biçimi): {ETİKET: url}}` çıkarır."""
    sonuc: dict[str, dict[str, str]] = {}
    for ay_m in re.finditer(r"<strong>([^:<]+):\s*</strong>(.*?)</p>", yil_html, re.S):
        ay_ham = ay_m.group(1).strip()
        ay_adi = AY_ESLEME.get(ay_ham)
        if ay_adi is None:
            continue
        linkler: dict[str, str] = {}
        for link_m in re.finditer(r'<a href="([^"]+)"[^>]*>([^<]+)</a>', ay_m.group(2)):
            etiket = link_m.group(2).strip()
            linkler[etiket] = _eib_li_duzelt(link_m.group(1))
        sonuc[ay_adi] = linkler
    return sonuc


def hedef_ay_linki(yil: int, ay_linkleri: dict[str, dict[str, str]], bugun: date) -> tuple[str, str] | None:
    """O yılın "EİB Aylık" (Ocak'ta yalnızca "EİB") bağlantısını seçer.

    Geçmiş yıllar için ARALIK (yılın tamamı o dosyada kümülatif), cari yıl
    için yayımlanan EN SON ay (kaynak 404 vermiyor, mevcut ayların listesi
    zaten JSON içeriğinden biliniyor — TİM'deki gibi ayrıca probe gerekmez).
    """
    if not ay_linkleri:
        return None
    mevcut_aylar = [a for a in AY_ADLARI if a in ay_linkleri]
    if not mevcut_aylar:
        return None
    if yil == bugun.year:
        ay = mevcut_aylar[-1]
    else:
        ay = "ARALIK" if "ARALIK" in ay_linkleri else mevcut_aylar[-1]
    linkler = ay_linkleri[ay]
    url = linkler.get("EİB Aylık") or linkler.get("EİB")
    if url is None:
        raise RuntimeError(f"EİB: {yil} {ay} için 'EİB Aylık' bağlantısı bulunamadı ({sorted(linkler)})")
    return ay, url


def cekilecek_yillar(bugun: date) -> list[int]:
    return list(range(ILK_YIL, bugun.year + 1))


def _ham_veri_sayfasini_bul(kitap):
    """Ham GTİP satır verisini taşıyan sayfayı bulur. Sayfa ADI yıldan yıla
    değişiyor ('GTIP ULKE GB TARIH', 'Sayfa1', ...); başlık satırının ilk
    hücresi 'AYLAR' olan sayfa aranır (büyük/küçük harf toleranslı)."""
    for ad in kitap.sheetnames:
        sayfa = kitap[ad]
        ilk_satir = next(sayfa.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if ilk_satir and ilk_satir[0] and str(ilk_satir[0]).strip().upper() == "AYLAR":
            return sayfa
    raise RuntimeError(f"EİB: ham veri sayfası bulunamadı (sayfalar: {kitap.sheetnames})")


def _sutun_haritasi(baslik_satiri) -> dict[str, int]:
    return {str(h).strip().lower(): i for i, h in enumerate(baslik_satiri) if h}


def _dogrula(aile: dict, su_alt: dict, yil: int) -> None:
    """Kısmi okuma ya da şablon değişimi sessizce geçmemeli."""
    gorulen_urunler = {u for (_, u) in aile.keys()}
    if SU_URUNLERI_GRUBU not in gorulen_urunler:
        raise RuntimeError(f"EİB {yil}: '{SU_URUNLERI_GRUBU}' grubu bulunamadı — şablon değişmiş olabilir")
    if not any(g in gorulen_urunler for g in HAYVANSAL_GRUPLARI):
        raise RuntimeError(f"EİB {yil}: hayvansal ürün gruplarından hiçbiri bulunamadı — şablon değişmiş olabilir")
    gorulen_altlar = {a for (_, a) in su_alt.keys()}
    if not any(ad in gorulen_altlar for ad in SU_URUNLERI_ADLARI):
        raise RuntimeError(f"EİB {yil}: su ürünleri alt kırılımından (Levrek/Çipura/...) hiçbiri bulunamadı")


def yil_verisini_cikar(baytlar: bytes, yil: int) -> dict[tuple[str, str], dict[str, float]]:
    """Bir yılın ham GTİP satırlarından aile/ürün bazlı aylık noktaları çıkarır.

    Dönüş: `{(kalem, ölçüt): {"YYYY-MM-01": değer}}`. `ölçüt` "fobusd" (USD)
    ya da "agirlik" (TON — kaynak KG verir, burada /1000 ile çevrilir).
    """
    kitap = openpyxl.load_workbook(BytesIO(baytlar), data_only=True, read_only=True)
    try:
        sayfa = _ham_veri_sayfasini_bul(kitap)
        satirlar = sayfa.iter_rows(values_only=True)
        baslik = next(satirlar)
        sutun = _sutun_haritasi(baslik)
        ay_i = sutun["aylar"]
        urun_i = sutun["ürün grubu"]
        alt1_i = next(v for k, v in sutun.items() if k.strip() == "alt grup 1")
        agirlik_i = sutun["agirlik"]
        fobusd_i = sutun["fobusd"]

        aile: dict[tuple[str, str], list[float]] = defaultdict(lambda: [0.0, 0.0])
        su_alt: dict[tuple[str, str], list[float]] = defaultdict(lambda: [0.0, 0.0])
        for row in satirlar:
            ay = row[ay_i]
            if ay not in AY_NO:
                continue
            urun = row[urun_i]
            fob = row[fobusd_i] or 0.0
            agr = row[agirlik_i] or 0.0
            aile[(ay, urun)][0] += fob
            aile[(ay, urun)][1] += agr
            if urun == SU_URUNLERI_GRUBU:
                alt1 = row[alt1_i]
                su_alt[(ay, alt1)][0] += fob
                su_alt[(ay, alt1)][1] += agr
    finally:
        kitap.close()

    _dogrula(aile, su_alt, yil)

    noktalar: dict[tuple[str, str], dict[str, float]] = {}

    def tarih_yaz(ay_adi: str) -> str:
        return f"{yil:04d}-{AY_NO[ay_adi]:02d}-01"

    def yaz(kalem: str, ay_adi: str, fob: float, agr: float) -> None:
        noktalar.setdefault((kalem, "fobusd"), {})[tarih_yaz(ay_adi)] = fob
        noktalar.setdefault((kalem, "agirlik"), {})[tarih_yaz(ay_adi)] = agr / 1000.0

    for (ay_adi, urun), (fob, agr) in aile.items():
        if urun == SU_URUNLERI_GRUBU or urun in HAYVANSAL_GRUPLARI:
            yaz(urun, ay_adi, fob, agr)

    for ay_adi in AY_ADLARI:
        varsa = [aile.get((ay_adi, g)) for g in HAYVANSAL_GRUPLARI]
        if not any(v is not None for v in varsa):
            continue
        yaz(
            HAYVANSAL_TOPLAM, ay_adi,
            sum(v[0] for v in varsa if v is not None),
            sum(v[1] for v in varsa if v is not None),
        )

    for ay_adi in AY_ADLARI:
        toplam = aile.get((ay_adi, SU_URUNLERI_GRUBU))
        if toplam is None:
            continue
        adlandirilan_fob = adlandirilan_agr = 0.0
        for ad in SU_URUNLERI_ADLARI:
            deger = su_alt.get((ay_adi, ad))
            if deger is not None:
                yaz(ad, ay_adi, deger[0], deger[1])
                adlandirilan_fob += deger[0]
                adlandirilan_agr += deger[1]
        yaz(DIGER_SU_URUNLERI, ay_adi, toplam[0] - adlandirilan_fob, toplam[1] - adlandirilan_agr)

    return noktalar


def _icerik_cek(session=None) -> str:
    http = session or requests
    yanit = http.get(ICERIK_URL, params={"SI_Id": ISTATISTIK_SI_ID}, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"EİB istatistik sayfası HTTP {yanit.status_code}")
    veri = yanit.json()
    if "icerik" not in veri:
        raise RuntimeError("EİB istatistik sayfası: beklenmeyen JSON gövdesi")
    return veri["icerik"]


def _dosya_indir(url: str, session=None) -> bytes:
    http = session or requests
    yanit = http.get(url, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"EİB yıllık bülten HTTP {yanit.status_code} ({url})")
    return yanit.content


def _tum_noktalari_getir(onbellek: dict, session=None, bugun: date | None = None) -> dict[tuple[str, str], dict[str, float]]:
    if "noktalar" in onbellek:
        return onbellek["noktalar"]
    bugun = bugun or date.today()

    yillar_html = yillara_ayir(_icerik_cek(session))

    noktalar: dict[tuple[str, str], dict[str, float]] = {}
    en_az_bir_yil = False
    for yil in cekilecek_yillar(bugun):
        yil_html = yillar_html.get(str(yil))
        if not yil_html:
            continue
        secim = hedef_ay_linki(yil, ay_baglantilarini_cikar(yil_html), bugun)
        if secim is None:
            continue
        _, url = secim
        yillik = yil_verisini_cikar(_dosya_indir(url, session), yil)
        for anahtar, seri_noktalari in yillik.items():
            noktalar.setdefault(anahtar, {}).update(seri_noktalari)
        en_az_bir_yil = True

    if not en_az_bir_yil:
        raise RuntimeError(f"EİB: {ILK_YIL} sonrası indirilecek yıllık bülten bulunamadı")

    onbellek["noktalar"] = noktalar
    return noktalar


def seri_cek(seri, onbellek: dict | None = None, session=None, bugun: date | None = None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` koşu boyunca paylaşılırsa 32 seri (14 su ürünleri + 18
    hayvansal) aynı 1-2 yıllık yıllık bülten indirmesini paylaşır.
    """
    onbellek = {} if onbellek is None else onbellek
    noktalar = _tum_noktalari_getir(onbellek, session, bugun)

    anahtar = (seri.eib_kalem, seri.eib_olcut)
    if anahtar not in noktalar or not noktalar[anahtar]:
        raise RuntimeError(f"EİB: kalem/ölçüt kombinasyonu bulunamadı: {anahtar!r} ({seri.id})")

    df = pd.DataFrame(sorted(noktalar[anahtar].items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
