"""ifo Institute (Münih) "ifo Business Climate Germany" Excel istemcisi.

Kaynak: https://www.ifo.de/en/ifo-time-series her ay güncellenen
`gsk-e-YYYYMM.xlsx` dosyasına bağlantı verir. O sayfanın kendisi Fastly/
Signal Sciences WAF JS-challenge'ı arkasında (sunucu tarafı `requests` ile
HTTP 200 ama yalnızca ~3KB'lık bir JS-yükleyici kabuk döner, gerçek HTML
istemci tarafında JS çalıştırılınca oluşur — ölçüldü 2026-09-18); bu
kabuktan bağlantı ayıklamak mümkün değil, headless tarayıcı bağımlılığı da
bu projede yok (`requirements-ingest.txt`).

Buna karşılık statik dosya sunucusu (`/sites/default/files/secure/
timeseries/`) WAF'ın arkasında DEĞİL — dosya adı öngörülebilir olduğu için
(`gsk-e-YYYYMM.xlsx`, ayın ifo yayın tarihinden ~3-4 hafta sonra çıkar) en
güncel aydan geriye doğru en fazla `MAKS_GERI_AY` ay denenir; var olmayan
ay da HTTP 200 ile aynı ~3KB kabuğu döndürdüğü için (soft-404) doğrulama
durum koduyla değil dosyanın ZIP/XLSX sihirli baytlarıyla (`PK\x03\x04`)
yapılır.

"ifo Business Climate" sayfasında "Germany" bloğu iki sütun grubu taşır:
"Index, 2015=100, seasonally adjusted" (Business Climate / Situation /
Expectations) ve ardından "Balances, seasonally adjusted" — aynı üç etiket
İKİ kez geçer. Kataloğun istediği endeks (index) değerleri her zaman İLK
geçen (soldaki) sütun grubunda olduğu için etiket araması ilk eşleşmeyi
alır (ifo'nun dokümante ettiği sabit sütun sırası).
"""

from __future__ import annotations

import io
import re
from datetime import date

import openpyxl
import pandas as pd
import requests

from core.catalog import GECERLI_IFO_SERILERI, Seri

TABAN = "https://www.ifo.de"
DOSYA_SABLONU = TABAN + "/sites/default/files/secure/timeseries/gsk-e-{yil:04d}{ay:02d}.xlsx"
MAKS_GERI_AY = 6
ZAMAN_ASIMI = 60
_BASLIKLAR = {"User-Agent": "Mozilla/5.0"}
_XLSX_SIHIRLI_BAYT = b"PK\x03\x04"

_AY_ETIKETI = re.compile(r"^\s*(\d{2})/(\d{4})\s*$")

# ifo_seri -> sütun başlığında aranan (strip edilmiş) etiket
SERI_ETIKETLERI = {
    "iklim": "Business Climate",
    "durum": "Business Situation",
    "beklenti": "Business Expectations",
}
assert set(SERI_ETIKETLERI.keys()) == GECERLI_IFO_SERILERI


def _ay_geri(yil: int, ay: int, kac: int) -> tuple[int, int]:
    toplam = (yil * 12 + (ay - 1)) - kac
    return toplam // 12, toplam % 12 + 1


def _en_guncel_kitabi_indir(session=None, bugun: date | None = None) -> openpyxl.Workbook:
    """En güncel aydan geriye doğru dener; ilk gerçek XLSX yanıtını döndürür."""
    http = session or requests
    bugun = bugun or date.today()
    denenenler = []
    for i in range(MAKS_GERI_AY):
        yil, ay = _ay_geri(bugun.year, bugun.month, i)
        url = DOSYA_SABLONU.format(yil=yil, ay=ay)
        yanit = http.get(url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
        denenenler.append(f"{url} (HTTP {yanit.status_code}, {len(yanit.content)}B)")
        if yanit.status_code == 200 and yanit.content[:4] == _XLSX_SIHIRLI_BAYT:
            return openpyxl.load_workbook(io.BytesIO(yanit.content), data_only=True)
    raise RuntimeError(
        f"ifo Business Climate XLSX'i son {MAKS_GERI_AY} ay için bulunamadı: "
        + "; ".join(denenenler)
    )


def _kitabi_getir(onbellek: dict, session=None, bugun: date | None = None):
    if "kitap" in onbellek:
        return onbellek["kitap"]
    onbellek["kitap"] = _en_guncel_kitabi_indir(session=session, bugun=bugun)
    return onbellek["kitap"]


def _sutun_indeksi(ws, etiket: str) -> int:
    """İlk 10 satırı tarayıp `etiket`i taşıyan İLK (soldaki) sütunu bulur.

    ifo, aynı üç etiketi (Business Climate/Situation/Expectations) önce
    Endeks grubunda sonra Denge (balance) grubunda tekrarlar; soldaki ilk
    eşleşme her zaman Endeks grubudur (bkz. modül docstring'i).
    """
    for satir in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        for i, hucre in enumerate(satir):
            if isinstance(hucre, str) and hucre.strip() == etiket:
                return i
    raise RuntimeError(f"ifo '{ws.title}' sayfasında '{etiket}' sütunu bulunamadı")


def seri_cek(
    seri: Seri, *, onbellek: dict | None = None, session=None, bugun: date | None = None,
) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse tek Excel dosyası (iklim/durum/beklenti üçü de aynı
    dosyayı paylaşır) koşu boyunca bir kez indirilir/ayrıştırılır.
    """
    onbellek = {} if onbellek is None else onbellek
    kitap = _kitabi_getir(onbellek, session=session, bugun=bugun)
    ifo_seri = seri.ifo_seri
    if ifo_seri not in SERI_ETIKETLERI:
        raise RuntimeError(f"Bilinmeyen ifo_seri: {ifo_seri!r} (seri={seri.id})")
    etiket = SERI_ETIKETLERI[ifo_seri]
    ws = kitap["ifo Business Climate"]
    sutun = _sutun_indeksi(ws, etiket)

    noktalar: list[tuple[str, float]] = []
    for satir in ws.iter_rows(values_only=True):
        ilk = satir[0]
        eslesme = _AY_ETIKETI.match(ilk) if isinstance(ilk, str) else None
        if not eslesme:
            continue
        deger = satir[sutun] if sutun < len(satir) else None
        if deger is None or not isinstance(deger, (int, float)):
            continue
        ay, yil = eslesme.groups()
        noktalar.append((f"{yil}-{ay}-01", float(deger)))

    if not noktalar:
        raise RuntimeError(f"ifo '{ifo_seri}' için veri noktası bulunamadı (seri={seri.id})")

    df = pd.DataFrame(sorted(noktalar), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
