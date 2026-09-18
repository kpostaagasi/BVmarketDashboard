"""ODMD (Otomotiv Distribütörleri ve Mobilite Derneği) aylık marka bazında
perakende satış istemcisi.

Kaynak: `odmd.org.tr` "Pazar - Perakende Satışlar" raporu — her ay bir XLSX
(ASP.NET WebForms postback indirmesi) yayımlanır: marka × (Otomobil/Hafif
Ticari/Toplam) × (Yerli/İthal/Toplam) perakende satış adedi. İndirme OTURUM
gerektirir: önce liste sayfası (`neuralnetwork.aspx?type=36`) GET edilip
`ASP.NET_SessionId` çerezi alınmalı, sonra indirme ucu (`wf_docudownload.aspx`)
AYNI `requests.Session` ile çağrılmalı — aksi halde sunucu boş bir "oturum
yok" HTML'i (200 ama içeriksiz) döner, XLSX değil. `session=None` ile çağrılan
(bare `requests` modülüne düşen) kullanım bu yüzden çalışmaz; `ingest.run`
zaten koşu boyunca tek bir `requests.Session` paylaştırır.

Ölçüldü (2026-09-18, canlı): Ağustos 2026 dosyasında FIAT Otomobil/Hafif
Ticari/Toplam = 2.125 / 4.054 / 6.179, VOLKSWAGEN = 5.624 / 1.796 / 7.420 —
referans platformla birebir.

Şirket bazlı toplam kartlar (ör. "DOAS VW Group Perakende Toplam") birden
çok markanın toplamıdır; `Seri.odmd_marka` bir demettir. Doğuş Otomotiv'in
Škoda dağıtımındaki payı ODMD'nin tam rakamının YARISIdır (şirketin kendi
sunumundaki "(½)" notuyla tutarlı, bkz. `doas_vwgroup_sales.html` referans
kartı "DOAS - SKODA Perakende - Toplam (½)"); bu markalar `Seri.
odmd_yarim_marka`ya ayrı yazılır ve 0,5 ağırlıkla toplanır.
"""

from __future__ import annotations

import io
import re
from datetime import date

import openpyxl
import pandas as pd
import requests

TABAN = "https://www.odmd.org.tr/web_2837_1"
LISTE_URL = f"{TABAN}/neuralnetwork.aspx?type=36"
INDIRME_URL = f"{TABAN}/wf_docudownload.aspx"
ZAMAN_ASIMI = 30
VARSAYILAN_GECMIS_YIL = 5
# ~13 satır/sayfa (aylık + yıllık karışık); 8 sayfa VARSAYILAN_GECMIS_YIL=5
# için rahat paydır (yaklaşık 60 ay + yıllık satırlar).
_MAKS_SAYFA = 8

AY_ADLARI = (
    "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
    "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
)
_AY_INDEKS = {ad: i + 1 for i, ad in enumerate(AY_ADLARI)}

# Rapor listesindeki satır başlığı: "2026 Ağustos Perakende Satışlar (Yerli &
# İthal)". Yıllık kümülatif satırlar ("2026 Yılı (Ocak-Ağustos) ...") KASITLI
# olarak eşleşmez — tek ay yerine yıl başından kümülatif taşıdıkları için
# aylık seri üretiminde kullanılamazlar (satır adı ay içermiyor).
_SATIR_DESENI = re.compile(r"(\d{4}) (" + "|".join(AY_ADLARI) + r") Perakende Satışlar")
_RAPOR_ID_DESENI = re.compile(r"sRep\((\d+)\)")


def _satirlari_ayikla(html: str) -> list[tuple[str, str]]:
    """Bir liste sayfasından `(primary_id, "YYYY.MM")` çiftlerini çıkarır."""
    sonuc = []
    for satir_html in html.split('<tr class="tr_r_rep">')[1:]:
        id_eslesme = _RAPOR_ID_DESENI.search(satir_html)
        ay_eslesme = _SATIR_DESENI.search(satir_html)
        if id_eslesme and ay_eslesme:
            yil, ay_adi = ay_eslesme.groups()
            sonuc.append((id_eslesme.group(1), f"{yil}.{_AY_INDEKS[ay_adi]:02d}"))
    return sonuc


def rapor_id_eslemesi(sayfalar_html: list[str]) -> dict[str, str]:
    """Birden çok liste sayfası HTML'inden `"YYYY.MM" -> primary_id` eşlemesi.

    Aynı ay birden çok sayfada görünmez (sayfalar kronolojik ayrık); ilk
    görülen kazanır.
    """
    eslesme: dict[str, str] = {}
    for html in sayfalar_html:
        for rapor_id, anahtar in _satirlari_ayikla(html):
            eslesme.setdefault(anahtar, rapor_id)
    if not eslesme:
        raise RuntimeError(
            f"ODMD rapor listesinde satır bulunamadı ({LISTE_URL}) — "
            "sayfa yapısı değişmiş olabilir"
        )
    return eslesme


def _indeks_cek(session=None) -> dict[str, str]:
    """Liste sayfasını ve sayfalamasını (`_MAKS_SAYFA` kadar) çeker."""
    http = session or requests
    sayfalar: list[str] = []
    for n in range(1, _MAKS_SAYFA + 1):
        if n == 1:
            yanit = http.get(LISTE_URL, timeout=ZAMAN_ASIMI)
        else:
            yanit = http.get(
                f"{TABAN}/sortial.aspx",
                params={"linkpos": n, "target": "categorial1", "type": "36", "detail": "single"},
                timeout=ZAMAN_ASIMI,
            )
        if yanit.status_code != 200:
            if n == 1:
                raise RuntimeError(f"ODMD rapor listesi HTTP {yanit.status_code}")
            break
        sayfalar.append(yanit.text)
    return rapor_id_eslemesi(sayfalar)


def _dosya_indir(rapor_id: str, session=None) -> bytes:
    http = session or requests
    yanit = http.get(
        INDIRME_URL,
        params={
            "primary_id": rapor_id, "type": "36", "target": "categorial1",
            "detail": "single", "downloadfirst": "yes",
        },
        headers={"Referer": LISTE_URL},
        timeout=ZAMAN_ASIMI,
    )
    icerik_tipi = yanit.headers.get("content-type") or ""
    if yanit.status_code != 200 or "spreadsheet" not in icerik_tipi:
        raise RuntimeError(
            f"ODMD rapor indirme başarısız (rapor_id={rapor_id}, "
            f"HTTP {yanit.status_code}, content-type={icerik_tipi!r}) — "
            "oturum çerezi eksik olabilir (liste sayfası önce GET edilmeli)"
        )
    return yanit.content


_MARKA_SUTUNLARI = {"otomobil": 3, "hafif_ticari": 6, "toplam": 9}


def marka_satirlari(dosya_baytlari: bytes) -> dict[str, dict[str, float]]:
    """XLSX baytlarından `{marka: {"otomobil":, "hafif_ticari":, "toplam":}}` çıkarır.

    Sayfa şablonu sabit: `MARKA` satırı OTOMOBİL/HAFİF TİCARİ/TOPLAM üç ana
    sütun başlığı taşır, her biri altında YERLİ/İTHAL/TOPLAM alt sütunu
    vardır — yalnızca TOPLAM alt sütunları okunur (0-tabanlı sütun 3/6/9).
    Öz-doğrulama: dosyanın kendi "TOPLAM:" satırı, marka satırlarının
    toplamıyla karşılaştırılır; tutmazsa RuntimeError (şablon değişmiş
    olabilir — sessizce eksik/yanlış veri üretmektense ingest kırılmalı).
    Ayrıca dönen sözlükte gerçek marka anahtarlarının yanında `"TOPLAM"`
    sözde-anahtarı da vardır — dosyanın kendi TOPLAM satırı, toplam pazar
    kartları için (bkz. `Seri.odmd_marka = ("TOPLAM",)`).
    """
    kitap = openpyxl.load_workbook(io.BytesIO(dosya_baytlari), data_only=True)
    satirlar = list(kitap[kitap.sheetnames[0]].iter_rows(values_only=True))

    # Başlık hücresinin metni ("MARKA") eski dosyalarda (ör. 2021 öncesi)
    # boş — yapısal olarak anlamlı olan üç sütun etiketidir (index 1/4/7),
    # "MARKA" yalnızca kozmetik bir etiket; varlığı şart koşulmaz.
    baslik_satiri = next(
        (
            s for s in satirlar
            if s and len(s) > 7 and (s[1], s[4], s[7]) == ("OTOMOBİL", "HAFİF TİCARİ", "TOPLAM")
        ),
        None,
    )
    if baslik_satiri is None:
        raise RuntimeError(
            "ODMD perakende dosyasında beklenen OTOMOBİL/HAFİF TİCARİ/TOPLAM "
            "sütun başlığı bulunamadı — şablon değişmiş olabilir"
        )

    markalar: dict[str, dict[str, float]] = {}
    toplam_satiri = None
    baslik_gecildi = False
    for satir in satirlar:
        if satir is baslik_satiri:
            baslik_gecildi = True
            continue
        if not baslik_gecildi or not satir or not satir[0]:
            continue
        ad = str(satir[0]).strip()
        if ad.upper().startswith("TOPLAM"):
            toplam_satiri = satir
            break
        markalar[ad] = {ad_: float(satir[kol] or 0) for ad_, kol in _MARKA_SUTUNLARI.items()}

    if toplam_satiri is None:
        raise RuntimeError("ODMD perakende dosyasında TOPLAM satırı bulunamadı")

    for ad_, kol in _MARKA_SUTUNLARI.items():
        beklenen = float(toplam_satiri[kol] or 0)
        bulunan = sum(m[ad_] for m in markalar.values())
        if abs(bulunan - beklenen) > 0.5:
            raise RuntimeError(
                f"ODMD öz-doğrulama — {ad_}: marka toplamı {bulunan:,.0f}, "
                f"dosyanın TOPLAM satırı {beklenen:,.0f}"
            )

    # Toplam pazar kartları (tüm markaların toplamı) için dosyanın kendi
    # "TOPLAM:" satırı ayrı bir sözde-marka anahtarıyla eklenir — yukarıdaki
    # öz-doğrulamadan SONRA eklenir, aksi halde marka toplamına kendi kendini
    # katıp doğrulamayı bozardı. `Seri.odmd_marka = ("TOPLAM",)` ile okunur.
    markalar["TOPLAM"] = {ad_: float(toplam_satiri[kol] or 0) for ad_, kol in _MARKA_SUTUNLARI.items()}

    return markalar


def _deger(seri, markalar: dict[str, dict[str, float]]) -> float:
    kategori = seri.odmd_kategori
    toplam = sum(markalar.get(m, {}).get(kategori, 0.0) for m in seri.odmd_marka)
    yarim = seri.odmd_yarim_marka or ()
    if yarim:
        toplam += 0.5 * sum(markalar.get(m, {}).get(kategori, 0.0) for m in yarim)
    return toplam


def seri_cek(seri, onbellek: dict | None = None, session=None,
             bugun: date | None = None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse rapor indeksi ve ayrıştırılmış marka satırları koşu
    boyunca paylaşılır — TOASO/DOAS'ın onlarca serisi aynı ~60 aylık XLSX
    kümesini okur, önbelleksiz her seri kendi ayını yeniden indirirdi.

    Bir marka BAZI aylarda dosyada hiç yoksa (ör. CUPRA Türkiye'ye 2021'den
    sonra girdi) bu gerçek iş durumudur — o ay için katkısı 0 sayılır,
    hata değildir. Marka İŞLENEN AYLARIN HİÇBİRİNDE bulunamazsa (yazım
    hatası koruması) RuntimeError.
    """
    bugun = bugun or date.today()
    onbellek = {} if onbellek is None else onbellek

    if "indeks" not in onbellek:
        onbellek["indeks"] = _indeks_cek(session)
    eslesme = onbellek["indeks"]

    en_eski_yil = bugun.year - VARSAYILAN_GECMIS_YIL
    gereken_markalar = (*seri.odmd_marka, *(seri.odmd_yarim_marka or ()))
    hic_bulunmayan = set(gereken_markalar)

    kendi: dict[str, float] = {}
    for anahtar, rapor_id in eslesme.items():
        yil = int(anahtar[:4])
        if yil < en_eski_yil:
            continue
        if anahtar not in onbellek:
            baytlar = _dosya_indir(rapor_id, session)
            onbellek[anahtar] = marka_satirlari(baytlar)
        markalar = onbellek[anahtar]
        hic_bulunmayan -= markalar.keys()

        ay = int(anahtar[5:7])
        kendi[f"{yil}-{ay:02d}-01"] = _deger(seri, markalar)

    if hic_bulunmayan:
        raise RuntimeError(
            f"ODMD'de hiç bulunamayan marka: {sorted(hic_bulunmayan)} ({seri.id}) — "
            "katalogdaki marka adı dosyayla eşleşmiyor olabilir"
        )
    if not kendi:
        raise RuntimeError(
            f"ODMD'de {seri.odmd_marka} için hiç veri bulunamadı ({seri.id})"
        )


    df = pd.DataFrame(sorted(kendi.items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
