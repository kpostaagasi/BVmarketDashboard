"""GPH (Global Ports Holding) Aylık Trafik İstatistikleri istemcisi.

Global Yatırım Holding (GLYHO) — GPH'nin dolaylı ana ortağı — yolcu/sefer
istatistiklerini KAP'ta SERBEST METİN olarak duyurur (ölçüldü 2026-09-19:
KAP `oda_ExplanationTextBlock`, "Bildirim Ekleri (0)" — ek YOK), ama aynı
verinin YAPILANDIRILMIŞ XLSX dosyasını kendi Yatırımcı İlişkileri sitesinde
yayımlar: `https://globalyatirim.com.tr/tr/raporlar/yolcu-istatistikleri/`
sayfasındaki "İndir" bağlantıları. Bu yüzden KAP değil, bu XLSX kaynak
alınıyor (bkz. görev kısıtı: "KAP veya şirket IR sayfaları").

Dosya adı kalıbı istikrarsız (`GPH_AylikTrafik_<Ay><Yıl>.xlsx`,
`gph-aylik-trafik-<ay>-<yıl>.xlsx`, ara sıra `-1` sonekli tekrar yüklemeler)
— bu yüzden URL TAHMİN EDİLMİYOR, sayfa her koşuda taranıp "İndir"
bağlantıları çözülüyor.

**TEK dosya TÜM tarihçeyi taşır**: en güncel XLSX, o ana kadarki HER AYIN
kendi sekmesini (`Ağustos-2026`, `Temmuz-2026`, …, `Eylül-21`) içerir — SGK
gibi, ay ay indirmeye gerek yok.

Sekme düzeni iki döneme ayrılıyor (ölçüldü, KAP Bildirim 1574870: "Ocak 2026
itibarıyla, konsolide edilmeyen GPH limanlarını da içeren ilave bir
gösterime yer verilmeye başlanmıştır"):

* **2026-01 öncesi**: yalnızca KONSOLİDE toplamlar var; satır etiketi
  "Toplam Sefer Sayısı" / "Toplam Yolcu Sayısı" (parantezsiz).
* **2026-01 ve sonrası**: satır etiketine " (Konsolide)" eklendi VE ayrıca
  "Konsolide Edilmeyen Limanlar" bölümü (Seferler/Yolcu Sayısı alt
  satırlarıyla) ve "... (Tüm GPH Portföyü)" toplam satırları eklendi.

Sütun düzeni HER İKİ dönemde de aynı: başlık satırı ("Kruvaziyer
Limanları" B sütununda) sonrası E sütunu (0-indeksli col4) HER ZAMAN o
sekmenin KENDİ AYININ cari-yıl değeridir (yıl sürüklenmesi yok — "Ocak-2025"
sekmesinin col4'ü her zaman 2025 Ocak'ın kendi rakamıdır, hangi yılda
indirildiği önemsiz). Parser bu yüzden ETİKET eşlemesiyle satırı bulup
col4'ü okur; satır SAYISI (32 ya da 66) döneme göre değişse de konum sabit.

"Toplam Yolcu (Tüm GPH Portföyü)" AYRI SERİ DEĞİLDİR: Konsolide + Konsolide
Edilmeyen toplamının basit bir toplamıdır (ölçüldü: Ağustos 2026'da
1.768.845 + 397.324 = 2.166.169, kaynağın kendi satırıyla birebir) —
`core/stats.py` katmanında hesaplanmalı.

Kalemler `gph_metrik` alanıyla seçilir (bkz.
`core.catalog.GECERLI_GPH_METRIKLERI`).
"""

from __future__ import annotations

import re

import pandas as pd
import requests

from core.catalog import GECERLI_GPH_METRIKLERI

IR_SAYFASI = "https://globalyatirim.com.tr/tr/raporlar/yolcu-istatistikleri/"
ZAMAN_ASIMI = 60

# gph_metrik -> (satır etiketi TESPİTİ, "Konsolide Edilmeyen Limanlar" alt satırı mı).
# Konsolide satırlar `startswith` + "Tüm GPH Portföyü" HARİÇ TUTULARAK bulunur
# (iki dönemde de "Toplam Sefer/Yolcu Sayısı" ile başlar, biri
# " (Konsolide)" sonekli, biri değil — bkz. modül docstring'i).
_KONSOLIDE_ETIKETLERI = {
    "sefer-konsolide": "Toplam Sefer Sayısı",
    "yolcu-konsolide": "Toplam Yolcu Sayısı",
}
_KONSOLIDE_EDILMEYEN_ALT_ETIKETI = {"yolcu-konsolide-edilmeyen": "Yolcu Sayısı"}
_KONSOLIDE_EDILMEYEN_BASLIGI = "Konsolide Edilmeyen Limanlar"
_TUM_PORTFOY_ISTISNASI = "Tüm GPH Portföyü"

assert set(_KONSOLIDE_ETIKETLERI) | set(_KONSOLIDE_EDILMEYEN_ALT_ETIKETI) == GECERLI_GPH_METRIKLERI

# Sekme adından (yıl, ay) çıkarımı: "Ağustos-2026", "Eylul-2025", "Tem-22",
# "Mart-23_Eski Raporlama" (atlanır — birincil sekmeyle çakışır) gibi biçimler.
_AY_KISALTMA = {
    "ocak": 1, "subat": 2, "şubat": 2, "mart": 3, "nisan": 4, "mayis": 5, "mayıs": 5,
    "haziran": 6, "haz": 6, "temmuz": 7, "tem": 7, "agustos": 8, "ağustos": 8,
    "eylul": 9, "eylül": 9, "ekim": 10, "kasim": 11, "kasım": 11, "aralik": 12, "aralık": 12,
    "nis": 4, "may": 5, "subat_kisa": 2,
}
_SEKME_DESENI = re.compile(r"^([A-Za-zÇĞİÖŞÜçğıöşü]+)-(\d{2,4})$")


def sekme_donemi(sekme_adi: str) -> tuple[int, int] | None:
    """`'Ağustos-2026'` -> `(2026, 8)`. Tanımayan/çift sekme (ör. eski
    raporlama varyantı) için None döner — çağıran atlamalı."""
    m = _SEKME_DESENI.match(sekme_adi.strip())
    if not m:
        return None
    ay_adi, yil_ham = m.groups()
    ay = _AY_KISALTMA.get(ay_adi.casefold())
    if ay is None:
        return None
    yil = int(yil_ham)
    if yil < 100:
        yil += 2000
    return yil, ay


_XLSX_BAGLANTISI = re.compile(r'href="\s*([^"]+\.xlsx)"', re.I)


def _guncel_dosya_url(session: requests.Session) -> str:
    """IR sayfasını tarar, EN GÜNCEL "İndir" bağlantısını (ilk XLSX) döner.

    Sayfa sunucu tarafında render ediliyor (ölçüldü 2026-09-19: düz
    `requests.get` ile bağlantılar HTML'de mevcut) — tarayıcıya gerek yok.
    Dosya adı kalıbı istikrarsız (bkz. modül docstring'i), bu yüzden liste
    SAYFA SIRASIYLA (en yeni ilk) okunur, ada göre sıralama YAPILMAZ.
    """
    yanit = session.get(IR_SAYFASI, timeout=ZAMAN_ASIMI, headers={"User-Agent": "Mozilla/5.0"})
    yanit.raise_for_status()
    eslesmeler = _XLSX_BAGLANTISI.findall(yanit.text)
    if not eslesmeler:
        raise RuntimeError("GPH: IR sayfasında hiç .xlsx bağlantısı bulunamadı — şablon değişmiş olabilir")
    return eslesmeler[0].strip()


def _kitabi_getir(onbellek: dict, session: requests.Session | None = None):
    if "kitap" in onbellek:
        return onbellek["kitap"]
    import io

    import openpyxl

    http = session or requests.Session()
    dosya_url = onbellek.get("dosya_url") or _guncel_dosya_url(http)
    yanit = http.get(dosya_url, timeout=ZAMAN_ASIMI)
    yanit.raise_for_status()
    kitap = openpyxl.load_workbook(io.BytesIO(yanit.content), data_only=True, read_only=True)
    onbellek["kitap"] = kitap
    return kitap


def _sekme_degerini_oku(ws, metrik: str) -> float | None:
    satirlar = list(ws.iter_rows(values_only=True))
    if metrik in _KONSOLIDE_ETIKETLERI:
        etiket = _KONSOLIDE_ETIKETLERI[metrik]
        for satir in satirlar:
            deger_etiket = satir[1]
            if deger_etiket is None or _TUM_PORTFOY_ISTISNASI in str(deger_etiket):
                continue
            if str(deger_etiket).strip().startswith(etiket):
                return satir[4]
        return None
    # "Konsolide Edilmeyen Limanlar" alt satırı — 2026-01 öncesi hiç yok.
    alt_etiket = _KONSOLIDE_EDILMEYEN_ALT_ETIKETI[metrik]
    baslik_bulundu = False
    for satir in satirlar:
        if satir[1] and str(satir[1]).strip() == _KONSOLIDE_EDILMEYEN_BASLIGI:
            baslik_bulundu = True
            continue
        if baslik_bulundu and satir[2] and str(satir[2]).strip() == alt_etiket:
            return satir[4]
        if baslik_bulundu and satir[1]:  # bir sonraki bölüm başlığına geçildi
            baslik_bulundu = False
    return None


def tum_noktalari_ayikla(kitap, metrik: str) -> list[tuple[str, float]]:
    noktalar = []
    for sekme_adi in kitap.sheetnames:
        donem = sekme_donemi(sekme_adi)
        if donem is None:
            continue
        yil, ay = donem
        deger = _sekme_degerini_oku(kitap[sekme_adi], metrik)
        if deger is None:
            continue  # metrik bu dönemde henüz yayımlanmıyor (gerçek eksiklik)
        noktalar.append((f"{yil:04d}-{ay:02d}-01", float(deger)))
    if not noktalar:
        raise RuntimeError(f"GPH: {metrik!r} için hiçbir sekmede veri bulunamadı — şablon değişmiş olabilir")
    return sorted(noktalar)


def seri_cek(seri, *, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı)."""
    onbellek = onbellek if onbellek is not None else {}
    kitap = _kitabi_getir(onbellek, session)
    metrik = seri.gph_metrik
    if metrik not in GECERLI_GPH_METRIKLERI:
        raise RuntimeError(f"GPH: bilinmeyen gph_metrik={metrik!r}")
    cache_key = f"noktalar:{metrik}"
    if cache_key not in onbellek:
        onbellek[cache_key] = tum_noktalari_ayikla(kitap, metrik)
    return pd.DataFrame(onbellek[cache_key], columns=["date", "value"])
