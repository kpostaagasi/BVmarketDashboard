"""DHMİ (Devlet Hava Meydanları İşletmesi) büyük havalimanları trafik istemcisi.

Kaynak: https://www.dhmi.gov.tr/Sayfalar/Istatistikler.aspx, "Havalimanları
Karşılaştırmalı İstatistikleri" sekmesi. Her ay bir SharePoint liste öğesi
(ör. `Attachments/442/`) altında beş PDF + tek bir "TÜMÜ.xlsx" eki
yayımlanır. Bu XLSX AYLIK DEĞİL — YIL BAŞINDAN İLGİLİ AYA KADAR
KÜMÜLATİFTİR ve aynı anda İKİ YILI karşılaştırır: sayfanın varsayılan
(postback'siz) yüklemesi CARİ YILIN tüm yayımlanmış ayları + BİR ÖNCEKİ
YILIN AYNI ayları için birer sütun taşır (ör. Ağustos bülteni sütun
başlıkları: "2025 AĞUSTOS SONU" | "2026 YILI AĞUSTOS SONU (Kesin
Olmayan)"). Ölçüldü (2026-09-18): yıl seçici (`ddlyil`, ASP.NET/SharePoint
WebPart postback'i + `__EVENTVALIDATION`) düz `requests` ile
tetiklenemiyor (sunucu "Hata" sayfası döner) — bu yüzden yalnızca
varsayılan (CARİ YIL) sayfa kullanılır; bu TEK istek CARİ YIL (Ocak..son
yayımlanan ay) ve BİR ÖNCEKİ YILIN aynı ayları olmak üzere ~1.5–2 yıllık
gerçek pencere verir.

AYLIK akış = ardışık iki bültenin AYNI YIL sütunundaki kümülatif değerin
farkı (Ocak ayı kümülatifi = Ocak'ın kendisi, fark alınmaz).

marketvisuals.net'in "Havalimanı Bazlı" kartları DHMİ'nin "TÜRKİYE GENELİ"
(~60 havalimanı) ya da "DHMİ TOPLAMI" (özel işletmeli havalimanları hariç)
satırlarıyla EŞLEŞMİYOR — ölçüldü: ikisi de referans karta göre %20'den
fazla sapıyor. Kartlar yalnızca SEÇİLİ 6 havalimanının (İstanbul, Sabiha
Gökçen, Ankara Esenboğa, İzmir Adnan Menderes, Antalya, Muğla Dalaman)
toplamıdır; bu 6'nın Temmuz 2026 aylık farkları toplamı referans sayfanın
5 KPI'ının (yolcu toplam/dış hat, kargo, uçak toplam/ticari) TAMAMIYLA tam
eşleşiyor (bkz. `tests/test_dhmi.py` ve görev raporu).

TAV ÇAKIŞMASI: Antalya/Ankara Esenboğa/İzmir Adnan Menderes'i TAV da
işletiyor. Ölçüldü: TAV'ın kendi "Antalya"/"Ankara"/"Izmir" toplam yolcu
serisi (`ingest/tav.py`, `havacilik/tav-*-toplam`) bu üç havalimanı için
DHMİ'nin ham satırıyla AYNI DEĞERİ verir (TAV, işlettiği havalimanının
resmi DHMİ trafiğini kendi IR bülteninde yeniden yayımlıyor) — ör. Temmuz
2026: TAV Antalya=5.555.485=DHMİ Antalya, TAV Ankara=1.443.339=DHMİ Ankara
Esenboğa, TAV İzmir=1.412.062=DHMİ Adnan Menderes. Bu adaptör yine de bu
üç havalimanını KENDİ DHMİ dosyasından çeker (TAV'ınkiyle birleştirmez):
(a) mimari bir Seri'nin TEK `kaynak_tipi`'nden beslenmesini gerektiriyor —
iki farklı adaptörü tek seride birleştirmenin şablonu yok; (b) TAV'ın
kapsamadığı İstanbul/Sabiha Gökçen/Muğla Dalaman zaten ayrı DHMİ okuması
gerektiriyor. Üretilen 5 seri (`havacilik/dhmi-*`) TAV'ın STANDALONE
Antalya/Ankara/İzmir serilerini MÜKERRER ETMEZ — TAV'ınkiler tek
havalimanı, buradakiler 6 havalimanının TOPLAMIdır ve farklı kartları
besler.

Sözleşme: seri_cek(seri, *, onbellek=None, session=None) -> DataFrame[date,value]
"""

from __future__ import annotations

import html
import re
from io import BytesIO

import openpyxl
import pandas as pd
import requests

from core.catalog import GECERLI_DHMI_OLCUTLERI

SAYFA = "https://www.dhmi.gov.tr/Sayfalar/Istatistikler.aspx"
ZAMAN_ASIMI = 60

# marketvisuals.net kartlarının kapsadığı 6 havalimanı: kısa ad -> DHMİ ham
# satır adı (yıldızlı/parantezli ekler temizlenmiş). Bkz. modül docstring'i.
_HAVALIMANLARI = {
    "istanbul": "İstanbul",
    "sabiha-gokcen": "İstanbul Sabiha Gökçen",
    "ankara-esenboga": "Ankara Esenboğa",
    "adnan-menderes": "İzmir Adnan Menderes",
    "antalya": "Antalya",
    "mugla-dalaman": "Muğla Dalaman",
}

# Ölçüt -> (okunacak sayfa adı, Toplam/Dış Hat sütun endeksi [0=İç,1=Dış,2=Toplam]).
_OLCUT_SHEET_KOLON: dict[str, tuple[str, int]] = {
    "yolcu-toplam": ("YOLCU", 2),
    "yolcu-dis-hat": ("YOLCU", 1),
    "kargo-toplam": ("KARGO", 2),
    "ucak-toplam": ("TÜM UÇAK", 2),
    "ucak-ticari": ("TİCARİ UÇAK", 2),
}
assert set(_OLCUT_SHEET_KOLON) == GECERLI_DHMI_OLCUTLERI

_AYLAR = (
    "OCAK", "ŞUBAT", "MART", "NİSAN", "MAYIS", "HAZİRAN", "TEMMUZ",
    "AĞUSTOS", "EYLÜL", "EKİM", "KASIM", "ARALIK",
)

_DOSYA_DESENI = re.compile(
    r'href="(https://www\.dhmi\.gov\.tr/Lists/Istatislikler/Attachments/\d+/[^"]*?\.xlsx)"'
)


def dosya_listesi(session=None) -> list[str]:
    """İstatistikler sayfasını kazır, "TÜMÜ.xlsx" eklerinin URL'lerini döner.

    Bulunma sırası kronolojik VARSAYILMAZ (SharePoint liste ID'leri revizyon
    öğeleri için boşluk bırakabiliyor, ölçüldü) — her dosyanın ait olduğu ay
    `_donem_coz` ile dosyanın KENDİ başlığından çözülür.
    """
    http = session or requests
    yanit = http.get(SAYFA, timeout=ZAMAN_ASIMI)
    yanit.raise_for_status()
    metin = html.unescape(yanit.text)
    return [m.group(1) for m in _DOSYA_DESENI.finditer(metin)]


def _dosya_indir(url: str, onbellek: dict, session=None) -> bytes:
    if url not in onbellek:
        http = session or requests
        yanit = http.get(url, timeout=ZAMAN_ASIMI)
        yanit.raise_for_status()
        onbellek[url] = yanit.content
    return onbellek[url]


def _donem_coz(baslik: str) -> tuple[int, str]:
    """'2026 YILI AĞUSTOS SONU\\n(Kesin Olmayan)' -> (2026, 'AĞUSTOS')."""
    yil_e = re.search(r"(\d{4})", baslik)
    ay = next((a for a in _AYLAR if a in baslik), None)
    if not yil_e or not ay:
        raise RuntimeError(f"DHMİ dönem başlığı çözülemedi: {baslik!r}")
    return int(yil_e.group(1)), ay


def _sayfa_noktalari(baytlar: bytes, sheet: str) -> dict[tuple[int, str], dict[str, tuple[float, float, float]]]:
    """Bir TÜMÜ.xlsx'in bir sayfasından HER İKİ yıl için 6 havalimanının
    (İç Hat, Dış Hat, Toplam) kümülatif üçlülerini çıkarır.

    Dönüş: {(yil, ay_adi): {havalimani_kisa_adi: (ic, dis, toplam)}}.
    """
    kitap = openpyxl.load_workbook(BytesIO(baytlar), data_only=True)
    if sheet not in kitap.sheetnames:
        raise RuntimeError(f"DHMİ dosyasında '{sheet}' sayfası yok: {kitap.sheetnames}")
    satirlar = list(kitap[sheet].iter_rows(values_only=True))
    if len(satirlar) < 7:
        raise RuntimeError(f"DHMİ '{sheet}' sayfası beklenenden kısa ({len(satirlar)} satır)")

    alt_baslik = tuple(satirlar[2][1:7])
    beklenen = ("İç Hat", "Dış Hat", "Toplam", "İç Hat", "Dış Hat", "Toplam")
    if alt_baslik != beklenen:
        raise RuntimeError(f"DHMİ '{sheet}' sütun şablonu değişmiş: {alt_baslik}")

    donem_satiri = satirlar[1]
    yil_sol, ay_sol = _donem_coz(str(donem_satiri[1]))
    yil_sag, ay_sag = _donem_coz(str(donem_satiri[4]))
    if ay_sol != ay_sag:
        raise RuntimeError(f"DHMİ '{sheet}': iki yıl sütunu farklı ayı gösteriyor ({ay_sol} / {ay_sag})")

    def satir_bul(ham_ad: str) -> tuple:
        for satir in satirlar[3:]:
            ad = satir[0]
            if isinstance(ad, str) and ad.replace("(*)", "").replace("(**)", "").strip() == ham_ad:
                return satir
        raise RuntimeError(f"DHMİ '{sheet}' sayfasında '{ham_ad}' havalimanı satırı yok")

    sonuc: dict[tuple[int, str], dict[str, tuple[float, float, float]]] = {
        (yil_sol, ay_sol): {}, (yil_sag, ay_sag): {},
    }
    for kisa_ad, ham_ad in _HAVALIMANLARI.items():
        satir = satir_bul(ham_ad)
        sonuc[(yil_sol, ay_sol)][kisa_ad] = (satir[1], satir[2], satir[3])
        sonuc[(yil_sag, ay_sag)][kisa_ad] = (satir[4], satir[5], satir[6])
    kitap.close()
    return sonuc


def seri_cek(seri, *, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı)."""
    onbellek = {} if onbellek is None else onbellek
    sheet, kolon = _OLCUT_SHEET_KOLON[seri.dhmi_olcut]

    if "dosyalar" not in onbellek:
        onbellek["dosyalar"] = dosya_listesi(session=session)
    dosyalar = onbellek["dosyalar"]
    if not dosyalar:
        raise RuntimeError("DHMİ istatistikler sayfasında TÜMÜ.xlsx eki bulunamadı")

    kumulatif: dict[int, dict[str, float]] = {}
    for url in dosyalar:
        baytlar = _dosya_indir(url, onbellek, session=session)
        for (yil, ay), havalimanlari in _sayfa_noktalari(baytlar, sheet).items():
            toplam = sum(deger[kolon] for deger in havalimanlari.values())
            kumulatif.setdefault(yil, {})[ay] = toplam

    satirlar = []
    for yil, aylar in kumulatif.items():
        onceki = 0.0
        for ay in _AYLAR:
            if ay not in aylar:
                continue
            cari = aylar[ay]
            satirlar.append((f"{yil:04d}-{_AYLAR.index(ay) + 1:02d}-01", cari - onceki))
            onceki = cari

    df = pd.DataFrame(satirlar, columns=["date", "value"]).sort_values("date")
    return df.reset_index(drop=True)
