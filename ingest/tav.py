"""TAV Havalimanları Yatırımcı İlişkileri aylık yolcu trafiği istemcisi.

Kaynak gerçekleri 2026-09-18'de canlı ölçüldü (sıfırdan yeniden keşfetmeye
çalışmayın):

- Dosya listesi `LISTE_SAYFASI`'nin düz HTML'inde gömülü: `a[href$=".xlsx"]`
  bağlantıları JS gerektirmeden `requests` ile okunabiliyor. Her bağlantı
  "Traffic Results" kategorisinin bir yılına karşılık gelir, dosya adı
  `DocumentsDDMMYYYYHHMMSS_.xlsx` biçiminde tarih damgalı ve TÜRETİLEMEZ.
- **Bu dosyalar YIL BAZLI DEĞİL, KÜMÜLATİF PENCERE'dir**: en yeni dosya TEK
  BAŞINA Ocak 2020'den bugüne (Sheet1 = son ay) TÜM aylık geçmişi ayrı
  sayfalar (`Sheet1` + `MMYY` sekmeleri) halinde taşıyor — diğer (eski)
  dosyalar bunun katı alt kümesi (daha az ay). Bu yüzden yalnızca EN YENİ
  dosya indirilir; `dosya_listesi` dosya adındaki zaman damgasına göre
  yeniden-eskiye sıralar, `seri_cek` yalnızca `[0]`'ı kullanır.
- Her sayfanın İKİ bloğu var: "Passengers / Yolcu" (yolcu sayısı — bu
  modülün konusu) ve altında "Air Traffic Movements / Ucus Sayisi" (uçuş
  sayısı — kapsam dışı, contract yalnızca yolcu alanları tanımlıyor).
  Bloklar arasında tam boş bir satır var; yolcu bloğu bu satırda durur.
- Her blokta varlık satırını (havalimanı adı ya da "TAV TOTAL") opsiyonel
  olarak "International / Dis Hat" (bazı sayfalarda yalnızca
  "International") + "Domestic / Ic Hat" (bazen yalnızca "Domestic") satır
  çifti izliyor. **Her varlığın kırılımı yok**: Ölçüldü — Antalya, Izmir,
  Ankara, Milas - Bodrum, Gazipasa Alanya, Almaty, Madinah / Medine ve TAV
  TOTAL'ın kırılımı var; Georgia / Gürcistan, Tunisia / Tunus, North
  Macedonia / Kuzey Makedonya, Zagreb'in YOK (tek satır, yalnızca toplam).
  Madinah/Medine'nin kırılımı da TÜM geçmişte yok: Aralık 2022 ve Ocak
  2023 sayfalarında tek satır, Temmuz-Ağustos 2026'da kırılımlı — TAV bu
  havalimanı için ayrımı sonradan yayımlamaya başlamış. Bu yüzden
  dis-hat/ic-hat segmentleri sabit bir varlık kümesine değil, HER SAYFADA
  GERÇEKTEN GÖRÜLEN satırlara göre üretilir.
- "TAV TOTAL" satırının değeri zaten dosyada var (havalimanlarının
  toplamı elle üretilmiyor): Ağustos 2026 Sheet1 satır 30-32 → Toplam
  14.742.288, Dış Hat 10.521.657, İç Hat 4.220.631 —
  marketvisuals.net/tav_traffic.html'in `TAVHL Toplam` kartlarıyla
  birebir eşleşiyor.
- Aylık değer, o sayfanın kendi "Passengers / Yolcu" başlık satırında
  RAPORLANAN YILA (başlıktan çözülür) eşit hücrenin bulunduğu sütundur.
  Bu sütunun konumu sayfadan sayfaya DEĞİŞİR (bazı sayfalar 2 karşılaştırma
  yılı gösterir → C sütunu, bazıları 3 yıl → D sütunu) — bu yüzden sabit
  sütun harfi yerine "başlık satırında raporlanan yıla eşit ilk hücre"
  dinamik olarak aranır. YTD (Ocak-X) bloğu aynı yılı bir kez daha taşır;
  ilk (en soldaki) eşleşme her zaman AYLIK bloktur, YTD değil.
- Sayfa adı (`Sheet1`, `MMYY`) yerine sayfanın KENDİ başlık metni
  (A1 hücresi, "TAV Traffic Figures – <Ay> <Yıl>\\nTAV Havalimanları Yolcu
  Sayıları – <Türkçe Ay> <Yıl>") ay/yıl kaynağı olarak kullanılıyor —
  sekme adı GÜVENİLMEZ: Aralık 2022 verisini taşıyan sekmenin adı "1212"
  (olması gereken "1222"), yani MMYY'yi olduğu gibi ayrıştırmak yanlış
  yıla (2012) düşerdi. Başlık metni de her zaman kusursuz değil: Nisan
  2023 sayfasının İngilizce kısmı "Apil 2023" yazıyor (eksik harf) — bu
  yüzden İngilizce ay adı eşleşmezse Türkçe ada (`Nisan`) düşülüyor.
  Ölçüldü: 80 sayfadan 79'u (Şubat 2020 - Ağustos 2026) bu şekilde
  sorunsuz çözülüyor.
- TEK bilinen istisna: `0120` (Ocak 2020) sayfasının başlığı "January
  2020" olarak sorunsuz çözülüyor, ama "Passengers / Yolcu" başlık
  satırındaki karşılaştırma sütunları 2018/2019 — raporlanan yılın
  (2020) kendi sütunu YOK (muhtemelen önceki bir şablondan kopyalanıp
  güncellenmemiş kalıntı). Bu TEK sayfa "yıl sütunu bulunamadı" olarak
  atlanır (`_sayfa_noktalari` → None); diğer 79 ay etkilenmez. Bu, adı
  ya da tarihi çözülemeyen bir sayfa değil — RuntimeError'a çıkan diğer
  tüm yollar (başlık yok, başlık ay/yıl içermiyor, "Passengers" satırı
  yok, sayısal olmayan hücre, yinelenen varlık/anahtar) hâlâ geçerli.
"""

from __future__ import annotations

import io
import re

import openpyxl
import pandas as pd
import requests

from core.catalog import GECERLI_TAV_SEGMENTLERI

LISTE_SAYFASI = "https://ir.tav.aero/en-EN/financials-and-operationals"
ZAMAN_ASIMI = 60

_BASLIK_ISARETI = "TAV Traffic Figures"
_YOLCU_ONEKI = "Passengers"
_DOSYA_ADI = re.compile(r"Documents(\d{2})(\d{2})(\d{4})(\d{2})(\d{2})(\d{2})_")

AY_EN = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)
AY_TR = (
    "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
    "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
)
# İngilizce ad önce denenir (başlıkta önce geçiyor); bozuksa ("Apil" gibi)
# Türkçe ada düşülür. Bkz. modül docstring'i.
AY_ADLARI = tuple(zip(AY_EN, range(1, 13))) + tuple(zip(AY_TR, range(1, 13)))

# Varlık satırının hemen ardından gelebilecek alt segment etiketleri.
# Bazı sayfalarda (ör. Madinah/Medine, 2026) bilingual son ek düşüyor.
_ALT_ETIKETLER = {
    "International / Dis Hat": "dis-hat",
    "International": "dis-hat",
    "Domestic / Ic Hat": "ic-hat",
    "Domestic": "ic-hat",
}
assert set(_ALT_ETIKETLER.values()) | {"toplam"} == GECERLI_TAV_SEGMENTLERI


def dosya_listesi(session=None) -> list[str]:
    """Liste sayfasını kazır, xlsx bağlantılarını YENİDEN-ESKİYE sıralı döner.

    URL'ler tarih damgalı (`DocumentsDDMMYYYYHHMMSS_.xlsx`) ve
    TÜRETİLEMEZ; her koşuda sayfa yeniden kazınmalı. `[0]` en güncel
    dosyadır ve TEK BAŞINA tüm geçmişi taşır (bkz. modül docstring'i) —
    `seri_cek` yalnızca bunu indirir.
    """
    http = session or requests
    yanit = http.get(
        LISTE_SAYFASI, headers={"User-Agent": "Mozilla/5.0"}, timeout=ZAMAN_ASIMI
    )
    if yanit.status_code != 200:
        raise RuntimeError(
            f"TAV finansal/operasyonel sayfası HTTP {yanit.status_code}"
        )
    baglantilar = re.findall(r'href="([^"]+\.xlsx)"', yanit.text)

    def damga(url: str):
        eslesme = _DOSYA_ADI.search(url)
        if not eslesme:
            return None
        gun, ay, yil, saat, dakika, saniye = (int(x) for x in eslesme.groups())
        return (yil, ay, gun, saat, dakika, saniye)

    damgali = [(damga(u), u) for u in baglantilar]
    damgali = [(d, u) for d, u in damgali if d is not None]
    if not damgali:
        raise RuntimeError(
            "TAV finansal/operasyonel sayfasında tarih damgalı .xlsx bağlantısı bulunamadı"
        )
    damgali.sort(key=lambda ikili: ikili[0], reverse=True)
    return [u for _, u in damgali]


def _sayfa_basligi(satirlar: list[tuple]) -> str | None:
    for satir in satirlar[:8]:
        deger = satir[0] if satir else None
        if isinstance(deger, str) and _BASLIK_ISARETI in deger:
            return deger
    return None


def _sayfa_tarihi(sayfa_adi: str, baslik: str) -> tuple[int, int]:
    """Başlık metninden (yıl, ay) çözer.

    İngilizce ve Türkçe ay adları ayrı ayrı denenir; ikisi de eşleşirse
    (normal durum) aynı çifte düşer, yalnızca biri eşleşirse (ör. "Apil"
    yazım hatası) o kullanılır. Sıfır ya da birden fazla FARKLI çift
    bulunursa (başlık bozuk/belirsiz) RuntimeError.
    """
    bulunan: set[tuple[int, int]] = set()
    for ad, ay_no in AY_ADLARI:
        eslesme = re.search(rf"{re.escape(ad)}\.?\s+(\d{{4}})", baslik)
        if eslesme:
            bulunan.add((int(eslesme.group(1)), ay_no))
    if len(bulunan) != 1:
        raise RuntimeError(
            f"TAV trafik bülteni: '{sayfa_adi}' sayfasının ay/yıl başlığı "
            f"çözülemedi: {baslik!r}"
        )
    return bulunan.pop()


def _yolcu_basligi_satiri(satirlar: list[tuple]) -> int | None:
    for i, satir in enumerate(satirlar):
        deger = satir[0] if satir else None
        if isinstance(deger, str) and deger.strip().startswith(_YOLCU_ONEKI):
            return i
    return None


def _yil_sutunu(baslik_satiri: tuple, yil: int) -> int | None:
    """Başlık satırında raporlanan yıla eşit İLK (en soldaki) sütun.

    YTD bloğu aynı yılı bir kez daha taşıdığı için ilk eşleşme alınır —
    bu her zaman AYLIK bloktur (bkz. modül docstring'i).
    """
    for sutun, deger in enumerate(baslik_satiri):
        if isinstance(deger, int) and not isinstance(deger, bool) and deger == yil:
            return sutun
    return None


def _hucre_oku(satir: tuple, sutun: int, baglam: str) -> float | None:
    """Hücreyi okur; boşsa None (yayımlanmamış/eksik — atlanır, sıfır
    yazılmaz), sayısal değilse RuntimeError."""
    if sutun >= len(satir):
        return None
    hucre = satir[sutun]
    if hucre is None:
        return None
    if isinstance(hucre, bool) or not isinstance(hucre, (int, float)):
        raise RuntimeError(f"TAV trafik bülteni: sayısal olmayan hücre {baglam}: {hucre!r}")
    return float(hucre)


def _sayfa_noktalari(sayfa) -> dict[tuple[str, str], tuple[str, float]] | None:
    """Bir sayfanın `{(varlık, segment): (tarih, değer)}` noktalarını çıkarır.

    Sayfanın ay/yıl başlığı ve "Passengers" satırı çözülemezse RuntimeError
    (sessiz eksik veri yasak). Yalnızca raporlanan yılın kendi sütunu
    başlıkta YOKSA (tek bilinen örnek: 0120/Ocak 2020) None döner — bu ay
    atlanır, diğerleri etkilenmez (bkz. modül docstring'i).
    """
    satirlar = list(sayfa.iter_rows(values_only=True))
    baslik = _sayfa_basligi(satirlar)
    if baslik is None:
        raise RuntimeError(f"TAV trafik bülteni: '{sayfa.title}' sayfasında başlık bulunamadı")
    yil, ay = _sayfa_tarihi(sayfa.title, baslik)
    tarih = f"{yil}-{ay:02d}-01"

    hdr_idx = _yolcu_basligi_satiri(satirlar)
    if hdr_idx is None:
        raise RuntimeError(f"TAV trafik bülteni ({tarih}): 'Passengers' başlık satırı bulunamadı")
    sutun = _yil_sutunu(satirlar[hdr_idx], yil)
    if sutun is None:
        return None  # Ölçüldü: yalnızca 0120 — bkz. modül docstring'i.

    noktalar: dict[tuple[str, str], tuple[str, float]] = {}
    gorulen_varlik: set[str] = set()
    i = hdr_idx + 1
    while i < len(satirlar):
        ham_ad = satirlar[i][0] if satirlar[i] else None
        if ham_ad is None:
            break  # yolcu bloğu bitti (Air Traffic Movements'tan önceki boş satır)
        ad = str(ham_ad).strip()
        if ad in _ALT_ETIKETLER:
            raise RuntimeError(
                f"TAV trafik bülteni ({tarih}): segment etiketi ('{ad}') varlık adından önce geldi"
            )
        if ad in gorulen_varlik:
            raise RuntimeError(f"TAV trafik bülteni ({tarih}): yinelenen varlık '{ad}'")
        gorulen_varlik.add(ad)

        deger = _hucre_oku(satirlar[i], sutun, f"'{ad}' ({tarih})")
        if deger is not None:
            noktalar[(ad, "toplam")] = (tarih, deger)
        i += 1

        if i < len(satirlar):
            dis_ad_ham = satirlar[i][0] if satirlar[i] else None
            dis_ad = str(dis_ad_ham).strip() if dis_ad_ham is not None else None
            if dis_ad in _ALT_ETIKETLER and _ALT_ETIKETLER[dis_ad] == "dis-hat":
                dis_deger = _hucre_oku(satirlar[i], sutun, f"'{ad}' dış hat ({tarih})")
                if dis_deger is not None:
                    noktalar[(ad, "dis-hat")] = (tarih, dis_deger)
                i += 1
                ic_ad_ham = satirlar[i][0] if i < len(satirlar) and satirlar[i] else None
                ic_ad = str(ic_ad_ham).strip() if ic_ad_ham is not None else None
                if ic_ad not in _ALT_ETIKETLER or _ALT_ETIKETLER[ic_ad] != "ic-hat":
                    raise RuntimeError(
                        f"TAV trafik bülteni ({tarih}): '{ad}' için Dış Hat'tan sonra "
                        f"İç Hat satırı beklenirken {ic_ad!r} geldi"
                    )
                ic_deger = _hucre_oku(satirlar[i], sutun, f"'{ad}' iç hat ({tarih})")
                if ic_deger is not None:
                    noktalar[(ad, "ic-hat")] = (tarih, ic_deger)
                i += 1
    return noktalar


def yolcu_noktalari(baytlar: bytes) -> dict[tuple[str, str], dict[str, float]]:
    """XLSX baytlarından `{(varlık, segment): {"YYYY-MM-01": değer}}` çıkarır.

    `Sheet1` (son ay) + tüm `MMYY` sayfaları işlenir; `disclaimer` hariç.
    Aynı (varlık, segment, tarih) iki farklı sayfadan üretilirse (şablon
    bozukluğu/çakışma) RuntimeError.
    """
    kitap = openpyxl.load_workbook(io.BytesIO(baytlar), data_only=True)
    try:
        noktalar: dict[tuple[str, str], dict[str, float]] = {}
        herhangi_sayfa_ok = False
        for ad in kitap.sheetnames:
            if ad == "disclaimer":
                continue
            sayfa_noktalari = _sayfa_noktalari(kitap[ad])
            if sayfa_noktalari is None:
                continue
            herhangi_sayfa_ok = True
            for anahtar, (tarih, deger) in sayfa_noktalari.items():
                gunluk = noktalar.setdefault(anahtar, {})
                if tarih in gunluk:
                    raise RuntimeError(
                        f"TAV trafik bülteni: {anahtar!r} için {tarih} birden fazla sayfada üretildi"
                    )
                gunluk[tarih] = deger
        if not herhangi_sayfa_ok:
            raise RuntimeError("TAV trafik bülteninde tek bir tanınan sayfa şablonu bulunamadı")
        return noktalar
    finally:
        kitap.close()


def _dosya_indir(url: str, session=None) -> bytes:
    http = session or requests
    yanit = http.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=ZAMAN_ASIMI)
    if yanit.status_code != 200:
        raise RuntimeError(f"TAV trafik bülteni HTTP {yanit.status_code}: {url}")
    return yanit.content


def seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse dosya listesi + ayrıştırma sonucu koşu boyunca
    paylaşılır: 28 seri aynı tek dosyayı okuduğu için yoksa 28 indirme +
    28 dosya-listesi kazıması olurdu.
    """
    onbellek = {} if onbellek is None else onbellek
    if "noktalar" not in onbellek:
        dosyalar = onbellek.get("dosyalar")
        if dosyalar is None:
            dosyalar = dosya_listesi(session)
            onbellek["dosyalar"] = dosyalar
        baytlar = _dosya_indir(dosyalar[0], session)
        onbellek["noktalar"] = yolcu_noktalari(baytlar)
    noktalar = onbellek["noktalar"]

    anahtar = (seri.tav_varlik, seri.tav_segment)
    if anahtar not in noktalar:
        raise RuntimeError(
            f"TAV trafik bülteninde varlık/segment bulunamadı: {anahtar!r} ({seri.id})"
        )

    df = pd.DataFrame(sorted(noktalar[anahtar].items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
