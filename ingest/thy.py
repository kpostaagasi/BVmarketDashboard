"""THY (Türk Hava Yolları) yatırımcı ilişkileri trafik bülteni istemcisi.

Kaynak gerçekleri 2026-09-18'de canlı ölçüldü:

- Dosya listesi TÜRETİLEMEZ: geçmiş yıllardaki xlsx bağlantıları düzensiz
  adlandırılmış (`turkish-airlines-trafik-ara_2025.xlsx`,
  `turkish-airlines-trafik-aralik-2022.xlsx`, biri `/trafik/` alt dizini
  olmadan `/documents/turkish-airlines-trafik-aralik_2024.xlsx`). Liste
  `TRAFIK_SAYFASI`dan kazınır; sayfa sunucu tarafında render edilir, JS
  gerekmez.
- BÜLTEN ŞABLONU YILA GÖRE DEĞİŞİYOR (iki farklı format ölçüldü):
  * Yeni şablon (cari yıl dosyasının kendi sayfaları, ör. 2026 dosyasının
    '2025' ve '2026' sayfaları): ay başlığı satır 2'de, segment adı
    ("Toplam"/"Yurt İçi"/"Yurt Dışı") metrik satırıyla AYNI satırda A
    sütununda, yalnızca segmentin ilk metrik satırında yazılı.
  * Eski şablon (yıl sonu arşiv dosyaları, ör. 2025/2024/2023/2022 dosyaları):
    segment adı ("TOPLAM"/"YURT İÇİ"/"YURT DIŞI", büyük harf) kendi başına
    bir satırda B sütununda, ay başlığı bir sonraki satırda; ayrıca bizim
    izlemediğimiz üç ek ölçüt var (Ücretli Yolcu Km, Uçulan Km, Dıştan Dışa
    Transfer Yolcu Sayısı) — bunlar sessizce atlanır (bilinmeyen ÖLÇÜT
    şablon farkıdır, hata değil; bilinmeyen SEGMENT hâlâ hatadır).
  Bu yüzden ayrıştırma satır pozisyonuna değil, hücre İÇERİĞİNE göre çalışır:
  her satırda hem A hem B sütunu segment sözlüğünde aranır (segment
  güncellenir), ay başlığı `AYLAR` ile birebir eşleşen satır olarak bulunur.
- HER DOSYA İKİ-DÖRT YIL SAYFASI TAŞIR (`Notlar`/`Notes` hariç, sayfa adı
  yılın kendisi, ör. `2022 dosyası → ['Notes','2019','2021','2022']`).
  2020 hiçbir dosyada yok (COVID döneminde bülten farklı yayımlanmış
  olabilir — kaynak sağlamıyor, sessizce atlanır çünkü sayfa hiç yok).
- Her sayfada satır 21 (yeni şablon) / 32 (eski şablon) civarında 'BÖLGESEL'
  başlıklı bir alt tablo var; bu tablo 'Yurt İçi'/'Toplam' gibi ana
  segment×ölçüt çiftlerini TEKRAR listeliyor (bölgesel kırılım öncesi özet
  satırı olarak). Ayrıştırma 'BÖLGESEL' satırında DURUR — durmazsa aynı
  (segment, ölçüt) anahtarı ikinci kez yazılmaya çalışılır ve
  `trafik_noktalari` hata verir.
- Ay hücresi boşsa (yayımlanmamış ay, cari yıl dosyasının gelecek ayları)
  hücre `None`; bu ay o (segment, ölçüt) sözlüğüne hiç eklenmez (sıfır
  uydurulmaz).
- Ölçüldü: Ağustos 2026 Toplam/Yurt İçi/Yurt Dışı Yolcu Sayısı sırasıyla
  10.001.394 / 3.591.405 / 6.409.989 (referans: marketvisuals.net/thy_traffic.html).
- Aynı yıl birden fazla dosyada geçebilir (ör. 2025 hem cari yıl dosyasının
  '2025' sayfasında hem 2025 arşiv dosyasında var); `dosya_listesi` sırası
  cari yıldan geçmişe doğru olduğu için birleştirme sırasında en son işlenen
  (en eski/arşiv) dosya kazanır — arşiv dosyası yıl sonu kesin rakamları
  taşır, cari yıl dosyasındaki devir değerleri küçük revizyonlar içerebilir.
"""

from __future__ import annotations

import io
import re

import openpyxl
import pandas as pd
import pdfplumber
import requests

from core.catalog import GECERLI_THY_METRIKLERI
from ingest.ir_sunum import tr_sayi

TABAN = "https://investor.turkishairlines.com"
TRAFIK_SAYFASI = f"{TABAN}/tr/mali-ve-operasyonel-veriler/trafik"
ZAMAN_ASIMI = 60
_BASLIKLAR = {"User-Agent": "Mozilla/5.0"}

BOLGESEL_ETIKETI = "BÖLGESEL"

AYLAR = (
    "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
    "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
)

# Bültendeki tam ölçüt etiketi -> katalogdaki kısa ad (`thy_olcut`).
OLCUT_ESLEME = {
    "Konma Sayısı (Yolcu Seferleri)": "konma",
    "Arzedilen Koltuk Km ('000)": "ask",
    "Yolcu Doluluk Oranı (%)": "doluluk",
    "Yolcu Sayısı": "yolcu",
    "Kargo + Posta (Ton)": "kargo",
}

# Eski şablonun (2022-2025 arşiv dosyaları) bilinen ama katalogda karşılığı
# olmayan ek ölçütleri: gerçek şablon farkı, ayrıştırma hatası değil.
# Buradaki listede OLMAYAN bir etiket (yazım hatası, gerçek şablon
# bozulması) hâlâ RuntimeError'a düşer — sessizce atlanmaz.
OLCUT_IZLENMEYEN = {
    "Ücretli Yolcu Km ('000)",
    "Uçulan Km ('000)",
    "Dıştan Dışa Transfer Yolcu Sayısı",
}

# Bültendeki segment etiketi (iki şablonun iki farklı yazımı) -> katalogdaki
# ad (`thy_segment`). A sütunu (yeni şablon) YALNIZCA segment adı taşır —
# bu yüzden A sütunundaki her metin bu sözlükte olmak ZORUNDADIR (aksi hata).
SEGMENT_ESLEME = {
    "Toplam": "Toplam", "TOPLAM": "Toplam",
    "Yurt İçi": "Yurt İçi", "YURT İÇİ": "Yurt İçi",
    "Yurt Dışı": "Yurt Dışı", "YURT DIŞI": "Yurt Dışı",
}


def dosya_listesi(session=None) -> list[str]:
    """Trafik sayfasındaki xlsx bağlantılarını mutlak URL olarak döner.

    URL deseni yıldan yıla tutarsız olduğu için türetilemez; sayfa
    kazınmalı. Sayfa hiç xlsx bağlantısı içermiyorsa (şablon değişmiş,
    sayfa taşınmış) sessizce boş liste dönmek yerine hata verilir.
    """
    http = session or requests
    yanit = http.get(TRAFIK_SAYFASI, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"THY trafik sayfası HTTP {yanit.status_code} döndü")
    html = yanit.content.decode("utf-8", errors="replace")
    baglantilar = re.findall(r'href="([^"]+\.xlsx)"', html)
    if not baglantilar:
        raise RuntimeError("THY trafik sayfasında xlsx bağlantısı bulunamadı")
    return [
        baglanti if baglanti.startswith("http") else f"{TABAN}{baglanti}"
        for baglanti in baglantilar
    ]


def _satiri_isle(
    satir: tuple, yil: int, mevcut_segment: str | None,
    gorulen: set[tuple[str, str]], noktalar: dict[tuple[str, str], dict[str, float]],
) -> tuple[str | None, bool]:
    """Bir veri satırını işler; `(güncel segment, BÖLGESEL'e ulaşıldı mı)` döner."""
    col0, col1 = satir[0], satir[1]
    veri_var = any(deger is not None for deger in satir[2:14])

    # A sütunu (yeni şablonda segmentin ilk metrik satırı) yalnızca segment
    # adı taşır; burada tanınmayan bir metin gerçek şablon bozulmasıdır.
    metin0 = str(col0).strip() if col0 is not None else ""
    if metin0:
        if metin0.upper() == BOLGESEL_ETIKETI:
            return mevcut_segment, True
        segment = SEGMENT_ESLEME.get(metin0)
        if segment is None:
            raise RuntimeError(f"THY bülteninde {yil}: bilinmeyen segment {metin0!r}")
        mevcut_segment = segment

    # B sütunu (eski şablonda segment başlığı, her iki şablonda ölçüt
    # etiketi ya da dipnot metni) — burada tanınmayan metin serbesttir,
    # aşağıda veri varsa ayrıca ölçüt olarak doğrulanır.
    metin1 = str(col1).strip() if col1 is not None else ""
    if metin1:
        if metin1.upper() == BOLGESEL_ETIKETI:
            return mevcut_segment, True
        segment = SEGMENT_ESLEME.get(metin1)
        if segment is not None:
            mevcut_segment = segment

    if not veri_var:
        # Segment-only başlık satırı (eski şablon), boş ayraç satırı ya da
        # dipnot metni (ör. "* Veriler tüm uçuşları içermektedir.") — hiçbiri
        # ölçüt verisi taşımaz.
        return mevcut_segment, False

    if not metin1:
        raise RuntimeError(f"THY bülteninde {yil}: ölçüt etiketi olmayan veri satırı")
    olcut = OLCUT_ESLEME.get(metin1)
    if olcut is None:
        if metin1 in OLCUT_IZLENMEYEN:
            return mevcut_segment, False
        raise RuntimeError(f"THY bülteninde {yil}: bilinmeyen ölçüt {metin1!r}")

    if mevcut_segment is None:
        raise RuntimeError(
            f"THY bülteninde {yil}: segment belirlenmeden ölçüt satırı geldi "
            f"({metin1!r})"
        )
    anahtar = (mevcut_segment, olcut)
    if anahtar in gorulen:
        raise RuntimeError(
            f"THY bülteninde {yil}: yinelenen segment/ölçüt anahtarı {anahtar} "
            "('BÖLGESEL' bloğu durdurulmadan okunmuş olabilir)"
        )
    gorulen.add(anahtar)

    aylik: dict[str, float] = {}
    for ay, hucre in enumerate(satir[2:14], start=1):
        if hucre is None:
            continue  # yayımlanmamış ay — sıfır uydurulmaz, atlanır
        if not isinstance(hucre, (int, float)):
            raise RuntimeError(
                f"THY bülteninde {yil}: sayısal olmayan hücre {metin1!r} "
                f"{yil}-{ay:02d}: {hucre!r}"
            )
        aylik[f"{yil}-{ay:02d}-01"] = float(hucre)
    noktalar.setdefault(anahtar, {}).update(aylik)
    return mevcut_segment, False


def trafik_noktalari(baytlar: bytes) -> dict[tuple[str, str], dict[str, float]]:
    """XLSX baytlarından `{(segment, ölçüt): {"YYYY-MM-01": değer}}` çıkarır.

    Dosyadaki `Notlar`/`Notes` dışında adı yıl olan TÜM sayfalar okunur
    (bir dosya 2-4 yıl taşıyabilir). Her sayfada 'BÖLGESEL' satırında durulur.
    """
    kitap = openpyxl.load_workbook(io.BytesIO(baytlar), data_only=True)
    try:
        noktalar: dict[tuple[str, str], dict[str, float]] = {}
        yil_sayfasi_var = False
        for ad in kitap.sheetnames:
            if not ad.isdigit():
                continue
            yil_sayfasi_var = True
            yil = int(ad)
            sayfa = kitap[ad]
            ay_basligi_bulundu = False
            mevcut_segment: str | None = None
            gorulen: set[tuple[str, str]] = set()
            for satir in sayfa.iter_rows(
                min_row=1, max_row=sayfa.max_row, min_col=1, max_col=14,
                values_only=True,
            ):
                if tuple(satir[2:14]) == AYLAR:
                    ay_basligi_bulundu = True
                    continue
                mevcut_segment, durduruldu = _satiri_isle(
                    satir, yil, mevcut_segment, gorulen, noktalar
                )
                if durduruldu:
                    break
            if not ay_basligi_bulundu:
                raise RuntimeError(f"THY bülteninde {yil}: ay başlığı bulunamadı")
        if not yil_sayfasi_var:
            raise RuntimeError("THY bülteninde yıl sayfası bulunamadı")
        return noktalar
    finally:
        kitap.close()


def _tum_noktalari_getir(
    onbellek: dict, session=None,
) -> dict[tuple[str, str], dict[str, float]]:
    """Dosya listesini ve her dosyanın ayrıştırılmış noktalarını önbellekler.

    15 seri (3 segment × 5 ölçüt) aynı dosya kümesini paylaşır; önbelleksiz
    her seri kendi dosyalarını yeniden indirir.
    """
    if "noktalar" in onbellek:
        return onbellek["noktalar"]

    http = session or requests
    birlesik: dict[tuple[str, str], dict[str, float]] = {}
    # Sıra `dosya_listesi`nin döndürdüğü sırayla (cari yıldan geçmişe);
    # aynı yıl birden fazla dosyada geçtiğinde en son işlenen (en eski,
    # dolayısıyla o yılın kendi yıl-sonu arşivi) kazanır — arşiv dosyası
    # kesin rakamları taşır.
    for url in dosya_listesi(session):
        yanit = http.get(url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
        if yanit.status_code != 200:
            raise RuntimeError(f"THY dosyası indirilemedi ({url}): HTTP {yanit.status_code}")
        for anahtar, aylik in trafik_noktalari(yanit.content).items():
            birlesik.setdefault(anahtar, {}).update(aylik)

    onbellek["noktalar"] = birlesik
    return birlesik


PDF_TOPLAM_OLCUTLERI = {"ucak-sayisi", "uculan-nokta"}

# Aylık trafik PDF bülteninin TOPLAM tablosundaki iki ölçüt (XLSX'te YOK):
# filo/uçulan-hatlar sayfalarının "canlı" (yalnızca bugünü gösteren, geçmişi
# olmayan) sürümleri yerine, her ay yayımlanan PDF'in kendi TOPLAM tablosu
# kullanılır — gerçek, tarihli kaynak (bkz. modül sonu).
_TRAFIK_PDF_BASLIGI = re.compile(r"^([A-ZÇĞİÖŞÜ]+)\s+(\d{4})\s+TRAFİK$", re.M)
_UCAK_SAYISI_DESENI = re.compile(r"^Uçak Sayısı (\d+) (\d+) %", re.M)
_UCULAN_NOKTA_DESENI = re.compile(r"^Uçulan Nokta (\d+) (\d+) %", re.M)
_AY_BUYUK_HARF = {ay.upper(): i + 1 for i, ay in enumerate(AYLAR)}


def pdf_toplam_noktalari(metin: str) -> tuple[str, float, float]:
    """Aylık trafik PDF'inin TOPLAM tablosundan `(tarih, uçak_sayısı,
    uçulan_nokta)` çıkarır. Yalnızca PDF'in ilk (TOPLAM) bloğu okunur —
    YURT DIŞI/YURT İÇİ bloklarında bu iki ölçüt hiç yok (yalnızca filo
    genelinde ve ağ genelinde anlamlı).
    """
    baslik = _TRAFIK_PDF_BASLIGI.search(metin)
    if not baslik:
        raise RuntimeError("THY trafik PDF'inde '<AY> <YIL> TRAFİK' başlığı bulunamadı")
    ay_adi, yil = baslik.group(1), int(baslik.group(2))
    ay = _AY_BUYUK_HARF.get(ay_adi)
    if ay is None:
        raise RuntimeError(f"THY trafik PDF'inde bilinmeyen ay adı: {ay_adi!r}")
    tarih = f"{yil}-{ay:02d}-01"

    ucak_m = _UCAK_SAYISI_DESENI.search(metin)
    nokta_m = _UCULAN_NOKTA_DESENI.search(metin)
    if not ucak_m or not nokta_m:
        raise RuntimeError(
            "THY trafik PDF'inde TOPLAM tablosunda 'Uçak Sayısı'/'Uçulan Nokta' "
            "satırları bulunamadı (şablon değişmiş olabilir)"
        )
    return tarih, float(ucak_m.group(2)), float(nokta_m.group(2))


def pdf_toplam_seri_cek(seri, session=None) -> pd.DataFrame:
    """Trafik sayfasındaki EN GÜNCEL aylık PDF bülteninden "ucak-sayisi" ya
    da "uculan-nokta" serisini çeker (tek nokta — bülten yalnızca cari ayı
    ve bir önceki yılın aynı ayını gösterir, geriye dönük arşiv sayfa
    başına tek dosya değil; bkz. modül docstring'i, dosya_listesi ile aynı
    kısıt). Ölçüldü (2026-09-20): Ağustos 2026 Uçak Sayısı=566, Uçulan
    Nokta=358 — marketvisuals.net kart özetiyle birebir eşleşiyor.
    """
    http = session or requests
    yanit = http.get(TRAFIK_SAYFASI, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"THY trafik sayfası HTTP {yanit.status_code} döndü")
    baglantilar = re.findall(r'href="([^"]+/trafik/[^"]+\.pdf)"', yanit.text)
    if not baglantilar:
        raise RuntimeError("THY trafik sayfasında aylık PDF bağlantısı bulunamadı")
    pdf_url = baglantilar[0]
    if not pdf_url.startswith("http"):
        pdf_url = f"{TABAN}{pdf_url}"

    pdf_yaniti = http.get(pdf_url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if pdf_yaniti.status_code != 200:
        raise RuntimeError(f"THY trafik PDF'i indirilemedi ({pdf_url}): HTTP {pdf_yaniti.status_code}")
    with pdfplumber.open(io.BytesIO(pdf_yaniti.content)) as pdf:
        metin = pdf.pages[0].extract_text() or ""
    tarih, ucak_sayisi, uculan_nokta = pdf_toplam_noktalari(metin)

    deger = {"ucak-sayisi": ucak_sayisi, "uculan-nokta": uculan_nokta}[seri.thy_olcut]
    df = pd.DataFrame([(tarih, deger)], columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)


def seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """THY trafik bültenlerinden tek bir (segment, ölçüt) serisini çeker.

    `onbellek` verilirse dosya listesi ve ayrıştırılmış noktalar koşu
    boyunca paylaşılır: 15 seri aynı ~5 dosyayı okuduğu için yoksa 75
    indirme olurdu. "ucak-sayisi"/"uculan-nokta" XLSX bülteninde YOK —
    aylık PDF bülteninin TOPLAM tablosundan ayrı bir yolla çekilir (bkz.
    `pdf_toplam_seri_cek`).
    """
    if seri.thy_olcut in PDF_TOPLAM_OLCUTLERI:
        return pdf_toplam_seri_cek(seri, session=session)
    onbellek = {} if onbellek is None else onbellek
    tum_noktalar = _tum_noktalari_getir(onbellek, session=session)
    anahtar = (seri.thy_segment, seri.thy_olcut)
    aylik = tum_noktalar.get(anahtar)
    if not aylik:
        raise RuntimeError(
            f"THY bülteninde segment/ölçüt bulunamadı: {anahtar} ({seri.id})"
        )

    df = pd.DataFrame(sorted(aylik.items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)


# --- Çeyreklik Yatırımcı Sunumu (kaynak_tipi: thy_ir) ---
#
# Kaynak gerçekleri 2026-09-20'de canlı ölçüldü (2Ç'26 sunumu,
# ir-presentation-2q26tr.pdf, 35 sayfa):
#
# - "Sunumlar" sayfası (`SUNUMLAR_SAYFASI`) en yeni çeyrekten en eskiye
#   sıralı `ir-presentation-NqYYtr.pdf` bağlantıları listeler (ör.
#   "2q26", "1q26", "4q25", "3q25", ...); ilk eşleşme kullanılır. Aynı
#   sayfada "…-basin-sunumu.pdf" (basın sunumu) ve "…infografik…pdf" gibi
#   FARKLI belgeler de var — yalnızca `ir-presentation-` önekli olanlar
#   eşleşir.
# - Çoğu tablo pdfplumber `extract_text()` ile TEMİZ satırlar üretiyor
#   (tablo çizgileri yok ama hücreler tek satıra akıyor, sütun sırası
#   sabit). İKİ istisna GRAFİK (bar chart) olarak gömülü, düz metin akışı
#   YOK: "Net Borç / EBITDA Oranı" (sayfa ~20) ve "Birim Gelir Gelişimi"
#   içindeki "Passenger RASK" satırı (sayfa ~3-4, RASK2 ve Yield AYNI
#   sayfada ama "Özet Finansal Veriler" tablosunda TEMİZ olarak da var,
#   o yüzden Passenger RASK dışında grafiğe hiç gidilmiyor). Bu ikisi için
#   metin akış SIRASI değer büyüklüğüne göre karışabiliyor (uzun çubuk =
#   üstte = önce okunur) — bu yüzden `extract_words()` ile x0 KONUM
#   eşlemesi kullanılır (aynı çubuğun etiketi ve değeri sayfada yatayda
#   hizalı), sıra eşlemesi DEĞİL. Ölçüldü: Net Borç/FAVÖK 8 değer/8 etiket
#   x0-farkı ≤5pt (S12A 2Ç'26 değeri 2,1x, etiketle x0 farkı 0,0pt);
#   Passenger RASK 2Ç'25/2Ç'26 değerleri x0-farkı ~2pt.
# - Sayfa numaraları çeyrekten çeyreğe KAYABİLİR (bu sunumda bir "Filo"
#   bölüm ayırıcı sayfası var, önceki çeyreklerde olmayabilir) — bu yüzden
#   sabit sayfa indeksi YERİNE her tablo kendi başlık metniyle aranır
#   (`_sayfa_metni_bul`).
# - "Özet Finansal Veriler" / "Operasyonel Gider Kırılımı" / "Birim Gider
#   (CASK) Kırılımı" / "KPI" / "Bölgesel Birim Gelir Değişimi" tabloları
#   yalnızca 2Ç(çeyrek) ve 6A(kümülatif yarı yıl) sütunları taşır; 6A
#   OKUNMAZ (farklı bir periyot uzunluğu — çeyreklik seriyle karıştırılırsa
#   yanlış damgalanmış nokta üretir). Her tablo İKİ nokta verir: cari
#   çeyrek ve bir önceki yılın aynı çeyreği (ör. 2Ç'25 ve 2Ç'26).
# - "Bilanço" tablosu YEDİ sütun taşır (2020-2025 yıl sonu + cari çeyrek
#   sonu, ör. "30.06.2026") — yıl sonu noktaları çeyreğin ilk ayı kuralıyla
#   O YILIN 4. ÇEYREĞİNE (Ekim) damgalanır (bilanço bir dönem sonu anlık
#   görüntüsüdür, FY rakamı fiilen Aralık sonudur).
# - "Filo" tablosunda alt tip satırları (ör. "A350-9") BOŞ hücreleri
#   (sıfır Finansal Kira/Opr. Kira) atlayarak farklı sütun sayısında
#   geliyor — bu yüzden alt tip satırları OKUNMAZ, yalnızca her gövde
#   grubunun ("Geniş Gövde"/"Dar Gövde"/"Kargo") "Toplam" alt toplam satırı
#   ve "Genel Toplam" satırı okunur (bunlar HER ZAMAN tam sütun sayısında).
#   Kargo grubunun "Toplam" satırında "Ortalama Filo Yaşı" sütunu BOŞ
#   (yalnızca 5 sayı, 6 değil) — son sütun bu yüzden OPSİYONEL.
# - Ölçüldü (2026-09-20): RASK2 8,93 / Yield 9,83 / Yakıt Fiyatı 1.480 /
#   Net Borç-FAVÖK S12A 2,1x / Bilanço Toplam Varlık 50.671 / Filo Genel
#   Toplam 552 — marketvisuals.net kart özetleriyle birebir eşleşiyor.

TABAN_YATIRIMCI = TABAN
SUNUMLAR_SAYFASI = f"{TABAN}/tr/mali-ve-operasyonel-veriler/sunumlar"
_SUNUM_BAGLANTISI = re.compile(r'href="([^"]*ir-presentation-(\d)q(\d{2})tr\.pdf)"', re.I)

_CEYREK_ILK_AY = {1: "01", 2: "04", 3: "07", 4: "10"}


def _ceyrek_tarihi(yil: int, ceyrek: int) -> str:
    """Katalog sözleşmesi: çeyreksel seri tarihi çeyreğin İLK ayı."""
    return f"{yil}-{_CEYREK_ILK_AY[ceyrek]}-01"


def sunum_bilgisi(session=None) -> tuple[str, int, int]:
    """Sunumlar sayfasından en güncel çeyreklik yatırımcı sunumu PDF'inin
    `(url, yıl, çeyrek)` bilgisini döner. Sayfa en yeniden en eskiye
    sıralı; ilk eşleşme kullanılır."""
    http = session or requests
    yanit = http.get(SUNUMLAR_SAYFASI, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"THY sunumlar sayfası HTTP {yanit.status_code} döndü")
    m = _SUNUM_BAGLANTISI.search(yanit.text)
    if not m:
        raise RuntimeError("THY sunumlar sayfasında çeyreklik yatırımcı sunumu bağlantısı bulunamadı")
    url = m.group(1)
    return (url if url.startswith("http") else f"{TABAN}{url}"), 2000 + int(m.group(3)), int(m.group(2))


def _sayfa_metni_bul(pdf, icerir: str) -> tuple[int, str]:
    """Verilen alt dizeyi içeren İLK sayfayı bulur; `(sayfa indeksi, metin)`
    döner. Sabit sayfa numarası yerine içerikle arama: sunum çeyrekten
    çeyreğe sayfa sırasını değiştirebiliyor."""
    for i, sayfa in enumerate(pdf.pages):
        metin = sayfa.extract_text() or ""
        if icerir in metin:
            return i, metin
    raise RuntimeError(f"THY sunumunda {icerir!r} içeren sayfa bulunamadı (şablon değişmiş olabilir)")


# "Özet Finansal Veriler" sayfası: RASK2, Yield, Kargo Geliri satırları.
# Dipnot rakamları etikete bitişik gelir (ör. "RASK22", "(R/Y)3") — `\d?`
# ile toleranslı.
_OZET_DESENLERI = {
    "rask2": re.compile(r"([\d,]+)\s+([\d,]+)\s+%[\d,]+\s+RASK2\d?\s*\(AKTK dahil\)"),
    "yield": re.compile(r"([\d,]+)\s+([\d,]+)\s+%[\d,]+\s+Yolcu Birim Gelir \(R/Y\)\d?\s*\(Usc\)"),
    "kargo-geliri": re.compile(r"([\d.]+)\s+([\d.]+)\s+%[\d,]+\s+Kargo Geliri\b"),
}

# "Operasyonel Gider Kırılımı" sayfası: Akaryakıt (Yakıt Gideri) satırı —
# tablo 4 rakam+2 yüzde ile biter ("... %55,1 %32,1"), bu yüzden ardından
# gelen ikinci grup da eşleştirilerek yanlışlıkla "Akaryakıt Hariç" benzeri
# bir satıra kaymadığından emin olunur.
_GIDER_DESENI = re.compile(r"([\d.]+)\s+([\d.]+)\s+%[\d,]+\s+Akaryakıt\s+[\d.]+\s+[\d.]+\s+%[\d,]+\s+%")

# "Birim Gider (CASK) Kırılımı" sayfası. "CASK" ve "Akaryakıt" etiketlerinin
# hemen ardından SAYI gelmesi zorunlu tutularak "CASK2 (AKTK dahil)2" ve
# "Akaryakıt Hariç CASK..." satırlarıyla karışması önlenir.
_CASK_DESENLERI = {
    "yakit-cask": re.compile(r"([\d,]+)\s+([\d,]+)\s+%[\d,]+\s+Akaryakıt\s+[\d,]"),
    "cask": re.compile(r"([\d,]+)\s+([\d,]+)\s+%[\d,]+\s+CASK\s+[\d,]"),
    "cask-ex-fuel": re.compile(r"([\d,]+)\s+([\d,]+)\s+%[\d,]+\s+Akaryakıt Hariç CASK\s+[\d,]"),
}

# "Önemli Performans Göstergeleri (KPI)" sayfası: Akaryakıt (USD/ton).
_KPI_DESENI = re.compile(r"([\d.]+)\s+([\d.]+)\s+%[\d,]+\s+Akaryakıt \(USD/ton\)")

# "Bilanço" sayfası: yedi yıl/dönem sütunu (2020..2025 + dönem sonu tarihi).
_BILANCO_ETIKET_ESLEME = {
    "Toplam Varlıklar": "varlik-toplam",
    "Toplam Yükümlülükler": "yukumluluk-toplam",
    "Toplam Özkaynaklar": "ozkaynak-toplam",
    "Kira Yükümlülükleri": "yukumluluk-kira",
    "Banka Kredileri": "yukumluluk-banka-kredisi",
    "Yolcu Uçuş Yükümlülükleri": "yukumluluk-yolcu-ucus",
    "Ticari Borçlar": "yukumluluk-ticari-borc",
    "Diğer Yükümlülükler": "yukumluluk-diger",
}
_BILANCO_BASLIK_DESENI = re.compile(
    r"(?:Varlıklar|Yükümlülükler) \(mn USD\) 2020 2021 2022 2023 2024 2025 (\d{2})\.(\d{2})\.(\d{4})"
)
_BILANCO_SATIR_DESENI = re.compile(
    r"^(" + "|".join(re.escape(k) for k in _BILANCO_ETIKET_ESLEME) + r")"
    r"\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)$",
    re.M,
)

# "Filo" sayfası: gövde grubu alt toplamları (Geniş Gövde / Dar Gövde /
# Kargo sırasıyla — sayfada bu sırayla geçer) + Genel Toplam.
_FILO_BASLIK_DESENI = re.compile(r"Filo \((\d{2})\.(\d{2})\.(\d{4}) İtibarıyla\)")
_FILO_TOPLAM_DESENI = re.compile(r"^Toplam (\d+) (\d+) (\d+) (\d+) ([\d,]+)(?: ([\d,]+))?$", re.M)
_FILO_GENEL_TOPLAM_DESENI = re.compile(r"^Genel Toplam (\d+) (\d+) (\d+) (\d+) ([\d,]+)(?: ([\d,]+))?$", re.M)
_FILO_GOVDE_SIRASI = ("filo-genis-govde", "filo-dar-govde", "filo-kargo")
_FILO_MULKIYET_SIRASI = ("filo-sahip-olunan", "filo-finansal-kira", "filo-operasyonel-kira")

# "Dolar Bazında Bölgesel Birim Gelir Değişimi" sayfası: RASK2 satırı sayfada
# İKİ kez geçer (üç bölgelik iki grup); her geçişte 3 bölge × (2Ç,6A) çifti.
_BOLGE_SIRASI = (
    "rask2-degisim-amerika", "rask2-degisim-avrupa", "rask2-degisim-uzak-dogu",
    "rask2-degisim-afrika", "rask2-degisim-orta-dogu", "rask2-degisim-ic-hat",
)
_BOLGE_GRUP1_DESENI = re.compile(r"Amerika Avrupa Uzak Doğu")
_BOLGE_GRUP2_DESENI = re.compile(r"Afrika Orta Doğu İç Hat")
_BOLGE_RASK2_DESENI = re.compile(r"RASK2 (-?%[\d,]+) (-?%[\d,]+)")

# "Net Borç / EBITDA Oranı" grafiği: sekiz yıl/dönem kategorisi (2019..2025
# yıl sonu + S12A cari çeyrek). "2Ç'26" gibi ek eksen alt-etiketi kasıtlı
# OLARAK adaylara dahil edilmez — S12A'nın x0 konumu değerle zaten tam
# hizalı (ölçüldü: fark 0,0pt), "2Ç'26" alt-etiketi ayrı bir sütun değil.
_NET_BORC_DEGER_DESENI = re.compile(r"^\d[,.]?\d*x$")
_NET_BORC_ETIKET_DESENI = re.compile(r"^(20\d\d|S12A)$")
_NET_BORC_TOLERANS_PT = 20.0


def _passenger_rask_degerleri(sayfa) -> tuple[float, float]:
    """'Birim Gelir Gelişimi' sayfasındaki 'Passenger RASK' bar grafiğinden
    `(2Ç önceki yıl, 2Ç cari yıl)` değer çiftini x0 konum eşlemesiyle
    çıkarır (metin akış SIRASI değer büyüklüğüne göre karışabildiği için
    güvenilmez — bkz. modül docstring'i)."""
    kelimeler = sayfa.extract_words(use_text_flow=False, keep_blank_chars=False)
    return _passenger_rask_kelimelerden(kelimeler)


def _passenger_rask_kelimelerden(kelimeler: list[dict]) -> tuple[float, float]:
    """`_passenger_rask_degerleri`'nin saf mantığı — `extract_words()`
    çıktısı üzerinden test edilebilir (gerçek PDF gerektirmez)."""
    baslik = next((k for k in kelimeler if k["text"] == "Passenger"), None)
    sonraki_baslik = next((k for k in kelimeler if k["text"] == "Revenue"), None)
    if baslik is None or sonraki_baslik is None:
        raise RuntimeError("THY sunumunda 'Passenger RASK' grafiği bulunamadı (şablon değişmiş olabilir)")
    bant = [k for k in kelimeler if baslik["top"] <= k["top"] < sonraki_baslik["top"]]

    ceyrek_etiketleri = [k for k in bant if re.fullmatch(r"2Ç'\d\d", k["text"])]
    if len(ceyrek_etiketleri) < 2:
        raise RuntimeError("THY sunumunda Passenger RASK eksen etiketleri (2Ç'XX) eksik")
    etiket_onceki, etiket_simdiki = ceyrek_etiketleri[0], ceyrek_etiketleri[1]

    adaylar = [k for k in bant if re.fullmatch(r"\d+,\d+", k["text"])]
    if len(adaylar) < 2:
        raise RuntimeError("THY sunumunda Passenger RASK sayısal değerleri bulunamadı")

    def en_yakin(etiket):
        aday = min(adaylar, key=lambda a: abs(a["x0"] - etiket["x0"]))
        mesafe = abs(aday["x0"] - etiket["x0"])
        if mesafe > 15.0:
            raise RuntimeError(
                f"THY sunumunda Passenger RASK değeri {etiket['text']!r} etiketine "
                f"konumca eşleşmedi (en yakın {aday['text']!r}, {mesafe:.1f}pt)"
            )
        return tr_sayi(aday["text"])

    return en_yakin(etiket_onceki), en_yakin(etiket_simdiki)


def _net_borc_favok_noktalari(sayfa, yil: int, ceyrek: int) -> dict[str, float]:
    """'Net Borç / EBITDA Oranı' grafiğinden `{"YYYY-10-01": değer}` (2019-
    2025 yıl sonu) + `{cari çeyrek tarihi: değer}` (S12A) döner. Grafik
    metin akışı DEĞER BÜYÜKLÜĞÜNE göre karışık geliyor — x0 konum eşlemesi
    zorunlu (bkz. modül docstring'i)."""
    kelimeler = sayfa.extract_words(use_text_flow=False, keep_blank_chars=False)
    return _net_borc_favok_kelimelerden(kelimeler, yil, ceyrek)


def _net_borc_favok_kelimelerden(kelimeler: list[dict], yil: int, ceyrek: int) -> dict[str, float]:
    """`_net_borc_favok_noktalari`'nin saf mantığı — `extract_words()`
    çıktısı üzerinden test edilebilir (gerçek PDF gerektirmez)."""
    degerler = [k for k in kelimeler if _NET_BORC_DEGER_DESENI.match(k["text"])]
    etiketler = [k for k in kelimeler if _NET_BORC_ETIKET_DESENI.match(k["text"])]
    if len(degerler) != 8 or len(etiketler) != 8:
        raise RuntimeError(
            f"THY sunumunda Net Borç/EBITDA grafiğinde {len(degerler)} değer, "
            f"{len(etiketler)} etiket bulundu (8/8 bekleniyor)"
        )
    sonuc: dict[str, float] = {}
    for deger_kelime in degerler:
        en_yakin = min(etiketler, key=lambda e: abs(e["x0"] - deger_kelime["x0"]))
        mesafe = abs(en_yakin["x0"] - deger_kelime["x0"])
        if mesafe > _NET_BORC_TOLERANS_PT:
            raise RuntimeError(
                f"THY sunumunda Net Borç/EBITDA değeri {deger_kelime['text']!r} hiçbir "
                f"etikete {_NET_BORC_TOLERANS_PT}pt toleransında eşleşmedi"
            )
        if en_yakin["text"] in sonuc:
            raise RuntimeError(f"THY sunumunda Net Borç/EBITDA etiketi {en_yakin['text']!r} birden fazla değere eşleşti")
        sonuc[en_yakin["text"]] = tr_sayi(deger_kelime["text"].rstrip("xX"))

    noktalar: dict[str, float] = {}
    for etiket, deger in sonuc.items():
        if etiket == "S12A":
            noktalar[_ceyrek_tarihi(yil, ceyrek)] = deger
        else:
            noktalar[f"{etiket}-10-01"] = deger
    return noktalar


def _bilanco_noktalari(metin: str) -> dict[str, dict[str, float]]:
    """Bilanço tablosundan `{thy_metrik: {tarih: değer}}` çıkarır (7 nokta:
    2020-2025 yıl sonu + cari çeyrek sonu)."""
    baslik = _BILANCO_BASLIK_DESENI.search(metin)
    if not baslik:
        raise RuntimeError("THY sunumunda Bilanço başlık satırı (2020..2025 sütunları) bulunamadı")
    gg, aa, yyyy = baslik.groups()
    ceyrek = {"01": 1, "02": 1, "03": 1, "04": 2, "05": 2, "06": 2, "07": 3, "08": 3, "09": 3, "10": 4, "11": 4, "12": 4}[aa]
    tarihler = ["2020-10-01", "2021-10-01", "2022-10-01", "2023-10-01", "2024-10-01", "2025-10-01", _ceyrek_tarihi(int(yyyy), ceyrek)]

    satirlar = list(_BILANCO_SATIR_DESENI.finditer(metin))
    if len(satirlar) != len(_BILANCO_ETIKET_ESLEME):
        raise RuntimeError(
            f"THY sunumunda Bilanço tablosunda {len(satirlar)} satır bulundu "
            f"({len(_BILANCO_ETIKET_ESLEME)} bekleniyor)"
        )
    sonuc: dict[str, dict[str, float]] = {}
    for m in satirlar:
        metrik = _BILANCO_ETIKET_ESLEME[m.group(1)]
        degerler = [tr_sayi(g) for g in m.groups()[1:]]
        sonuc[metrik] = dict(zip(tarihler, degerler))
    return sonuc


def _filo_noktalari(metin: str) -> dict[str, float]:
    """Filo tablosundan gövde tipi (3) + mülkiyet yapısı (3) kompozisyon
    bileşenlerini döner; iki çapraz toplam kontrolüyle doğrulanır."""
    govde_satirlari = _FILO_TOPLAM_DESENI.findall(metin)
    if len(govde_satirlari) != 3:
        raise RuntimeError(
            f"THY sunumunda Filo tablosunda {len(govde_satirlari)} 'Toplam' alt "
            "toplam satırı bulundu (3 bekleniyor: Geniş Gövde/Dar Gövde/Kargo)"
        )
    genel = _FILO_GENEL_TOPLAM_DESENI.search(metin)
    if not genel:
        raise RuntimeError("THY sunumunda Filo tablosunda 'Genel Toplam' satırı bulunamadı")

    govde_toplamlari = [int(satir[0]) for satir in govde_satirlari]
    genel_toplam = int(genel.group(1))
    if sum(govde_toplamlari) != genel_toplam:
        raise RuntimeError(
            f"THY sunumunda Filo gövde tipi alt toplamları ({sum(govde_toplamlari)}) "
            f"Genel Toplam ({genel_toplam}) ile uyuşmuyor"
        )
    mulkiyet_degerleri = [int(genel.group(2)), int(genel.group(3)), int(genel.group(4))]
    if sum(mulkiyet_degerleri) != genel_toplam:
        raise RuntimeError(
            f"THY sunumunda Filo mülkiyet kırılımı toplamı ({sum(mulkiyet_degerleri)}) "
            f"Genel Toplam ({genel_toplam}) ile uyuşmuyor"
        )
    sonuc = dict(zip(_FILO_GOVDE_SIRASI, (float(v) for v in govde_toplamlari)))
    sonuc.update(zip(_FILO_MULKIYET_SIRASI, (float(v) for v in mulkiyet_degerleri)))
    return sonuc


def _bolgesel_noktalari(metin: str) -> dict[str, float]:
    """Bölgesel RASK2 YoY değişim tablosundan altı bölgenin 2Ç sütununu
    döner (6A kümülatif sütunu atlanır — bkz. modül docstring'i)."""
    if not (_BOLGE_GRUP1_DESENI.search(metin) and _BOLGE_GRUP2_DESENI.search(metin)):
        raise RuntimeError("THY sunumunda Bölgesel Birim Gelir Değişimi bölge başlıkları bulunamadı")
    if _BOLGE_GRUP1_DESENI.search(metin).start() > _BOLGE_GRUP2_DESENI.search(metin).start():
        raise RuntimeError("THY sunumunda Bölgesel Birim Gelir Değişimi bölge grup sırası beklenenden farklı")
    eslesmeler = _BOLGE_RASK2_DESENI.findall(metin)
    if len(eslesmeler) != 6:
        raise RuntimeError(
            f"THY sunumunda Bölgesel Birim Gelir Değişimi RASK2 satırında {len(eslesmeler)} "
            "bölge çifti bulundu (6 bekleniyor)"
        )
    return {ad: tr_sayi(ceyrek) for ad, (ceyrek, _altiay) in zip(_BOLGE_SIRASI, eslesmeler)}


def _sunumu_ayristir(pdf, yil: int, ceyrek: int) -> dict[str, dict[str, float]]:
    """Açılmış sunum PDF'inden TÜM thy_metrik noktalarını tek geçişte
    çıkarır: `{thy_metrik: {tarih: değer}}`."""
    onceki_tarih = _ceyrek_tarihi(yil - 1, ceyrek)
    bu_tarih = _ceyrek_tarihi(yil, ceyrek)
    noktalar: dict[str, dict[str, float]] = {}

    _, metin = _sayfa_metni_bul(pdf, "Özet Finansal Veriler")
    for metrik, desen in _OZET_DESENLERI.items():
        m = desen.search(metin)
        if not m:
            raise RuntimeError(f"THY sunumunda '{metrik}' için Özet Finansal Veriler satırı bulunamadı")
        noktalar[metrik] = {onceki_tarih: tr_sayi(m.group(1)), bu_tarih: tr_sayi(m.group(2))}

    idx, _ = _sayfa_metni_bul(pdf, "Birim Gelir Gelişimi")
    onceki, simdiki = _passenger_rask_degerleri(pdf.pages[idx])
    noktalar["yolcu-rask"] = {onceki_tarih: onceki, bu_tarih: simdiki}

    _, metin = _sayfa_metni_bul(pdf, "Operasyonel Gider Kırılımı")
    m = _GIDER_DESENI.search(metin)
    if not m:
        raise RuntimeError("THY sunumunda Operasyonel Gider Kırılımı / Akaryakıt satırı bulunamadı")
    noktalar["yakit-gideri"] = {onceki_tarih: tr_sayi(m.group(1)), bu_tarih: tr_sayi(m.group(2))}

    # Başlık ("Birim Gider (CASK) Kırılımı") pdfplumber'da bindirmeli font
    # efektiyle bozuk geliyor (ölçüldü); "CASK2 (AKTK dahil)" veri satırı
    # etiketi bozulmadan geliyor ve bu sayfaya özgü.
    _, metin = _sayfa_metni_bul(pdf, "CASK2 (AKTK dahil)")
    for metrik, desen in _CASK_DESENLERI.items():
        m = desen.search(metin)
        if not m:
            raise RuntimeError(f"THY sunumunda '{metrik}' için CASK Kırılımı satırı bulunamadı")
        noktalar[metrik] = {onceki_tarih: tr_sayi(m.group(1)), bu_tarih: tr_sayi(m.group(2))}

    _, metin = _sayfa_metni_bul(pdf, "Önemli Performans Göstergeleri (KPI)")
    m = _KPI_DESENI.search(metin)
    if not m:
        raise RuntimeError("THY sunumunda KPI / Akaryakıt (USD/ton) satırı bulunamadı")
    noktalar["yakit-fiyati"] = {onceki_tarih: tr_sayi(m.group(1)), bu_tarih: tr_sayi(m.group(2))}

    # "Bilanço" kelimesi CAPEX sayfasının dipnotunda da geçiyor (ölçüldü);
    # tablo başlığına özgü "Varlıklar (mn USD)" kullanılır.
    _, metin = _sayfa_metni_bul(pdf, "Varlıklar (mn USD)")
    noktalar.update(_bilanco_noktalari(metin))

    _, metin = _sayfa_metni_bul(pdf, "İtibarıyla)")
    for metrik, deger in _filo_noktalari(metin).items():
        noktalar[metrik] = {bu_tarih: deger}

    _, metin = _sayfa_metni_bul(pdf, "Dolar Bazında Bölgesel Birim Gelir Değişimi")
    for metrik, deger in _bolgesel_noktalari(metin).items():
        noktalar[metrik] = {bu_tarih: deger}

    idx, _ = _sayfa_metni_bul(pdf, "EBITDA Oranı")
    noktalar["net-borc-favok"] = _net_borc_favok_noktalari(pdf.pages[idx], yil, ceyrek)

    eksik = GECERLI_THY_METRIKLERI - set(noktalar)
    if eksik:
        raise RuntimeError(f"THY sunumunda çıkarılamayan metrikler: {sorted(eksik)}")
    return noktalar


def _sunum_noktalarini_getir(onbellek: dict, session=None) -> dict[str, dict[str, float]]:
    """PDF indirme + ayrıştırmayı önbellekler: 30 seri aynı tek dosyayı
    okuduğu için yoksa 30 indirme olurdu."""
    if "noktalar" in onbellek:
        return onbellek["noktalar"]
    http = session or requests
    url, yil, ceyrek = sunum_bilgisi(session)
    yanit = http.get(url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"THY yatırımcı sunumu indirilemedi ({url}): HTTP {yanit.status_code}")
    with pdfplumber.open(io.BytesIO(yanit.content)) as pdf:
        onbellek["noktalar"] = _sunumu_ayristir(pdf, yil, ceyrek)
    return onbellek["noktalar"]


def sunum_seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """THY çeyreklik yatırımcı sunumundan tek bir `thy_metrik` serisini
    çeker. `onbellek` verilirse PDF indirme + ayrıştırma koşu boyunca
    paylaşılır (bkz. `_sunum_noktalarini_getir`)."""
    onbellek = {} if onbellek is None else onbellek
    tum_noktalar = _sunum_noktalarini_getir(onbellek, session=session)
    noktalar = tum_noktalar.get(seri.thy_metrik)
    if not noktalar:
        raise RuntimeError(f"THY sunumunda thy_metrik bulunamadı: {seri.thy_metrik!r} ({seri.id})")
    df = pd.DataFrame(sorted(noktalar.items()), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
