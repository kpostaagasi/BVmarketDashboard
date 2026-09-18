"""Turkcell Yatırımcı İlişkileri çeyreklik "Financial and Operational Data"
Excel istemcisi.

Kaynak gerçekleri 2026-09-18'de canlı ölçüldü:

- Çeyrek arşivi `CEYREK_SAYFASI`nın Next.js `__NEXT_DATA__` JSON'unda
  `{year, quarter, content: [{title, url}, ...]}` biçiminde tam olarak
  duruyor (sunucu tarafında render ediliyor, JS gerekmez) — 2015 Ç1'den
  bugüne kesintisiz. Dosya URL'i DOSYA ADINDAN türetilemez: çoğu çeyrek
  `Q{çeyrek}{yy}-FO-Info.xlsx` kalıbında ama 2025 Ç3 ters sıralı
  (`FO-Info-Q325.xlsx`), 2023 ve 2024'ün Ç4/yıl-sonu dosyaları
  `FY23-FO-Info.xlsx`/`FY24-FO-Info.xlsx` (2025 Ç4 yine `Q425-...`'e
  dönüyor). Bu yüzden URL, dosya adı eşleştirilerek değil, o dönemin
  `content` listesindeki başlığı "Financial and Operational Data" İÇEREN
  girdiden okunur.
- EN GÜNCEL dosyanın kendisi TEK BAŞINA tüm pencereyi taşıyor (TAV'a
  benzer, dosya geçmişi birleştirilmez): "Operational KPIs" ve
  "ARPU_IAS29" sayfaları Q122'den bugüne, "Segment Revenue-EBITDA_Unadj"
  ve "Revenue Breakdown_Unadj" Q124'ten (segment bazlı raporlama o
  tarihte başladı — daha eski çeyrek yok, uydurulmaz), "Subsidiaries_Unadj"
  Q123'ten bugüne uzanıyor.
- "ARPU_IAS29" sayfası enflasyon düzeltmeli (IAS29) değerleri taşır; bu,
  referans sitedeki ARPU kartlarının BİRİNCİL (latestVal) değeridir.
  Sayfanın "Nominal" karşılığı ayrı bir kart/seri DEĞİL — aynı grafiğin
  ikinci (karşılaştırma) çizgisi olarak duruyor, katalogda temsil
  edilmiyor.
- "Mobil ARPU (M2M Hariç) vs TÜFE — YoY" kartı ARPU büyümesini TÜFE ile
  aynı grafikte üst üste koyan KOMPOZİT bir görsel (referans sitede ayrı
  bir bölüm); tek bir ham hücreye karşılık gelmediği için katalog dışı
  bırakıldı.
- Ölçüldü (2026-Ç2, referans marketvisuals.net/tcell.html ile birebir):
  Mobil Postpaid Abone 32,5 mn; Turkcell Fiber 2.625,7 bin; Superbox/5G
  818,1 bin; Mobil Prepaid 7,5 mn; Mobil M2M 6,1 mn (ve ek olarak IPTV,
  Resell, churn, ARPU, segment gelir/FAVÖK, Paycell/Financell, KKTC/BeST
  serilerinin tamamı ilgili hücrelerle birebir doğrulandı).
"""

from __future__ import annotations

import io
import json
import re

import openpyxl
import pandas as pd
import requests

TABAN = "https://www.turkcell.com.tr"
DOSYA_TABANI = "https://ffo3gv1cf3ir.merlincdn.net"
CEYREK_SAYFASI = f"{TABAN}/en-en/about-us/investor-relations/quarterly-results"
ZAMAN_ASIMI = 60
_BASLIKLAR = {"User-Agent": "Mozilla/5.0"}

_CEYREK_DESENI = re.compile(r"^Q([1-4])(\d{2})$")
_CEYREK_AY = {1: "01", 2: "04", 3: "07", 4: "10"}
_NEXT_DATA = re.compile(
    r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>', re.S
)

# metrik_id -> (sayfa_adi, blok_ipucu_ya_da_None, satır_etiketi). Blok ipucu
# yalnızca aynı etiketin sayfa içinde birden fazla bloktA tekrarlandığı
# durumlarda gerekli (ör. "Turkcell Türkiye" hem Revenue hem EBITDA
# tablosunda, "Revenue (million TRY)" hem BeST hem Paycell hem Financell
# bloğunda geçiyor).
METRIK_KAYNAK: dict[str, tuple[str, str | None, str]] = {
    "mobil-postpaid-abone": ("Operational KPIs", None, "Mobile postpaid subscribers (million)"),
    "fiber-abone": ("Operational KPIs", None, "Turkcell Fiber (thousand)"),
    "superbox-abone": ("Operational KPIs", None, "Superbox (thousand)3"),
    "mobil-prepaid-abone": ("Operational KPIs", None, "Mobile prepaid subscribers (million)"),
    "mobil-m2m-abone": ("Operational KPIs", None, "Mobile M2M (million)"),
    "iptv-abone": ("Operational KPIs", None, "IPTV (thousand)"),
    "resell-sabit-genisbant-abone": ("Operational KPIs", None, "Resell Fixed Broadband (thousand)2"),
    "mobil-churn": ("Operational KPIs", None, "Mobile churn (%)"),
    "sabit-churn": ("Operational KPIs", None, "Fixed churn (%)"),
    "mobil-arpu-m2m-haric": ("ARPU_IAS29", None, "Mobile ARPU, blended (TRY) (excluding M2M)"),
    "residential-fiber-arpu": ("ARPU_IAS29", None, "Residential Fiber ARPU"),
    "mobil-arpu-blended": ("ARPU_IAS29", None, "Mobile ARPU, blended (TRY)"),
    "postpaid-arpu-m2m-haric": ("ARPU_IAS29", None, "Postpaid (excluding M2M)"),
    "prepaid-arpu": ("ARPU_IAS29", None, "Prepaid (TRY)"),
    "turkiye-segment-geliri": ("Segment Revenue-EBITDA_Unadj", "Revenue Breakdown", "Turkcell Türkiye"),
    "techfin-segment-geliri": ("Segment Revenue-EBITDA_Unadj", "Revenue Breakdown", "Techfin"),
    "turkiye-segment-favok": ("Segment Revenue-EBITDA_Unadj", "EBITDA Breakdown", "Turkcell Türkiye"),
    "techfin-segment-favok": ("Segment Revenue-EBITDA_Unadj", "EBITDA Breakdown", "Techfin"),
    "tuketici-geliri": ("Revenue Breakdown_Unadj", None, "Consumer Revenues"),
    "kurumsal-geliri": ("Revenue Breakdown_Unadj", None, "Corporate Revenues"),
    "toptan-geliri": ("Revenue Breakdown_Unadj", None, "Wholesale"),
    "paycell-geliri": ("Subsidiaries_Unadj", "Paycell Summary Data", "Revenue (million TRY)"),
    "financell-geliri": ("Subsidiaries_Unadj", "Financell Summary Data", "Revenue (million TRY)"),
    "kktc-geliri": ("Subsidiaries_Unadj", "Kuzey Kıbrıs Summary Data", "Revenue"),
    "kktc-abone": ("Subsidiaries_Unadj", "Kuzey Kıbrıs Summary Data", "Number of subscribers (million)"),
    "best-geliri": ("Subsidiaries_Unadj", "BeST Summary Data", "Revenue (million TRY)"),
    "best-abone": ("Subsidiaries_Unadj", "BeST Summary Data", "Total"),
}


def _blok_sinirlari(ws) -> list[tuple[str | None, int, int]]:
    """Sayfayı, B sütununda 'Summary Data' ile biten ya da 'Breakdown
    (Unadjusted)' içeren başlık satırlarına göre bloklara ayırır:
    `[(başlık, ilk_satır, son_satır), ...]`. Böyle bir başlık yoksa (ör.
    "Operational KPIs", "ARPU_IAS29") tüm sayfa TEK blok sayılır.
    """
    basliklar = []
    for r in range(1, ws.max_row + 1):
        deger = ws.cell(row=r, column=2).value
        if isinstance(deger, str) and (
            deger.endswith("Summary Data") or "Breakdown (Unadjusted)" in deger
        ):
            basliklar.append((deger, r))
    if not basliklar:
        return [(None, 1, ws.max_row)]
    return [
        (baslik, r, (basliklar[i + 1][1] - 1 if i + 1 < len(basliklar) else ws.max_row))
        for i, (baslik, r) in enumerate(basliklar)
    ]


def _seriyi_bul(ws, etiket: str, blok_ipucu: str | None) -> dict[str, float]:
    """`(sayfa, blok_ipucu, etiket)` üçlüsünün TÜM çeyreklerini
    `{"Q226": değer, ...}` olarak döner; bulunamazsa boş sözlük."""
    for baslik, ilk, son in _blok_sinirlari(ws):
        if blok_ipucu and (baslik is None or blok_ipucu not in baslik):
            continue
        baslik_satiri = None
        for r in range(ilk, min(ilk + 3, son) + 1):
            if sum(1 for h in ws[r] if isinstance(h.value, str) and _CEYREK_DESENI.match(h.value)) >= 2:
                baslik_satiri = r
                break
        if baslik_satiri is None:
            continue
        sutunlar = {
            h.value: h.column for h in ws[baslik_satiri]
            if isinstance(h.value, str) and _CEYREK_DESENI.match(h.value)
        }
        for r in range(baslik_satiri + 1, son + 1):
            if ws.cell(row=r, column=2).value == etiket:
                return {
                    donem: float(ws.cell(row=r, column=sutun).value)
                    for donem, sutun in sutunlar.items()
                    if isinstance(ws.cell(row=r, column=sutun).value, (int, float))
                }
    return {}


def _ceyrek_tarihi(etiket: str) -> str:
    eslesme = _CEYREK_DESENI.match(etiket)
    ceyrek, yil2 = eslesme.groups()
    return f"20{yil2}-{_CEYREK_AY[int(ceyrek)]}-01"


def _donem_kayitlarini_bul(next_data: dict) -> list[dict]:
    """`__NEXT_DATA__` ağacında `{year, quarter, content}` biçimli listeyi
    arar. Next.js sorgu indeksleri sayfa değiştikçe kayabildiği için
    indekse değil, kayıt ŞEKLİNE göre aranır."""
    for sorgu in next_data.get("props", {}).get("pageProps", {}).get("dehydratedState", {}).get("queries", []):
        kayitlar = sorgu.get("state", {}).get("data")
        if (
            isinstance(kayitlar, list) and kayitlar
            and isinstance(kayitlar[0], dict)
            and {"year", "quarter", "content"} <= kayitlar[0].keys()
        ):
            return kayitlar
    raise RuntimeError(
        "Turkcell çeyrek sonuçları sayfasında dönem listesi bulunamadı — "
        "sayfa yapısı değişmiş olabilir"
    )


def en_guncel_dosya_url(session=None) -> str:
    """En güncel çeyreğin "Financial and Operational Data" xlsx URL'i."""
    http = session or requests
    yanit = http.get(CEYREK_SAYFASI, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"Turkcell çeyrek sonuçları sayfası HTTP {yanit.status_code}")
    eslesme = _NEXT_DATA.search(yanit.text)
    if not eslesme:
        raise RuntimeError(
            "Turkcell çeyrek sonuçları sayfasında __NEXT_DATA__ bulunamadı — "
            "sayfa yapısı değişmiş olabilir"
        )
    kayitlar = _donem_kayitlarini_bul(json.loads(eslesme.group(1)))
    en_guncel = max(kayitlar, key=lambda k: (k["year"], k["quarter"]))
    icerikler = [
        ic["url"] for ic in en_guncel.get("content", [])
        if "Financial and Operational Data" in ic.get("title", "")
    ]
    if not icerikler:
        raise RuntimeError(
            f"Turkcell {en_guncel['year']} Ç{en_guncel['quarter']}: "
            "'Financial and Operational Data' bağlantısı yok"
        )
    url = icerikler[0]
    return url if url.startswith("http") else DOSYA_TABANI + url


def _kitabi_getir(onbellek: dict, session=None):
    if "kitap" in onbellek:
        return onbellek["kitap"]
    http = session or requests
    url = en_guncel_dosya_url(session)
    yanit = http.get(url, timeout=ZAMAN_ASIMI, headers=_BASLIKLAR)
    if yanit.status_code != 200:
        raise RuntimeError(f"Turkcell F&O Excel'i indirilemedi ({url}): HTTP {yanit.status_code}")
    onbellek["kitap"] = openpyxl.load_workbook(io.BytesIO(yanit.content), data_only=True)
    return onbellek["kitap"]


def seri_cek(seri, onbellek: dict | None = None, session=None) -> pd.DataFrame:
    """Tam pencereyi yeniden çeker (artımlı değil — revizyonlar yakalanmalı).

    `onbellek` verilirse tek dosya (27 serinin tamamı aynı çeyrek Excel'ini
    paylaşır) koşu boyunca bir kez indirilir/ayrıştırılır.
    """
    onbellek = {} if onbellek is None else onbellek
    kitap = _kitabi_getir(onbellek, session=session)
    sayfa_adi, blok_ipucu, etiket = METRIK_KAYNAK[seri.turkcell_metrik]
    if sayfa_adi not in kitap.sheetnames:
        raise RuntimeError(f"Turkcell F&O Excel'inde '{sayfa_adi}' sayfası yok")
    ceyrekler = _seriyi_bul(kitap[sayfa_adi], etiket, blok_ipucu)
    if not ceyrekler:
        raise RuntimeError(
            f"Turkcell {sayfa_adi}: '{etiket}' satırı bulunamadı "
            f"(blok={blok_ipucu}, seri={seri.id}) — şablon değişmiş olabilir"
        )
    noktalar = [(_ceyrek_tarihi(donem), deger) for donem, deger in ceyrekler.items()]
    df = pd.DataFrame(sorted(noktalar), columns=["date", "value"])
    if seri.start_date:
        df = df[df["date"] >= seri.start_date]
    return df.reset_index(drop=True)
